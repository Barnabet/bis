"""Independently reopen frozen public data and check the manifest's observations.

This is a bounded source/report reconciliation checker, not a domain adapter or
a report-generation score. It reads local originals only, does not import the
application, and never evaluates formulas or follows workbook external links.
The caller is responsible for verifying original-file hashes and report bytes.
"""
from __future__ import annotations

import calendar
import csv
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
import re

from openpyxl import load_workbook


def _json(value):
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, (list, tuple)):
        return [_json(item) for item in value]
    if isinstance(value, dict):
        return {str(key): _json(item) for key, item in value.items()}
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _number(value):
    if isinstance(value, bool) or value is None:
        raise ValueError(f'Expected a finite numeric value, got {value!r}')
    try:
        result = Decimal(str(value))
    except InvalidOperation as exc:
        raise ValueError(f'Expected a finite numeric value, got {value!r}') from exc
    if not result.is_finite():
        raise ValueError('Non-finite numeric value')
    return result


def _round(value):
    return value.quantize(Decimal('0.1'), rounding=ROUND_HALF_UP)


class _Results:
    def __init__(self):
        self.data = {'checks': [], 'errors': [], 'reconciliation_findings': []}

    def check(self, kind, identifier, actual, expected, location=None):
        item = {'kind': kind, 'id': identifier, 'location': location,
                'actual': _json(actual), 'expected': _json(expected),
                'passed': actual == expected}
        self.data['checks'].append(item)
        if not item['passed']:
            self.data['errors'].append({**item, 'message': 'Source control mismatch'})

    def numeric(self, kind, identifier, actual, expected, location=None):
        self.check(kind, identifier, _number(actual), _number(expected), location)

    def error(self, identifier, exc):
        self.data['errors'].append({'id': identifier, 'error_type': type(exc).__name__,
                                    'message': str(exc)})

    def finding(self, identifier, source, printed, location):
        delta = _number(source) - _number(printed)
        if delta:
            self.data['reconciliation_findings'].append({
                'id': identifier, 'location': location,
                'kind': 'published_data_reconciliation_discrepancy',
                'recomputed_source_value': _json(source),
                'published_report_value': _json(printed), 'difference': str(delta),
                'status': 'unresolved_source_discrepancy',
                'cause': 'Not established; no rounding tolerance is assumed.',
            })


# Explicit table contracts, independently transcribed from the published tables.
# These locators select level inputs; expected business values come only from
# frozen report observations in the manifest, never from application output.
_CENSUS_LEVELS = {
    'sales_usd_billions': ('J12', None),
    'month_on_month_pct': ('J12', 'K12'),
    'year_on_year_pct': ('J12', 'M12'),
    'retail_month_on_month_pct': ('J17', 'K17'),
    'retail_year_on_year_pct': ('J17', 'M17'),
    'nonstore_year_on_year_pct': ('J64', 'M64'),
    'food_services_year_on_year_pct': ('J67', 'M67'),
}


def _census(book, case, out):
    levels, changes = book['Table 1.'], book['Table 2.']
    year, month = map(int, case['period'].split('-'))
    for sheet, cell, phrase in [
        (levels, 'A4', 'millions of dollars'), (changes, 'A3', 'shown as percents')
    ]:
        out.check('identity', f'{sheet.title}:units', phrase in str(sheet[cell].value),
                  True, f'{sheet.title}!{cell}')
    for cell, expected in [('J6', 'Adjusted2'), ('J7', year), ('M7', year - 1),
                           ('J9', '(a)'), ('K9', '(p)'), ('L9', '(r)'),
                           ('M9', '(r)'), ('N9', '(r)')]:
        out.check('identity', f'levels:{cell}', levels[cell].value, expected,
                  f'Table 1.!{cell}')
    for cell, offset in [('J8', 0), ('K8', -1), ('L8', -2), ('M8', 0), ('N8', -1)]:
        actual = re.match(r'[A-Za-z]+', str(levels[cell].value).strip())
        wanted = calendar.month_abbr[(month + offset - 1) % 12 + 1]
        out.check('identity', f'levels:month:{cell}', actual[0] if actual else None,
                  wanted, f'Table 1.!{cell}')
    out.check('identity', 'changes:current_period',
              bool(re.match(rf'{calendar.month_abbr[month]}\.?\s+{year}\s+Advance$',
                            str(changes['C8'].value))), True, 'Table 2.!C8')
    # Category identity prevents a cell at the right address in a shifted or
    # repurposed table from being accepted merely because its number matches.
    for sheet, row, label, naics in [
        (levels, 11, 'Retail & food services', None),
        (levels, 12, 'total', None), (levels, 17, 'Retail', None),
        (levels, 64, 'Nonstore retailers', '454'),
        (levels, 67, 'Food services & drinking places', '722'),
        (changes, 14, 'Retail & food services', None),
        (changes, 15, 'total', None), (changes, 21, 'Retail', None),
        (changes, 51, 'Nonstore retailers', '454'),
        (changes, 53, 'Food services & drinking places', '722'),
    ]:
        text = str(sheet[f'B{row}'].value).strip().rstrip(' .,…')
        out.check('identity', f'{sheet.title}:category:{row}', text, label,
                  f'{sheet.title}!B{row}')
        if naics:
            out.check('identity', f'{sheet.title}:naics:{row}',
                      str(sheet[f'A{row}'].value), naics, f'{sheet.title}!A{row}')
    for control in case['verified_source_checks']:
        metric = control['metric']
        try:
            locator = control['xlsx_cell']
            address = f"{locator['sheet']}!{locator['cell']}"
            actual = _number(book[locator['sheet']][locator['cell']].value)
            out.numeric('frozen_source_value', metric, actual, locator['raw_value'], address)
            unit = 'USD billion' if metric == 'sales_usd_billions' else 'percent'
            out.check('identity', f'{metric}:unit', control['unit'], unit, address)
            display = _round(actual / 1000 if metric == 'sales_usd_billions' else actual)
            wanted = control['expected_from_published_report']
            out.numeric('report_numeric_control', metric, display, wanted, address)
            if metric in _CENSUS_LEVELS:
                current_cell, prior_cell = _CENSUS_LEVELS[metric]
                current = _number(levels[current_cell].value)
                if prior_cell is None:
                    recomputed = _round(current / 1000)
                else:
                    prior = _number(levels[prior_cell].value)
                    if prior == 0:
                        raise ValueError('Zero comparison level makes the percentage undefined')
                    recomputed = _round((current / prior - 1) * 100)
                out.numeric('independent_arithmetic', metric, recomputed, wanted,
                            {'sheet': 'Table 1.', 'current': current_cell, 'prior': prior_cell})
            elif metric != 'rolling_three_month_year_on_year_pct':
                raise ValueError(f'Unknown Census control {metric}')
        except (ValueError, KeyError, TypeError, ArithmeticError) as exc:
            out.error(metric, exc)


def _ons(path, case, out):
    with path.open(encoding='utf-8-sig', newline='') as stream:
        rows = list(csv.reader(stream))
    for index, label in [(0, 'Title'), (1, 'CDID'), (3, 'Unit'), (4, 'Release Date')]:
        out.check('identity', f'metadata_row:{index + 1}', rows[index][0], label,
                  f'CSV row {index + 1}')
    expected_date = datetime.strptime(case['release_date'], '%Y-%m-%d').strftime('%d-%m-%Y')
    vintage = case['vintage_verification']
    out.check('identity', 'latest_period', rows[-1][0], vintage['latest_csv_month'], 'CSV final row')
    out.check('identity', 'row_count', len(rows), vintage['csv_rows_including_metadata'])
    out.check('identity', 'series_count', len(rows[1]) - 1, vintage['series_count'])
    for control in case['verified_source_checks']:
        identifier = f"{control['cdid']}:{control['csv_row_label']}"
        try:
            column = control['csv_column_1_based'] - 1
            row = control['csv_row_1_based'] - 1
            location = {'row_1_based': row + 1, 'column_1_based': column + 1}
            out.check('identity', f'{identifier}:cdid', rows[1][column], control['cdid'], location)
            out.check('identity', f'{identifier}:unique_cdid', rows[1].count(control['cdid']), 1)
            out.check('identity', f'{identifier}:period', rows[row][0], control['csv_row_label'], location)
            out.check('identity', f'{identifier}:unique_period',
                      sum(bool(r) and r[0] == control['csv_row_label'] for r in rows), 1)
            out.check('identity', f'{identifier}:release_date', rows[4][column], expected_date, location)
            # MS6Y has a blank published Unit cell; its title identifies a share.
            expected_unit = '' if control['cdid'] == 'MS6Y' else '%'
            out.check('identity', f'{identifier}:unit', rows[3][column], expected_unit, location)
            title = control.get('series_title')
            if title is None:
                title = next(c['series_title'] for c in case['verified_source_checks']
                             if c['cdid'] == control['cdid'] and 'series_title' in c)
            out.check('identity', f'{identifier}:title', rows[0][column], title, location)
            out.numeric('frozen_source_value', identifier, rows[row][column], control['csv_value'], location)
            out.numeric('report_numeric_control', identifier, rows[row][column],
                        control['declared_report_value'], location)
        except (ValueError, KeyError, IndexError, TypeError, StopIteration) as exc:
            out.error(identifier, exc)


def _fmc(book, corpus, case, out):
    year, quarter = case['period'].split('-')
    period = f'{quarter} {year}'
    headers = ['Quarter, Year', 'Port Name', 'Laden Exports', 'Empty Exports',
               'Laden Imports', 'Empty Imports', 'Export Tonnage', 'Import Tonnage']
    out.check('identity', 'US Ports:headers', list(next(book['US Ports'].values)), headers)
    out.check('identity', 'Ocean Carriers:headers', list(next(book['Ocean Carriers'].values)),
              [headers[0], 'Carrier Name', *headers[2:]])
    for control in case['verified_source_checks']:
        identifier = f"{period}:{control['port']}:{control['cell']}"
        try:
            cell = book['US Ports'][control['cell']]
            out.check('identity', f'{identifier}:manifest_period', control['quarter'], period)
            out.check('identity', f'{identifier}:quarter', book['US Ports'].cell(cell.row, 1).value,
                      period, f'US Ports!A{cell.row}')
            out.check('identity', f'{identifier}:port', book['US Ports'].cell(cell.row, 2).value,
                      control['port'], f'US Ports!B{cell.row}')
            out.numeric('frozen_source_value', identifier, cell.value, control['actual'], f'US Ports!{cell.coordinate}')
            out.numeric('report_numeric_control', identifier, cell.value, control['pdf_printed'],
                        f'US Ports!{cell.coordinate}')
        except (ValueError, KeyError, TypeError) as exc:
            out.error(identifier, exc)
    matches = [r for r in corpus['source_reconciliation'] if r['quarter'] == period]
    if len(matches) != 1:
        raise ValueError(f'Expected one reconciliation record for {period}')
    reconciliation = matches[0]
    printed_containers = reconciliation['pdf_printed_container_totals']
    printed_tonnage = reconciliation['pdf_printed_tonnage_totals']
    for sheet, reference in reconciliation['sheet_checks'].items():
        try:
            selected = [(n, row) for n, row in enumerate(book[sheet].values, 1)
                        if row[0] == period]
            row_spec = reconciliation['source_rows'][sheet]
            actual_range = {'first_row': selected[0][0] if selected else None,
                            'last_row': selected[-1][0] if selected else None, 'count': len(selected)}
            out.check('identity', f'{sheet}:{period}:source_rows', actual_range, row_spec)
            totals = [Decimal(0)] * 6
            for row_number, row in selected:
                for offset, value in enumerate(row[2:8]):
                    number = _number(value)
                    if number < 0 or number != number.to_integral_value():
                        raise ValueError(f'{sheet} row {row_number} contains a non-integer or negative freight count')
                    totals[offset] += number
            location = {'sheet': sheet, 'quarter': period, 'rows': actual_range}
            out.check('reconciliation_numeric', f'{sheet}:{period}:container_sums', totals[:4],
                      [_number(n) for n in reference['container_category_sums']], location)
            out.numeric('reconciliation_numeric', f'{sheet}:{period}:container_grand_sum',
                        sum(totals[:4]), reference['container_grand_sum'], location)
            out.check('reconciliation_numeric', f'{sheet}:{period}:tonnage_sums', totals[4:],
                      [_number(n) for n in reference['tonnage_sums']], location)
            for index, actual in enumerate(totals[:4]):
                identifier = f'{sheet}:{period}:{headers[index + 2]}'
                out.numeric('reconciliation_numeric', identifier + ':difference',
                            actual - _number(printed_containers[index]),
                            reference['container_category_difference_vs_pdf'][index], location)
                out.finding(identifier, actual, printed_containers[index], location)
            out.numeric('reconciliation_numeric', f'{sheet}:{period}:grand_difference',
                        sum(totals[:4]) - _number(printed_containers[4]),
                        reference['container_grand_difference_vs_pdf'], location)
            out.finding(f'{sheet}:{period}:container_grand_total',
                        sum(totals[:4]), printed_containers[4], location)
            for index, actual in enumerate([*totals[4:], sum(totals[4:])]):
                identifier = f'{sheet}:{period}:tonnage:{index}'
                out.numeric('reconciliation_numeric', identifier, actual, printed_tonnage[index], location)
                out.finding(identifier, actual, printed_tonnage[index], location)
        except (ValueError, KeyError, TypeError, IndexError) as exc:
            out.error(f'{sheet}:{period}:reconciliation', exc)
    category_sum = sum(_number(n) for n in printed_containers[:4])
    difference = category_sum - _number(printed_containers[4])
    out.numeric('reconciliation_numeric', f'{period}:printed_categories_vs_grand', difference,
                reconciliation['published_container_categories_sum_minus_printed_grand_total'])
    out.finding(f'{period}:printed_container_categories_vs_printed_grand', category_sum,
                printed_containers[4], {'report': case['report_asset'], 'quarter': period})


def source_value_checks(corpus, case, source_path):
    """Return JSON-safe checks, errors and unresolved reconciliation findings.

    ``corpus`` and ``case`` are records from fixtures/public-corpus/manifest.json;
    ``source_path`` is a local original CSV/XLSX. No network or model calls occur.
    A nonempty findings list means published totals do not fully reconcile even
    when every supported source control passes. Missing/nonnumeric cells fail;
    neither blanks, suppression markers nor formula strings become zero.
    """
    out = _Results()
    try:
        if not case.get('verified_source_checks'):
            raise ValueError('No frozen source controls are available')
        path = Path(source_path)
        if corpus['id'] == 'ons_retail':
            _ons(path, case, out)
        elif corpus['id'] in {'census_marts', 'fmc_freight'}:
            # Formula text is exposed, not evaluated, and external relationships
            # are discarded. Numeric controls therefore cannot trust caches.
            book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
            try:
                if corpus['id'] == 'census_marts':
                    _census(book, case, out)
                else:
                    _fmc(book, corpus, case, out)
            finally:
                book.close()
        else:
            raise ValueError(f"Unsupported corpus {corpus['id']}")
    except Exception as exc:
        out.error(case.get('id', 'source'), exc)
    return out.data
