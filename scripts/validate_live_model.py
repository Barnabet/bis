#!/usr/bin/env python3
"""Opt-in, synthetic end-to-end validation against a real local CLIProxyAPI.

Run with the project's Python environment and explicitly configured provider:
    uv run python scripts/validate_live_model.py --run-live

This is not part of the default test suite. It creates a fresh isolated workspace,
resolves only predeclared synthetic fixture decisions, and makes at most three
logical model jobs / five real Responses requests (including bounded repairs).
It never reads credential files, prints headers, substitutes model output, changes
protocols, or accepts the resulting commentary on behalf of a human reviewer.
"""
from __future__ import annotations

import argparse
import copy
import csv
from datetime import datetime, timezone
from decimal import Decimal
import hashlib
from io import BytesIO, StringIO
import json
from pathlib import Path
import sys
from uuid import uuid4


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_ROOT = ROOT / "output"
EXPECTED_MODEL = "claude-opus-5"
MAX_CALLS = 5
QUOTE = "The renewal campaign drove the reported movement."
TERMINAL = {"completed", "blocked", "failed", "cancelled"}


def utc_now():
    return datetime.now(timezone.utc).isoformat()


def require(condition, message):
    if not condition:
        raise AssertionError(message)


def json_bytes(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, indent=2,
                      allow_nan=False).encode("utf-8")


def sha256(data):
    return hashlib.sha256(data).hexdigest()


def sanitized(value):
    """Defense in depth; captured request bodies contain synthetic evidence only."""
    if isinstance(value, dict):
        return {key: ("[REDACTED]" if key.casefold().replace("-", "_") in
                      {"authorization", "api_key", "access_token", "credential", "password"}
                      else sanitized(item)) for key, item in value.items()}
    if isinstance(value, list):
        return [sanitized(item) for item in value]
    return value


class Report:
    def __init__(self, output, data):
        self.output = output
        self.summary = {
            "started_at": utc_now(), "status": "running", "output": str(output),
            "workspace": str(data), "synthetic_only": True,
            "semantic_review_required": True,
            "semantic_review_status": "not_performed",
            "scope": "Real local provider, worker, evaluation and exporters; predeclared synthetic policy decisions only.",
            "limits": {"logical_model_jobs": 3, "outbound_responses_calls": MAX_CALLS},
            "transport_calls": [], "jobs": [], "checks": [], "retries": [],
        }
        self.save()

    def save(self):
        self.summary["updated_at"] = utc_now()
        temporary = self.output / "summary.json.tmp"
        temporary.write_bytes(json_bytes(sanitized(self.summary)))
        temporary.replace(self.output / "summary.json")

    def stage(self, name):
        self.summary["stage"] = name
        self.save()
        print(name, flush=True)

    def check(self, name, condition):
        self.summary["checks"].append({"name": name, "passed": bool(condition)})
        self.save()
        require(condition, name)

    def artifact(self, name, value):
        path = self.output / name
        path.parent.mkdir(parents=True, exist_ok=True)
        data = value if isinstance(value, bytes) else json_bytes(value)
        path.write_bytes(data)
        return {"path": str(path), "sha256": sha256(data), "size_bytes": len(data)}


class TransportAudit:
    """Observe and delegate the real transport; never inspect or save headers."""
    def __init__(self, report, canary):
        self.report, self.canary = report, canary
        self.original = None
        self.requests = []
        self.network_allowed = True

    def __enter__(self):
        from foundry import model_provider
        self.original = model_provider._responses_transport
        model_provider._responses_transport = self.call
        return self

    def __exit__(self, *_):
        from foundry import model_provider
        model_provider._responses_transport = self.original

    @property
    def count(self):
        return len(self.requests)

    def call(self, body, headers, **options):
        from foundry.errors import DomainError
        if not self.network_allowed:
            raise DomainError("LIVE_CACHE_MISS", "The cache replay attempted a new provider call.")
        if self.count >= MAX_CALLS:
            raise DomainError("LIVE_CALL_BUDGET", "The live validation request ceiling was reached.")
        request = json.loads(body)
        require(request.get("model") == EXPECTED_MODEL, "Outbound request changed the requested model")
        require(self.canary.casefold() not in body.decode("utf-8").casefold(),
                "Reserved canary reached the outbound request boundary")
        require(request.get("store") is False and request.get("tools") == [],
                "Outbound request must disable storage and tools")
        index = self.count + 1
        entry = {"number": index, "started_at": utc_now(), "request_sha256": sha256(body),
                 "model": request["model"], "schema_name": request["text"]["format"]["name"],
                 "origin": {key: options.get(key) for key in ("scheme", "host", "port")},
                 "headers_recorded": False, "status": "dispatched"}
        entry["sanitized_request"] = self.report.artifact(
            f"requests/call-{index:02d}.json", sanitized(request))
        self.requests.append({"body": request, "sha256": sha256(body)})
        self.report.summary["transport_calls"].append(entry)
        self.report.save()
        try:
            result = self.original(body, headers, **options)
            entry.update(status="response_received", response_sha256=sha256(result),
                         response_size_bytes=len(result))
            return result
        except Exception as exc:
            entry.update(status="failed", error_code=getattr(exc, "code", type(exc).__name__))
            raise
        finally:
            entry["finished_at"] = utc_now()
            self.report.save()


def period(label, year, quarter):
    month = (quarter - 1) * 3 + 1
    end_year, end_month = (year + 1, 1) if quarter == 4 else (year, month + 3)
    prior_end_year = end_year - 1
    return {"label": label, "start": f"{year}-{month:02d}-01",
            "end_exclusive": f"{end_year}-{end_month:02d}-01",
            "comparison": {"start": f"{year - 1}-{month:02d}-01",
                           "end_exclusive": f"{prior_end_year}-{end_month:02d}-01"},
            "timezone": "Europe/Paris", "as_of": f"{end_year}-{end_month:02d}-03T12:00:00+02:00"}


def csv_data(rows):
    out = StringIO(newline="")
    writer = csv.writer(out)
    writer.writerow(["transaction_id", "date", "region", "status", "amount", "currency"])
    writer.writerows(rows)
    return out.getvalue().encode("utf-8")


def reserved_target(canary):
    """Write expected observations manually; never use candidate preparation."""
    from docx import Document
    from foundry.runtime import DISCLOSURE
    document = Document()
    document.add_heading("Quarterly revenue", 0)
    for text in ("Current revenue (EUR): 160.00", "Comparison revenue (EUR): 100.00",
                 "Revenue change (EUR): +60.00", "Growth: 60.0%", f"Selected region: {canary}",
                 DISCLOSURE):
        document.add_paragraph(text)
    document.add_heading("Regional performance", 1)
    table = document.add_table(rows=1, cols=5)
    for cell, text in zip(table.rows[0].cells,
                          ["Region", "Current (EUR)", "Comparison (EUR)", "Change (EUR)", "Growth"]):
        cell.text = text
    for cell, text in zip(table.add_row().cells, [canary, "160.00", "100.00", "+60.00", "60.0%"]):
        cell.text = text
    out = BytesIO()
    document.save(out)
    return out.getvalue()


def run_job(service, worker, requested, report):
    report.stage(f"Running {requested['kind']} job {requested['id']}")
    require(requested["status"] == "queued", "A fresh operation must enqueue a new job")
    require(worker.run_one(), "The worker did not claim the requested job")
    completed = service.store.job(requested["id"])
    safe = {key: completed.get(key) for key in
            ("id", "kind", "status", "stage", "steps", "result", "error", "created_at", "updated_at")}
    report.summary["jobs"].append(safe)
    report.save()
    require(completed["status"] in TERMINAL, "Worker did not finish the expected job")
    require(completed["status"] == "completed",
            f"{completed['kind']} job did not complete: {(completed.get('error') or {}).get('code', completed['status'])}")
    return completed


def receipt_check(report, receipt, label):
    report.check(f"{label}: exact provider and model receipt",
                 receipt.get("provider") == "cliproxyapi" and receipt.get("model") == EXPECTED_MODEL
                 and receipt.get("requested_model") == EXPECTED_MODEL and receipt.get("protocol") == "responses")
    report.check(f"{label}: receipt has request and response identities",
                 all(isinstance(receipt.get(key), str) and receipt[key] for key in
                     ("response_id", "prompt_digest", "response_digest", "output_digest")))


def retry_check(service, worker, report, transport, completed, request_again, label):
    before = (transport.count, len(service.store.list("model_call")))
    retried = request_again()
    require(not worker.run_one(), "Idempotency retry unexpectedly queued another operation")
    after = (transport.count, len(service.store.list("model_call")))
    report.summary["retries"].append({"label": label, "original_job_id": completed["id"],
                                       "returned_job_id": retried["id"], "before": before, "after": after})
    report.check(f"{label}: completed idempotency retry uses the original result without another model call",
                 retried["id"] == completed["id"] and retried["status"] == "completed"
                 and retried["result"] == completed["result"] and before == after)


def check_snapshot_delta(report, before, after, label):
    report.check(f"{label}: immutable facts and datasets",
                 after["facts"] == before["facts"] and after["datasets"] == before["datasets"])
    old_nodes = [node for node in before["nodes"] if node["id"] != "commentary"]
    new_nodes = [node for node in after["nodes"] if node["id"] != "commentary"]
    report.check(f"{label}: non-commentary nodes, views and program are unchanged",
                 old_nodes == new_nodes and after["views"] == before["views"] and after["program"] == before["program"])
    report.check(f"{label}: separate review-required revision",
                 after["id"] != before["id"] and after["parent_id"] == before["id"]
                 and after["revision"] == before["revision"] + 1 and after["status"] == "review_required")
    codes = {finding["code"] for finding in after["findings"]}
    report.check(f"{label}: explicit review finding and no blocking findings",
                 "COMPOSITION_REVIEW_REQUIRED" in codes and not any(f["severity"] == "block" for f in after["findings"]))


def cache_replay(service, report, transport, job, snapshot):
    from foundry.composition import INSTRUCTIONS
    from foundry.storage import digest
    receipt = snapshot["metadata"]["composition"]["receipt"]
    captured = next(entry["body"] for entry in transport.requests if entry["sha256"] == receipt["prompt_digest"])
    payload = json.loads(captured["input"][0]["content"][0]["text"])
    schema = captured["text"]["format"]["schema"]
    name = captured["text"]["format"]["name"]
    identity = digest({"job_id": job["id"], "instructions": INSTRUCTIONS, "payload": payload,
                       "schema": schema, "name": name, "provider": job["payload"]["model_config"]})
    record = service.store.get("model_call", "model_" + identity)
    expected = json.loads(service.store.read_blob(record["response_digest"]))
    before = (transport.count, len(service.store.list("model_call")))
    transport.network_allowed = False
    try:
        replay = service._captured_provider(job).generate(INSTRUCTIONS, payload, schema, name)
    finally:
        transport.network_allowed = True
    after = (transport.count, len(service.store.list("model_call")))
    report.summary["cache_replay"] = {"job_id": job["id"], "model_call_id": record["id"],
                                      "before": before, "after": after, "network_dispatch_disabled": True}
    report.check("Completed composition response replays through host cache without network or added records",
                 replay == expected and before == after)


def validate(service, worker, report, transport, canary):
    from docx import Document
    from foundry.observations import inspect_target
    from foundry.runtime import render_runs

    report.stage("Registering independent synthetic historical examples")
    created = service.create_type("Live model synthetic validation", "Isolated, disposable CLIProxyAPI validation.", demo=True)
    report_type, program = created["report_type"], created["program"]
    report.summary["report_type_id"] = report_type["id"]
    report.summary["program_id"] = program["id"]
    examples = []
    for name, role in (("ambiguous", "authoring"), ("discriminating", "development")):
        directory = ROOT / "fixtures" / "learning" / name
        target = service.upload((directory / "report.docx").read_bytes(), f"{name}-report.docx", demo=True)
        source = service.upload((directory / "transactions.csv").read_bytes(), f"{name}-transactions.csv", demo=True)
        case_period = json.loads((directory / "period.json").read_text())
        examples.append(service.add_example(report_type["id"], target["id"], [source["id"]], case_period,
                                             role, f"Synthetic {name}", "Independently authored fixture; not customer evidence."))
    target_bytes = reserved_target(canary)
    reserve_period = period("Q3 2024", 2024, 3)
    source_bytes = csv_data([
        ["reserved-current", "2024-07-15", canary, "posted", "160.00", "EUR"],
        ["reserved-comparison", "2023-07-15", canary, "posted", "100.00", "EUR"],
    ])
    report.summary["reserved_fixture"] = {
        "canary": canary, "expected": {"current": "160.00", "comparison": "100.00", "change": "60.00", "growth": "60.0%"},
        "target": report.artifact("inputs/reserved-report.docx", target_bytes),
        "source": report.artifact("inputs/reserved-transactions.csv", source_bytes),
    }
    target = service.upload(target_bytes, "reserved-report.docx", demo=True)
    source = service.upload(source_bytes, "reserved-transactions.csv", demo=True)
    inspection = inspect_target(target)
    report.check("Manually authored reserved target has complete supported coverage",
                 bool(inspection["regions"]) and all(region["status"] == "mapped" for region in inspection["regions"])
                 and not inspection["issues"])
    reserved = service.add_example(report_type["id"], target["id"], [source["id"]], reserve_period,
                                   "reserved", "Synthetic reserved reconstruction", "Evaluator-only canary example.")
    report.summary["examples"] = [{key: example.get(key) for key in
                                   ("id", "digest", "corpus_role", "report_asset_id", "source_asset_ids", "period")}
                                  for example in [*examples, reserved]]
    report.check("Reserved example is hidden from ordinary source inspection",
                 reserved["inspection"] is None and service.asset_view(target["id"]).get("reserved")
                 and service.asset_view(source["id"]).get("reserved"))

    requirements = ("Use posted EUR transactions, explicit current and comparison periods, the supplied regional table, "
                    "and the region selected by the policy supported by the examples. Preserve the causal disclosure. "
                    "Do not approve decisions or publish the candidate.")
    initial_digest = program["digest"]
    request_learning = lambda: service.request_learning(program["id"], initial_digest, requirements, "openai", "live-learning")
    learned_job = run_job(service, worker, request_learning(), report)
    program = service.store.get("program", program["id"])
    learning = program["learning"]
    report.summary["learning"] = copy.deepcopy(learning)
    report.artifact("learning/program-before-review.json", program)
    report.save()
    receipt_check(report, learning["model_receipt"], "Learning")
    report.check("Only absolute-change policy is supported by both historical examples",
                 {h["value"] for h in learning["hypotheses"] if h["supported"]} == {"largest_absolute_change"}
                 and len(learning["hypotheses"]) == 3)
    report.check("Real model proposes the supported absolute-change policy",
                 learning["model_proposal"]["selection"] == "largest_absolute_change")
    report.check("Learning preserves unresolved human decisions and does not publish",
                 all(decision["resolution"] is None for decision in program["decisions"])
                 and program["state"] == "candidate" and program["evaluation"] is None
                 and service.store.get("report_type", report_type["id"])["active_program_id"] is None)
    report.check("Reserved content is absent from learned evidence and model context",
                 canary.casefold() not in json.dumps(learning).casefold()
                 and reserved["id"] not in learning["example_ids"]
                 and set(learning["example_ids"]) == {example["id"] for example in examples})
    retry_check(service, worker, report, transport, learned_job, request_learning, "Learning")

    report.stage("Reviewing predeclared synthetic choices and running real evaluation")
    report.check("All authoring/development historical regions are mapped",
                 all(region["status"] == "mapped" for region in learning["coverage"]))
    choices = {
        "selection": ("largest_absolute_change", "The two independent fixtures support this implemented policy only."),
        "completeness": ("assumed_complete", "These synthetic sources explicitly contain every regional record for each period."),
        "template": ("compatible_reviewed", "Use the disclosed compatible renderer profile; native application certification is separate."),
        "requirements": ("supported_scope_reviewed", "The scripted requirements fit the posted-EUR regional reporting adapter."),
    }
    report.summary["fixture_decisions"] = []
    for decision in list(program["decisions"]):
        require(decision["id"] in choices, f"Unreviewed new decision: {decision['id']}")
        value, rationale = choices[decision["id"]]
        require(value in {item["value"] for item in decision["alternatives"]}, "Declared fixture resolution is unavailable")
        program = service.resolve(program["id"], decision["id"], value, program["digest"])
        report.summary["fixture_decisions"].append({"id": decision["id"], "question": decision["question"],
                                                    "resolution": value, "rationale": rationale,
                                                    "actor_context": "Predeclared automated synthetic-fixture review"})
        report.save()
    before_non_model = (transport.count, len(service.store.list("model_call")))
    program = service.evaluate(program["id"], program["digest"])
    report.summary["evaluation"] = program["evaluation"]
    report.save()
    report.check("Real evaluation passes including two historical cases and the reserved canary",
                 program["evaluation"]["passed"] and program["evaluation"]["holdout_count"] == 1
                 and program["evaluation"]["reconstruction_count"] == 3)
    program = service.publish(program["id"], program["digest"], actor="synthetic_live_validation")
    report.summary["published_program"] = {key: program[key] for key in ("id", "version", "digest", "state")}

    report.stage("Generating an independent Q2 period through the published program")
    new_rows = []
    for region, prior, current in [("East", 10, 100), ("North", 1100, 1200), ("South", 100, 250), ("West", 200, 600)]:
        new_rows.extend([[f"q2-{region}-current", "2026-04-15", region, "posted", f"{current}.00", "EUR"],
                         [f"q2-{region}-comparison", "2025-04-15", region, "posted", f"{prior}.00", "EUR"]])
    new_bytes = csv_data(new_rows)
    report.summary["new_period_source"] = report.artifact("inputs/new-period-transactions.csv", new_bytes)
    source = service.upload(new_bytes, "new-period-transactions.csv", demo=True)
    generated_job = run_job(service, worker, service.request_run(report_type["id"], source["id"],
                            period("Q2 2026", 2026, 2), "live-new-period"), report)
    snapshot = service.store.get("snapshot", generated_job["result"]["snapshot_id"])
    report.artifact("snapshots/generated.json", snapshot)
    facts = snapshot["facts"]
    report.summary["new_period_facts"] = {key: {field: facts[key][field] for field in ("value", "display")}
                                          for key in ("revenue.current", "revenue.comparison", "revenue.change", "revenue.growth", "driver.region")}
    report.check("Independent new-period totals and selected region match manually specified expectations",
                 Decimal(facts["revenue.current"]["value"]) == Decimal("2150")
                 and Decimal(facts["revenue.comparison"]["value"]) == Decimal("1410")
                 and Decimal(facts["revenue.change"]["value"]) == Decimal("740")
                 and facts["revenue.growth"]["display"] == "52.5%" and facts["driver.region"]["value"] == "West")
    report.check("Evaluation, publication and new-period generation made no model calls",
                 before_non_model == (transport.count, len(service.store.list("model_call"))))

    report.summary["compositions"] = []
    for label in ("facts-only", "cited-note"):
        if label == "facts-only":
            source_ids = []
            objective = ("Summarize posted revenue and the selected region using provided fact placeholders. Include "
                         "revenue.current, revenue.comparison, revenue.change, revenue.growth and driver.region. "
                         "Use the phrase selected region. Keep quantities in fact references and omit causal interpretations.")
        else:
            document = Document()
            document.add_paragraph(QUOTE)
            note_bytes = BytesIO()
            document.save(note_bytes)
            note = service.upload(note_bytes.getvalue(), "current-period-renewal-note.docx", demo=True)
            report.summary["current_period_note"] = {"asset_id": note["id"], "digest": note["digest"],
                "artifact": report.artifact("inputs/current-period-renewal-note.docx", note_bytes.getvalue()),
                "status": "Synthetic attributed claim; not independently verified causal evidence"}
            source_ids = [note["id"]]
            objective = ("Summarize posted revenue and the selected region using fact placeholders. Include this exact "
                         f'quotation: "{QUOTE}" Introduce it in the same paragraph with "According to the supplied evidence" '
                         "and cite its supplied evidence ID. Present it as the source's claim, not an established cause.")
        parent = copy.deepcopy(snapshot)
        request_composition = lambda: service.request_composition(parent["id"], parent["revision"], source_ids,
                                                                   objective, f"live-{label}")
        completed = run_job(service, worker, request_composition(), report)
        snapshot = service.store.get("snapshot", completed["result"]["snapshot_id"])
        composition = snapshot["metadata"]["composition"]
        node = next(node for node in snapshot["nodes"] if node["id"] == "commentary")
        prose = render_runs(node["runs"], snapshot["facts"])
        report.summary["compositions"].append({"label": label, "snapshot_id": snapshot["id"],
            "parent_id": parent["id"], "revision": snapshot["revision"], "resolved_prose": prose,
            "runs": node["runs"], "composition": composition, "findings": snapshot["findings"],
            "semantic_review_required": True,
            "artifact": report.artifact(f"snapshots/{label}.json", snapshot),
            "prose_artifact": report.artifact(f"compositions/{label}.txt", prose.encode("utf-8"))})
        report.save()
        check_snapshot_delta(report, parent, snapshot, label)
        receipt_check(report, composition["receipt"], label)
        for attempt in composition["attempts"]:
            receipt_check(report, attempt["receipt"], f"{label} attempt {attempt['attempt']}")
        report.check(f"{label}: original parent snapshot remains unchanged", service.store.get("snapshot", parent["id"]) == parent)
        if label == "facts-only":
            report.check("Facts-only composition has no qualitative citations or extra source bindings",
                         composition["evidence"] == [] and composition["evidence_refs"] == []
                         and snapshot["source_assets"] == parent["source_assets"])
        else:
            evidence = composition["evidence"]
            evidence_ids = {item["id"] for item in evidence}
            report.check("Cited composition uses only the requested current-period source and exact source locators",
                         bool(evidence) and all(item["asset_id"] == note["id"] and item["artifact_sha256"] == note["digest"]
                         and item["locator"] for item in evidence)
                         and bool(composition["evidence_refs"]) and set(composition["evidence_refs"]) <= evidence_ids
                         and {asset["id"] for asset in snapshot["source_assets"]} == {source["id"], note["id"]})
            report.check("Causal quote is exact, attributed, cited and explicitly left for review",
                         QUOTE in prose and "According to the supplied evidence" in prose
                         and any(f["code"] == "CAUSAL_EVIDENCE_REVIEW" and f["severity"] == "review" for f in snapshot["findings"]))
        retry_check(service, worker, report, transport, completed, request_composition, label)
        if label == "facts-only":
            cache_replay(service, report, transport, completed, snapshot)

    report.stage("Exporting the final unaccepted revision in four real formats")
    before_exports = (transport.count, len(service.store.list("model_call")))
    report.summary["exports"] = []
    for format in ("docx", "xlsx", "pdf", "pptx"):
        completed = run_job(service, worker, service.request_export(snapshot["id"], format, "compatible", f"live-export-{format}"), report)
        export = service.store.get("export", completed["result"]["export_id"])
        data = service.store.read_blob(export["digest"])
        report.summary["exports"].append({"record": export, "artifact": report.artifact(f"exports/final-report.{format}", data)})
        report.check(f"{format.upper()}: real artifact binds the final draft snapshot",
                     len(data) > 100 and sha256(data) == export["digest"] and export["snapshot_id"] == snapshot["id"]
                     and export["snapshot_digest"] == snapshot["digest"] and export["status"] == "draft")
        report.check(f"{format.upper()}: export adds no transport or persisted model calls",
                     before_exports == (transport.count, len(service.store.list("model_call"))))
    model_records = service.store.list("model_call")
    model_jobs = [job for job in service.store.jobs() if job["kind"] in {"learning", "composition"}]
    report.summary["model_call_records"] = model_records
    report.summary["counts"] = {"logical_model_jobs": len(model_jobs), "transport_calls": transport.count,
                                 "persisted_model_calls": len(model_records)}
    report.check("Exactly three model jobs and at most five genuine Responses calls were used",
                 len(model_jobs) == 3 and 3 <= transport.count <= MAX_CALLS and len(model_records) == transport.count)
    report.check("Reserved canary never appeared in any outbound request",
                 all(canary.casefold() not in json.dumps(request["body"]).casefold() for request in transport.requests))
    report.check("Final report remains unaccepted and reserved example remains unexposed",
                 service.store.get("snapshot", snapshot["id"])["status"] == "review_required"
                 and service.example_view(reserved["id"])["inspection"] is None
                 and not service.store.list("example_exposure"))
    report.summary.update(status="passed_pending_semantic_review", finished_at=utc_now(),
                          final_snapshot_id=snapshot["id"], semantic_review_required=True,
                          review_notes=["Read both resolved commentaries and confirm meaning, scope and attribution.",
                                        "A matched causal quotation is not proof that its claim is true.",
                                        "Inspect the four export artifacts; native Office certification is separate."])
    report.save()


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--run-live", action="store_true", help="Explicitly permit the bounded real model calls.")
    parser.add_argument("--output", type=Path, help="Fresh result directory under this repository's ignored output/.")
    parser.add_argument("--data", type=Path, help="Fresh workspace under output/; defaults to <result directory>/workspace.")
    args = parser.parse_args(argv)
    if not args.run_live:
        parser.error("Live calls are opt-in. Configure CLIProxyAPI and claude-opus-5, then pass --run-live.")
    unique = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S.%fZ") + "-" + uuid4().hex[:6]
    output = (args.output or OUTPUT_ROOT / "cliproxyapi-live" / unique).expanduser().resolve()
    data = (args.data or output / "workspace").expanduser().resolve()
    for label, path in (("output", output), ("data", data)):
        if path == OUTPUT_ROOT.resolve() or not path.is_relative_to(OUTPUT_ROOT.resolve()):
            parser.error(f"--{label} must be a child directory of {OUTPUT_ROOT}.")
        if path.exists() and (not path.is_dir() or any(path.iterdir())):
            parser.error(f"Refusing to overwrite a nonempty or non-directory {label} path: {path}")
    if output == data or output.is_relative_to(data):
        parser.error("The workspace must not contain the result directory.")
    output.mkdir(parents=True, exist_ok=True)
    report = Report(output, data)
    # Import application dependencies only after explicit opt-in. No credential
    # files are read here; provider.status() exposes only redacted configuration.
    sys.path.insert(0, str(ROOT))
    try:
        from foundry import model_provider
        from foundry.jobs import Worker
        from foundry.service import Service
        from foundry.storage import Store
        model_status = model_provider.status()
        report.summary["model_status"] = model_status
        report.check("Explicitly configured CLIProxyAPI with exact claude-opus-5 and Responses protocol",
                     model_status["configured"] and model_status["provider"] == "cliproxyapi"
                     and model_status["model"] == EXPECTED_MODEL and model_status["protocol"] == "responses")
        service = Service(Store(data))
        worker = Worker(service)
        suffix = "".join(chr(ord("a") + int(char, 16)) for char in uuid4().hex[:10]).title()
        canary = "Reserved Canary Cedar " + suffix
        with TransportAudit(report, canary) as transport:
            validate(service, worker, report, transport, canary)
        print(f"Structural checks passed; semantic review remains required. Results: {output / 'summary.json'}")
        return 0
    except Exception as exc:
        report.summary.update(status="failed", finished_at=utc_now(), error={
            "type": type(exc).__name__, "code": getattr(exc, "code", None),
            "message": str(exc) if isinstance(exc, AssertionError) else
                getattr(exc, "message", "Validation failed; inspect the saved stage and job records.")})
        report.save()
        print(f"Validation stopped. Partial results: {output / 'summary.json'}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
