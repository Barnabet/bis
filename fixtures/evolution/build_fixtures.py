"""Author synthetic independent fixtures; never import or run Foundry.

Run with the bundled artifact Python. Monetary totals, changes, exact growth
fractions, displayed percentages and winners below are declared by the fixture
author. Arithmetic is checked independently before any artifact is written.
Future documents must be frozen before the live experiment; do not regenerate
them in response to an implementation result.
"""
from __future__ import annotations

import csv
import hashlib
import json
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, ROUND_HALF_UP, localcontext
from fractions import Fraction
from pathlib import Path
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
DISCLOSURE = "This report describes posted revenue movements. The source data does not establish the causes of those movements."

# Each row is region, current EUR, comparison EUR, declared change EUR,
# declared exact growth fraction (None for a zero comparison), displayed growth.
# Each total has the same fields after omitting region. Winners are separately
# declared rather than selected by this builder.
CORPORA = [
    {
        "id": "current_revenue", "expected_selection": "largest_current_revenue",
        "future_canary": "Dune Canary",
        "cases": [
            ("history_q3_2025", 2025, 3, [
                ("Atrium", "200", "100", "100", "1", "100.0%"),
                ("Briar", "60", "50", "10", "1/5", "20.0%"),
                ("Cypress", "15", "10", "5", "1/2", "50.0%"),
            ], ("275", "160", "115", "23/32", "71.9%"), ("Atrium", "100"), ["ambiguous_history"]),
            ("history_q4_2025", 2025, 4, [
                ("Atrium", "1000", "990", "10", "1/99", "1.0%"),
                ("Briar", "500", "100", "400", "4", "400.0%"),
                ("Cypress", "100", "10", "90", "9", "900.0%"),
            ], ("1600", "1100", "500", "5/11", "45.5%"), ("Atrium", "10"), ["discriminating_history"]),
            ("future_q1_2026", 2026, 1, [
                ("Atrium", "800", "1000", "-200", "-1/5", "-20.0%"),
                ("Briar", "650", "700", "-50", "-1/14", "-7.1%"),
                ("Cypress", "120", "100", "20", "1/5", "20.0%"),
            ], ("1570", "1800", "-230", "-23/180", "-12.8%"), ("Atrium", "-200"), ["overall_decline", "selected_region_decline", "mixed_direction"]),
            ("future_q2_2026", 2026, 2, [
                ("Atrium", "500", "100", "400", "4", "400.0%"),
                ("Briar", "500", "200", "300", "3/2", "150.0%"),
                ("Cypress", "125", "500", "-375", "-3/4", "-75.0%"),
            ], ("1125", "800", "325", "13/32", "40.6%"), ("Atrium", "400"), ["selection_tie", "normalized_region_tie_break"]),
            ("future_q3_2026", 2026, 3, [
                ("Atrium", "100", "200", "-100", "-1/2", "-50.0%"),
                ("Briar", "90", "60", "30", "1/2", "50.0%"),
                ("Cypress", "0", "15", "-15", "-1", "-100.0%"),
                ("Dune Canary", "300", "0", "300", None, "Not defined"),
            ], ("490", "275", "215", "43/55", "78.2%"), ("Dune Canary", "300"), ["new_region", "disappeared_region", "undefined_regional_growth", "historical_comparison_reused"]),
            ("future_q4_2026", 2026, 4, [
                ("Atrium", "1000.25", "1000", "0.25", "1/4000", "0.0%"),
                ("Briar", "1500.50", "500", "1000.50", "2001/1000", "200.1%"),
                ("Cypress", "0", "100", "-100", "-1", "-100.0%"),
                ("Dune Canary", "50.25", "0", "50.25", None, "Not defined"),
            ], ("2551.00", "1600", "951.00", "951/1600", "59.4%"), ("Briar", "1000.50"), ["fractional_currency", "split_transactions", "historical_comparison_reused"]),
        ],
    },
    {
        "id": "absolute_change", "expected_selection": "largest_absolute_change",
        "future_canary": "Orchard Canary",
        "cases": [
            ("history_q3_2025", 2025, 3, [
                ("Harbor", "300", "150", "150", "1", "100.0%"),
                ("Juniper", "80", "40", "40", "1", "100.0%"),
                ("Larch", "10", "8", "2", "1/4", "25.0%"),
            ], ("390", "198", "192", "32/33", "97.0%"), ("Harbor", "150"), ["ambiguous_history"]),
            ("history_q4_2025", 2025, 4, [
                ("Harbor", "2000", "1980", "20", "1/99", "1.0%"),
                ("Juniper", "900", "200", "700", "7/2", "350.0%"),
                ("Larch", "160", "10", "150", "15", "1500.0%"),
            ], ("3060", "2190", "870", "29/73", "39.7%"), ("Juniper", "700"), ["discriminating_history"]),
            ("future_q1_2026", 2026, 1, [
                ("Harbor", "900", "1000", "-100", "-1/10", "-10.0%"),
                ("Juniper", "100", "900", "-800", "-8/9", "-88.9%"),
                ("Larch", "300", "100", "200", "2", "200.0%"),
            ], ("1300", "2000", "-700", "-7/20", "-35.0%"), ("Juniper", "-800"), ["overall_decline", "selected_region_decline", "mixed_direction"]),
            ("future_q2_2026", 2026, 2, [
                ("Harbor", "100", "400", "-300", "-3/4", "-75.0%"),
                ("Juniper", "600", "300", "300", "1", "100.0%"),
                ("Larch", "50", "25", "25", "1", "100.0%"),
            ], ("750", "725", "25", "1/29", "3.4%"), ("Harbor", "-300"), ["selection_tie", "normalized_region_tie_break", "tie_opposite_signs"]),
            ("future_q3_2026", 2026, 3, [
                ("Harbor", "330", "300", "30", "1/10", "10.0%"),
                ("Juniper", "0", "80", "-80", "-1", "-100.0%"),
                ("Larch", "15", "10", "5", "1/2", "50.0%"),
                ("Orchard Canary", "250", "0", "250", None, "Not defined"),
            ], ("595", "390", "205", "41/78", "52.6%"), ("Orchard Canary", "250"), ["new_region", "disappeared_region", "undefined_regional_growth", "historical_comparison_reused"]),
            ("future_q4_2026", 2026, 4, [
                ("Harbor", "1900.10", "2000", "-99.90", "-999/20000", "-5.0%"),
                ("Juniper", "900.20", "900", "0.20", "1/4500", "0.0%"),
                ("Larch", "260.30", "160", "100.30", "1003/1600", "62.7%"),
                ("Orchard Canary", "1000.40", "0", "1000.40", None, "Not defined"),
            ], ("4061.00", "3060", "1001.00", "1001/3060", "32.7%"), ("Orchard Canary", "1000.40"), ["fractional_currency", "split_transactions", "historical_comparison_reused"]),
        ],
    },
    {
        "id": "percentage_change", "expected_selection": "largest_percentage_change",
        "future_canary": "Zephyr Canary",
        "cases": [
            ("history_q3_2025", 2025, 3, [
                ("Maple", "400", "200", "200", "1", "100.0%"),
                ("Quartz", "100", "80", "20", "1/4", "25.0%"),
                ("Willow", "30", "20", "10", "1/2", "50.0%"),
            ], ("530", "300", "230", "23/30", "76.7%"), ("Maple", "200"), ["ambiguous_history"]),
            ("history_q4_2025", 2025, 4, [
                ("Maple", "1500", "1490", "10", "1/149", "0.7%"),
                ("Quartz", "700", "200", "500", "5/2", "250.0%"),
                ("Willow", "120", "10", "110", "11", "1100.0%"),
            ], ("2320", "1700", "620", "31/85", "36.5%"), ("Willow", "110"), ["discriminating_history"]),
            ("future_q1_2026", 2026, 1, [
                ("Maple", "800", "1000", "-200", "-1/5", "-20.0%"),
                ("Quartz", "300", "400", "-100", "-1/4", "-25.0%"),
                ("Willow", "10", "100", "-90", "-9/10", "-90.0%"),
            ], ("1110", "1500", "-390", "-13/50", "-26.0%"), ("Willow", "-90"), ["overall_decline", "selected_region_decline", "all_regions_decline"]),
            ("future_q2_2026", 2026, 2, [
                ("Maple", "100", "200", "-100", "-1/2", "-50.0%"),
                ("Quartz", "300", "200", "100", "1/2", "50.0%"),
                ("Willow", "110", "100", "10", "1/10", "10.0%"),
            ], ("510", "500", "10", "1/50", "2.0%"), ("Maple", "-100"), ["selection_tie", "normalized_region_tie_break", "tie_opposite_signs"]),
            ("future_q3_2026", 2026, 3, [
                ("Maple", "440", "400", "40", "1/10", "10.0%"),
                ("Quartz", "0", "100", "-100", "-1", "-100.0%"),
                ("Willow", "45", "30", "15", "1/2", "50.0%"),
                ("Zephyr Canary", "600", "0", "600", None, "Not defined"),
            ], ("1085", "530", "555", "111/106", "104.7%"), ("Quartz", "-100"), ["new_region", "disappeared_region", "undefined_regional_growth_excluded_from_selection", "historical_comparison_reused"]),
            ("future_q4_2026", 2026, 4, [
                ("Maple", "1600.50", "1500", "100.50", "67/1000", "6.7%"),
                ("Quartz", "1050.25", "700", "350.25", "1401/2800", "50.0%"),
                ("Willow", "60", "120", "-60", "-1/2", "-50.0%"),
                ("Zephyr Canary", "900.25", "0", "900.25", None, "Not defined"),
            ], ("3611.00", "2320", "1291.00", "1291/2320", "55.6%"), ("Quartz", "350.25"), ["fractional_currency", "split_transactions", "exact_ratio_beats_equal_display_magnitude", "historical_comparison_reused"]),
        ],
    },
]


def ratio_decimal(value):
    if value is None:
        return None
    rational = Fraction(value)
    with localcontext() as ctx:
        ctx.prec = 40
        return format(Decimal(rational.numerator) / Decimal(rational.denominator), "f")


def named_values(values):
    current, comparison, change, ratio, display = values
    return {"current": current, "comparison": comparison, "change": change,
            "growth_fraction": ratio, "growth_value": ratio_decimal(ratio),
            "growth_display": display}


def validate_values(values):
    current, comparison, change, ratio, display = values
    c, p, d = map(Fraction, (current, comparison, change))
    assert c >= 0 and p >= 0 and c - p == d, values
    if p == 0:
        assert ratio is None and display == "Not defined", values
    else:
        assert Fraction(ratio) == d / p, values
        rounded = (Decimal(ratio_decimal(ratio)) * 100).quantize(Decimal("0.1"), rounding=ROUND_HALF_UP)
        assert f"{rounded:.1f}%" == display, values


def winner(rows, policy):
    # Independent Fraction ranking is validation of the declared winner only.
    choices = []
    for region, current, comparison, change, growth, _ in rows:
        if policy == "largest_current_revenue":
            metric = Fraction(current)
        elif policy == "largest_absolute_change":
            metric = abs(Fraction(change))
        elif growth is not None:
            metric = abs(Fraction(growth))
        else:
            continue
        choices.append((-metric, region.casefold(), region))
    return min(choices)[2]


def money(value):
    return f"{Decimal(value):,.2f}"


def write_docx(destination, period, expected):
    doc = Document()
    doc.core_properties.title = f"Quarterly revenue {period['label']}"
    doc.core_properties.subject = "Independent synthetic regional revenue reporting fixture"
    doc.core_properties.author = "Report Foundry validation fixtures"
    doc.core_properties.created = datetime(2026, 9, 13, tzinfo=timezone.utc)
    doc.core_properties.modified = datetime(2026, 9, 13, tzinfo=timezone.utc)
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.75)
    for style_name in ("Normal", "Title", "Heading 1"):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
    for border in doc.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    doc.styles["Normal"].font.size = Pt(10)
    doc.styles["Normal"].paragraph_format.space_after = Pt(7)
    doc.styles["Title"].font.size = Pt(22)
    doc.styles["Heading 1"].font.size = Pt(14)
    doc.add_paragraph(f"Quarterly revenue - {period['label']}", style="Title")
    totals = expected["totals"]
    doc.add_paragraph(f"Current revenue (EUR): {money(totals['current'])}")
    doc.add_paragraph(f"Comparison revenue (EUR): {money(totals['comparison'])}")
    doc.add_paragraph(f"Revenue change (EUR): {money(totals['change'])}")
    doc.add_paragraph(f"Growth: {totals['growth_display']}")
    doc.add_paragraph(f"Selected region: {expected['driver']['region']}")
    doc.add_paragraph("Regional performance", style="Heading 1")
    headers = ["Region", "Current (EUR)", "Comparison (EUR)", "Change (EUR)", "Growth (%)"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [1.35, 1.18, 1.35, 1.18, 1.15]
    for i, label in enumerate(headers):
        table.rows[0].cells[i].text = label
    for row in expected["regions"]:
        cells = table.add_row().cells
        values = [row["region"], money(row["current"]), money(row["comparison"]), money(row["change"]), row["growth_display"]]
        for cell, value in zip(cells, values):
            cell.text = value
    for ri, row in enumerate(table.rows):
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side in ("top", "left", "bottom", "right"):
                item = OxmlElement(f"w:{side}")
                item.set(qn("w:w"), "90")
                item.set(qn("w:type"), "dxa")
                margins.append(item)
            props.append(margins)
            borders = OxmlElement("w:tcBorders")
            for side in ("top", "left", "bottom", "right"):
                item = OxmlElement(f"w:{side}")
                item.set(qn("w:val"), "single")
                item.set(qn("w:sz"), "4")
                item.set(qn("w:color"), "D9D9D9")
                borders.append(item)
            props.append(borders)
            if ri == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "E4EBF2")
                props.append(shade)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.space_before = Pt(2)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.RIGHT
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    run.bold = ri == 0
    repeated = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeated)
    doc.add_paragraph("Editorial context", style="Heading 1")
    doc.add_paragraph(DISCLOSURE)
    doc.save(destination)


def make_period(year, quarter):
    month = 1 + 3 * (quarter - 1)
    end = date(year + (quarter == 4), 1 if quarter == 4 else month + 3, 1)
    prior_end = date(year - 1 + (quarter == 4), 1 if quarter == 4 else month + 3, 1)
    return {"label": f"Q{quarter} {year}", "start": date(year, month, 1).isoformat(),
            "end_exclusive": end.isoformat(),
            "comparison": {"start": date(year - 1, month, 1).isoformat(), "end_exclusive": prior_end.isoformat()},
            "timezone": "Europe/Paris", "as_of": datetime(end.year, end.month, end.day, 12, tzinfo=ZoneInfo("Europe/Paris")).isoformat()}


def write_source(destination, corpus_id, case_id, period, rows):
    with destination.open("w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["transaction_id", "date", "region", "status", "amount", "currency"])
        writer.writeheader()
        count = 0

        def emit(day, region, amount, status="posted"):
            nonlocal count
            count += 1
            writer.writerow({"transaction_id": f"{corpus_id}-{case_id}-{count:03}", "date": day,
                             "region": region, "status": status, "amount": f"{amount:.2f}", "currency": "EUR"})

        for window, field in ((period, 1), (period["comparison"], 2)):
            for index, row in enumerate(rows):
                amount = Decimal(row[field])
                if amount == 0:
                    continue  # True regional absence exercises the approved zero rule.
                start, end = date.fromisoformat(window["start"]), date.fromisoformat(window["end_exclusive"])
                first = (amount / 2).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
                region = f"  {row[0].swapcase()}  " if "q2_2026" in case_id and index == 0 else row[0]
                emit(start.isoformat(), region, first)
                emit((end - timedelta(days=1)).isoformat(), row[0], amount - first)
        # Large distractors would dominate every policy if exclusions failed.
        emit(period["start"], rows[0][0], Decimal("999999.99"), "cancelled")
        emit(period["comparison"]["start"], rows[-1][0], Decimal("888888.88"), "cancelled")
        emit(period["end_exclusive"], rows[0][0], Decimal("777777.77"))
        emit(period["comparison"]["end_exclusive"], rows[-1][0], Decimal("666666.66"))


def write_json(path, value):
    path.write_text(json.dumps(value, indent=2, ensure_ascii=False) + "\n")


def main():
    manifest = {"schema_version": 1, "corpora": []}
    policies = [corpus["expected_selection"] for corpus in CORPORA]
    arithmetic = []
    for corpus in CORPORA:
        entry = {key: corpus[key] for key in ("id", "expected_selection", "future_canary")}
        entry.update(history=[], future=[])
        supporting = set(policies)
        for case_id, year, quarter, rows, totals, selected, coverage in corpus["cases"]:
            for row in rows:
                validate_values(row[1:])
            validate_values(totals)
            assert sum(Fraction(row[1]) for row in rows) == Fraction(totals[0])
            assert sum(Fraction(row[2]) for row in rows) == Fraction(totals[1])
            assert [row[0] for row in rows] == sorted(row[0] for row in rows)
            chosen = next(row for row in rows if row[0] == selected[0])
            assert Fraction(chosen[3]) == Fraction(selected[1])
            assert winner(rows, corpus["expected_selection"]) == selected[0]
            if case_id.startswith("history"):
                supporting &= {policy for policy in policies if winner(rows, policy) == selected[0]}
            if year == 2026 and quarter in (3, 4):
                history = next(case for case in corpus["cases"] if case[1:3] == (2025, quarter))
                historical_current = {row[0]: Fraction(row[1]) for row in history[3]}
                future_comparison = {row[0]: Fraction(row[2]) for row in rows if Fraction(row[2]) != 0}
                assert historical_current == future_comparison
            period = make_period(year, quarter)
            expected = {"schema_version": 1, "case_id": f"{corpus['id']}/{case_id}",
                        "selection": corpus["expected_selection"], "totals": named_values(totals),
                        "driver": {"region": selected[0], **named_values(chosen[1:]), "change": selected[1]},
                        "regions": [{"region": row[0], **named_values(row[1:])} for row in rows],
                        "commentary": DISCLOSURE,
                        "coverage": coverage + ["split_transactions", "start_inclusive", "end_exclusive", "cancelled_excluded"],
                        "oracle": {"method": "manually_declared_values_checked_with_independent_fraction_arithmetic",
                                   "growth_decimal_precision": 40,
                                   "candidate_output_used": False, "runtime_imported": False}}
            folder = ROOT / corpus["id"] / case_id
            folder.mkdir(parents=True, exist_ok=True)
            write_json(folder / "expected.json", expected)
            write_json(folder / "period.json", period)
            write_source(folder / "transactions.csv", corpus["id"], case_id, period, rows)
            write_docx(folder / "report.docx", period, expected)
            record = {"id": case_id, **{key: str((folder / filename).relative_to(ROOT)) for key, filename in
                      (("report", "report.docx"), ("source", "transactions.csv"), ("period", "period.json"), ("expected", "expected.json"))}}
            entry["history" if case_id.startswith("history") else "future"].append(record)
            arithmetic.append({"case_id": expected["case_id"], "checks_passed": True,
                               "declared_driver": selected[0],
                               "independent_policy_winners": {policy: winner(rows, policy) for policy in policies}})
        assert supporting == {corpus["expected_selection"]}, (corpus["id"], supporting)
        manifest["corpora"].append(entry)
    write_json(ROOT / "manifest.json", manifest)
    write_json(ROOT / "arithmetic-check.json", {"schema_version": 1, "case_count": len(arithmetic),
               "method": "Fraction arithmetic and independent declared-winner ranking; no candidate output", "cases": arithmetic})
    paths = [ROOT / case[key] for corpus in manifest["corpora"] for split in ("history", "future") for case in corpus[split]
             for key in ("report", "source", "period", "expected")]
    write_json(ROOT / "frozen-sha256.json", {"schema_version": 1,
               "purpose": "Freeze all inputs and independent expectations before live authoring",
               "files": {str(path.relative_to(ROOT)): hashlib.sha256(path.read_bytes()).hexdigest() for path in sorted(paths)}})
    print(json.dumps({"corpora": len(manifest["corpora"]), "cases": len(arithmetic), "files_frozen": len(paths), "arithmetic_passed": True}))


if __name__ == "__main__":
    main()
