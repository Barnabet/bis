"""Integrated authoring/evaluation boundaries; model calls are explicit test doubles."""
from __future__ import annotations

import copy
from io import BytesIO
import json
from pathlib import Path
import shutil

import pytest
from fastapi.testclient import TestClient

from foundry.api import create_app
from foundry.errors import DomainError
from foundry.jobs import Worker
from foundry.runtime import default_period
from foundry.service import ROOT, Service, validate_stored_snapshot
from foundry.storage import Store


FIXTURES = ROOT / 'fixtures' / 'learning'


@pytest.fixture
def service(tmp_path):
    return Service(Store(tmp_path))


def _type(service):
    created = service.create_type('Historical policy reconstruction')
    return created['report_type'], created['program']


def _assets(service, name):
    folder = FIXTURES / name
    target = service.upload((folder / 'report.docx').read_bytes(), f'{name}-report.docx')
    source = service.upload((folder / 'transactions.csv').read_bytes(), f'{name}-transactions.csv')
    return target, source, json.loads((folder / 'period.json').read_text())


def _pair(service, report_type, name, role='authoring'):
    target, source, period = _assets(service, name)
    example = service.add_example(report_type['id'], target['id'], [source['id']], period, role)
    return example, target, source, period


def _finish(service, job):
    assert Worker(service).run_one()
    return service.store.job(job['id'])


def _learn(service, program, key='learn-test'):
    request = service.request_learning(program['id'], program['digest'], '', 'deterministic', key)
    completed = _finish(service, request)
    assert completed['status'] == 'completed', completed.get('error')
    return service.store.get('program', program['id']), completed


def _smoke_export(snapshot, format, output_dir, **kwargs):
    output_dir = Path(output_dir); output_dir.mkdir(parents=True, exist_ok=True)
    artifact = output_dir / f'lifecycle-smoke.{format}'
    artifact.write_bytes(b'Explicit lifecycle-only renderer test substitute. ' * 8)
    return {'path': str(artifact)}


def _review(service, program, selection='largest_absolute_change'):
    for decision in program['decisions']:
        value = selection if decision['id'] == 'selection' else decision['alternatives'][0]['value']
        program = service.resolve(program['id'], decision['id'], value, program['digest'])
    for region in program.get('learning', {}).get('coverage', []):
        if region['status'] not in {'mapped', 'out_of_scope'}:
            program = service.review_coverage(program['id'], program['digest'], region['id'], 'out_of_scope',
                                              'Fixture review explicitly excludes this unimplemented historical presentation region.')
    return program


def _evaluate(service, program, monkeypatch):
    with monkeypatch.context() as patch:
        patch.setattr('foundry.exporters.export_snapshot', _smoke_export)
        return service.evaluate(program['id'], program['digest'])


def _release(service, monkeypatch):
    report_type, program = _type(service)
    program = _evaluate(service, _review(service, program), monkeypatch)
    assert program['evaluation']['passed']
    return report_type, service.publish(program['id'], program['digest'])


def _snapshot(service, monkeypatch):
    report_type, _ = _release(service, monkeypatch)
    source = service.upload((ROOT / 'fixtures' / 'transactions.csv').read_bytes(), 'transactions.csv')
    job = service.request_run(report_type['id'], source['id'], default_period().model_dump(mode='json'), 'native-base')
    completed = _finish(service, job)
    assert completed['status'] == 'completed', completed.get('error')
    return service.store.get('snapshot', completed['result']['snapshot_id'])


def _docx(text):
    from docx import Document
    document = Document(); document.add_paragraph(text)
    out = BytesIO(); document.save(out)
    return out.getvalue()


class ModelDouble:
    def __init__(self):
        self.calls = []
    def status(self):
        return {'configured': True, 'provider': 'openai', 'model': 'test-only-model', 'limits': {'max_output_tokens': 2048}}
    def generate(self, instructions, payload, schema, name):
        self.calls.append(copy.deepcopy({'instructions': instructions, 'payload': payload, 'schema': schema, 'name': name}))
        if name == 'report_policy':
            output = {'selection': 'largest_absolute_change', 'rationale': 'Test-only provider proposal requiring human decision.',
                      'unresolved_questions': ['Confirm the consequential selection policy.']}
        else:
            output = {'paragraphs': [{'template': 'Posted revenue reached {{revenue.current}}, with {{driver.region}} selected for the regional commentary.',
                                      'evidence_refs': []}]}
        return {'output': output, 'receipt': {'provider': 'test_double', 'model': 'test-only-model',
                                             'response_id': f'test_response_{len(self.calls)}', 'usage': {'output_tokens': 25}}}


def _model(monkeypatch):
    model = ModelDouble()
    monkeypatch.setattr('foundry.model_provider.OpenAIProvider', lambda: model)
    return model


def test_example_pairs_pin_original_evidence_and_reject_invalid_bindings(service):
    report_type, _ = _type(service)
    target, source, period = _assets(service, 'ambiguous')
    example = service.add_example(report_type['id'], target['id'], [source['id']], period, 'authoring')
    assert example['asset_digests'] == {target['id']: target['digest'], source['id']: source['digest']}
    assert example['inspection']['observations']
    assert service.examples(report_type['id'])[0]['digest'] == example['digest']
    for target_id, source_ids in [(target['id'], [target['id']]), (target['id'], [source['id'], source['id']]),
                                  (source['id'], [target['id']])]:
        with pytest.raises(DomainError) as error:
            service.add_example(report_type['id'], target_id, source_ids, period, 'authoring')
        assert error.value.code == 'EXAMPLE_BINDING_INVALID'
    with pytest.raises(DomainError) as error:
        service.add_example(report_type['id'], target['id'], [source['id']], period, 'development')
    assert error.value.code == 'EXAMPLE_EXISTS'


def test_reserved_api_assets_and_same_bytes_alias_are_masked_until_reveal(tmp_path):
    with TestClient(create_app(tmp_path, start_worker=False)) as client:
        service = client.app.state.service
        report_type, _ = _type(service)
        example, target, source, _ = _pair(service, report_type, 'discriminating', 'reserved')
        alias = service.upload((FIXTURES / 'discriminating' / 'report.docx').read_bytes(), 'different-name.docx')
        original = service.store.get('example', example['id'])
        before = service.corpus(report_type['id'])['digest']
        for asset in [target, source, alias]:
            response = client.get(f'/api/assets/{asset["id"]}')
            assert response.status_code == 200 and response.json()['reserved'] is True
            assert response.json()['profile']['regions'] == []
            assert 'sample_rows' not in response.json()['profile']
            assert client.get(f'/api/assets/{asset["id"]}/download').status_code == 403
        bootstrap = client.get('/api/bootstrap').json()
        assert all(a.get('reserved') for a in bootstrap['assets'])
        assert client.get(f'/api/examples/{example["id"]}').json()['inspection'] is None
        revealed = client.post(f'/api/examples/{example["id"]}/reveal', json={'reason': 'Investigate this evaluated discrepancy.'})
        assert revealed.status_code == 200 and revealed.json()['effective_role'] == 'development'
        assert revealed.json()['inspection']['observations']
        assert service.corpus(report_type['id'])['digest'] != before
        assert service.store.get('example', example['id']) == original
        assert len(service.store.list('example_exposure')) == 1
        client.post(f'/api/examples/{example["id"]}/reveal', json={'reason': 'A repeated reveal.'})
        assert len(service.store.list('example_exposure')) == 1
        assert client.get(f'/api/assets/{alias["id"]}/download').content == (FIXTURES / 'discriminating' / 'report.docx').read_bytes()


def test_reserved_pair_cannot_overlap_authoring_or_be_reused_as_authoring(service):
    report_type, _ = _type(service)
    _, target, source, period = _pair(service, report_type, 'ambiguous')
    changed_period = copy.deepcopy(period); changed_period['label'] = 'Same evidence under another label'
    with pytest.raises(DomainError) as error:
        service.add_example(report_type['id'], target['id'], [source['id']], changed_period, 'reserved')
    assert error.value.code == 'CORPUS_OVERLAP'
    _, target, source, period = _pair(service, report_type, 'discriminating', 'reserved')
    with pytest.raises(DomainError) as error:
        service.add_example(report_type['id'], target['id'], [source['id']], period, 'authoring')
    assert error.value.code == 'RESERVED_EVIDENCE'


def test_learning_preserves_ambiguity_and_completed_request_identity(service):
    report_type, program = _type(service)
    _pair(service, report_type, 'ambiguous')
    prior_digest = program['digest']
    learned, job = _learn(service, program)
    assert len([h for h in learned['learning']['hypotheses'] if h['supported']]) == 3
    assert all(d['resolution'] is None for d in learned['decisions'])
    retry = service.request_learning(program['id'], prior_digest, '', 'deterministic', 'learn-test')
    assert retry['id'] == job['id'] and retry['status'] == 'completed'
    with pytest.raises(DomainError) as error:
        service.request_learning(program['id'], prior_digest, 'Different requirements', 'deterministic', 'learn-test')
    assert error.value.code == 'IDEMPOTENCY_CONFLICT'


@pytest.mark.parametrize('change', ['candidate', 'corpus'])
def test_learning_job_rejects_changed_candidate_or_corpus_before_execution(service, change):
    report_type, program = _type(service)
    _pair(service, report_type, 'ambiguous')
    job = service.request_learning(program['id'], program['digest'], '', 'deterministic', 'stale-learning')
    if change == 'candidate':
        service.resolve(program['id'], 'selection', 'largest_absolute_change', program['digest'])
    else:
        _pair(service, report_type, 'discriminating', 'development')
    completed = _finish(service, job)
    assert completed['status'] == 'blocked'
    assert completed['error']['code'] == ('VERSION_CONFLICT' if change == 'candidate' else 'CORPUS_CHANGED')
    assert 'learning' not in service.store.get('program', program['id'])


def test_cancelled_learning_cannot_commit_candidate_analysis(service):
    report_type, program = _type(service)
    _pair(service, report_type, 'ambiguous')
    queued = service.request_learning(program['id'], program['digest'], '', 'deterministic', 'cancel-learning')
    job = service.store.claim()
    def stage(name):
        if name == 'Recording candidate evidence and unresolved decisions':
            service.store.cancel(job['id'])
    with pytest.raises(DomainError) as error:
        service.execute_learning(job, stage)
    assert error.value.code == 'LEASE_LOST'
    assert service.store.job(queued['id'])['status'] == 'cancelled'
    assert service.store.get('program', program['id']) == program


def test_new_example_during_learning_fences_the_commit(service):
    report_type, program = _type(service)
    _pair(service, report_type, 'ambiguous')
    service.request_learning(program['id'], program['digest'], '', 'deterministic', 'changing-corpus')
    job = service.store.claim()
    def stage(name):
        if name == 'Recording candidate evidence and unresolved decisions':
            _pair(service, report_type, 'discriminating', 'development')
    with pytest.raises(DomainError) as error:
        service.execute_learning(job, stage)
    assert error.value.code == 'CORPUS_CHANGED'
    assert service.store.get('program', program['id']) == program


def test_distinguishing_examples_reconstruct_and_wrong_policy_fails(service, monkeypatch):
    report_type, program = _type(service)
    _pair(service, report_type, 'ambiguous')
    _pair(service, report_type, 'discriminating', 'development')
    learned, _ = _learn(service, program)
    assert [h['value'] for h in learned['learning']['hypotheses'] if h['supported']] == ['largest_absolute_change']
    wrong = _evaluate(service, _review(service, learned, 'largest_current_revenue'), monkeypatch)
    assert wrong['evaluation']['passed'] is False
    assert any(not c['passed'] and c['name'].startswith('Historical reconstruction') for c in wrong['evaluation']['checks'])
    correct = _evaluate(service, _review(service, wrong), monkeypatch)
    assert correct['evaluation']['passed'] is True
    assert correct['evaluation']['reconstruction_count'] == 2


def test_reserved_evaluation_is_not_an_authoring_input_or_diagnostic_leak(service, monkeypatch):
    report_type, program = _type(service)
    _pair(service, report_type, 'ambiguous')
    withheld, target, source, _ = _pair(service, report_type, 'discriminating', 'reserved')
    learned, _ = _learn(service, program)
    assert learned['learning']['example_ids'] != []
    assert withheld['id'] not in learned['learning']['example_ids']
    assert all(c['example_id'] != withheld['id'] for c in learned['learning']['coverage'])
    assert len([h for h in learned['learning']['hypotheses'] if h['supported']]) == 3
    evaluated = _evaluate(service, _review(service, learned, 'largest_current_revenue'), monkeypatch)
    reserved = [c for c in evaluated['evaluation']['checks'] if c['name'].startswith('Reserved pair')]
    assert evaluated['evaluation']['holdout_count'] == 1
    assert len(reserved) == 1 and reserved[0]['passed'] is False
    assert 'discrepancies' not in reserved[0] and 'South' not in json.dumps(reserved)
    assert '1600' not in json.dumps(reserved)
    assert service.asset_view(target['id'])['reserved'] and service.asset_view(source['id'])['reserved']
    service.reveal_example(withheld['id'], 'Inspect the failed reserved pair.')
    with pytest.raises(DomainError) as error:
        service.evaluate(evaluated['id'], evaluated['digest'])
    assert error.value.code == 'CORPUS_CHANGED'


@pytest.mark.parametrize('selection,expected', [('largest_current_revenue', 'North'), ('largest_percentage_change', 'East')])
def test_published_selection_changes_real_new_period_output(service, monkeypatch, selection, expected):
    report_type, program = _type(service)
    _pair(service, report_type, 'ambiguous')
    learned, _ = _learn(service, program)
    ready = _evaluate(service, _review(service, learned, selection), monkeypatch)
    assert ready['evaluation']['passed']
    service.publish(ready['id'], ready['digest'])
    _, source, period = _assets(service, 'discriminating')
    completed = _finish(service, service.request_run(report_type['id'], source['id'], period, 'new-period-policy'))
    assert completed['status'] == 'completed', completed.get('error')
    snapshot = service.store.get('snapshot', completed['result']['snapshot_id'])
    assert snapshot['facts']['driver.region']['value'] == expected
    assert snapshot['facts']['revenue.current']['value'] == '1600.00'
    assert snapshot['program']['digest'] == ready['digest']


def test_changed_oracle_fixture_invalidates_publication_without_changing_candidate(service, monkeypatch, tmp_path):
    report_type, program = _type(service)
    ready = _evaluate(service, _review(service, program), monkeypatch)
    fixture_root = tmp_path / 'changed-oracle'; shutil.copytree(ROOT / 'fixtures', fixture_root / 'fixtures')
    expected = fixture_root / 'fixtures' / 'expected.json'
    expected.write_text(expected.read_text() + '\n')
    monkeypatch.setattr('foundry.service.ROOT', fixture_root)
    monkeypatch.setattr(service, '_installed_code_identity', lambda: ready['code_identity'])
    with pytest.raises(DomainError) as error:
        service.publish(ready['id'], ready['digest'])
    assert error.value.code == 'EVALUATION_BASIS_CHANGED'
    assert service.store.get('report_type', report_type['id'])['active_program_id'] is None
    assert service.store.list('release') == []


def test_openai_authoring_is_a_recorded_proposal_and_cannot_resolve_policy(service, monkeypatch):
    model = _model(monkeypatch)
    report_type, program = _type(service)
    _pair(service, report_type, 'ambiguous')
    job = service.request_learning(program['id'], program['digest'], 'Highlight the movement.', 'openai', 'model-authoring')
    completed = _finish(service, job)
    assert completed['status'] == 'completed', completed.get('error')
    learned = service.store.get('program', program['id'])
    assert learned['learning']['engine'] == 'openai_assisted_hypothesis_search'
    assert learned['learning']['model_receipt']['provider'] == 'test_double'
    assert all(d['resolution'] is None for d in learned['decisions'])
    assert len(model.calls) == len(service.store.list('model_call')) == 1


def test_composition_creates_immutable_revision_and_all_exports_do_not_call_model(service, monkeypatch):
    original = _snapshot(service, monkeypatch)
    model = _model(monkeypatch)
    source = service.upload(_docx('The account team reported improved renewal discussions.'), 'management-note.docx')
    request = service.request_composition(original['id'], original['revision'], [source['id']], 'Summarize this reporting period.', 'compose-period')
    completed = _finish(service, request)
    assert completed['status'] == 'completed', completed.get('error')
    revised = service.store.get('snapshot', completed['result']['snapshot_id'])
    validate_stored_snapshot(revised)
    assert revised['parent_id'] == original['id'] and revised['revision'] == original['revision'] + 1
    assert revised['status'] == 'review_required'
    assert revised['facts'] == original['facts'] and revised['datasets'] == original['datasets']
    assert [n for n in revised['nodes'] if n['id'] != 'commentary'] == [n for n in original['nodes'] if n['id'] != 'commentary']
    assert service.store.get('snapshot', original['id']) == original
    assert revised['metadata']['composition']['receipt']['human_review_required'] is True
    assert len(model.calls) == 1
    same = service.request_composition(original['id'], original['revision'], [source['id']], 'Summarize this reporting period.', 'compose-period')
    assert same['id'] == request['id'] and same['status'] == 'completed'
    for format in ['docx', 'xlsx', 'pdf', 'pptx']:
        exported = _finish(service, service.request_export(revised['id'], format, 'compatible', f'composed-{format}'))
        assert exported['status'] == 'completed', exported.get('error')
        artifact = service.store.get('export', exported['result']['export_id'])
        assert len(service.store.read_blob(artifact['digest'])) > 100
    assert len(model.calls) == 1


def test_model_receipt_cache_reuses_exact_job_request_and_cancel_fences_result(service, monkeypatch):
    original = _snapshot(service, monkeypatch)
    model = _model(monkeypatch)
    service.request_composition(original['id'], original['revision'], [], 'Summarize this period.', 'cache-job')
    job = service.store.claim()
    provider = service._captured_provider(job)
    schema = {'type': 'object'}
    first = provider.generate('instructions', {'a': 'b'}, schema, 'report_commentary')
    again = service._captured_provider(job).generate('instructions', {'a': 'b'}, schema, 'report_commentary')
    assert first == again and len(model.calls) == 1
    service.store.cancel(job['id'])
    with pytest.raises(DomainError) as error:
        service.execute_composition(job, lambda name: service.store.stage(job['id'], job['token'], name))
    assert error.value.code == 'LEASE_LOST'
    assert service.store.list('snapshot') == [original]


def test_composition_rejects_historical_target_alias_as_qualitative_input(service, monkeypatch):
    original = _snapshot(service, monkeypatch)
    _model(monkeypatch)
    report_type = service.store.get('report_type', original['report_type_id'])
    _, target, _, _ = _pair(service, report_type, 'ambiguous')
    alias = service.upload(service.store.read_blob(target['digest']), 'renamed-commentary.docx')
    with pytest.raises(DomainError) as error:
        service.request_composition(original['id'], original['revision'], [alias['id']], 'Summarize this period.', 'historical-input')
    assert error.value.code == 'COMPOSITION_SOURCE_INVALID'


def test_newly_registered_historical_target_is_rechecked_before_composition(service, monkeypatch):
    original = _snapshot(service, monkeypatch)
    model = _model(monkeypatch)
    target, source, period = _assets(service, 'ambiguous')
    job = service.request_composition(original['id'], original['revision'], [target['id']], 'Summarize this period.', 'source-role-drift')
    service.add_example(original['report_type_id'], target['id'], [source['id']], period, 'authoring')
    completed = _finish(service, job)
    assert completed['status'] == 'blocked'
    assert completed['error']['code'] == 'COMPOSITION_SOURCE_INVALID'
    assert model.calls == []
    assert service.store.list('snapshot') == [original]


def test_changed_provider_configuration_blocks_queued_composition(service, monkeypatch):
    original = _snapshot(service, monkeypatch)
    model = _model(monkeypatch)
    job = service.request_composition(original['id'], original['revision'], [], 'Summarize this period.', 'provider-drift')
    previous_status = model.status()
    monkeypatch.setattr(model, 'status', lambda: {**previous_status, 'model': 'different-test-model'})
    completed = _finish(service, job)
    assert completed['status'] == 'blocked'
    assert completed['error']['code'] == 'MODEL_CONFIGURATION_CHANGED'
    assert model.calls == []
    assert service.store.list('snapshot') == [original]


def test_cancellation_during_provider_call_does_not_store_receipt_or_revision(service, monkeypatch):
    original = _snapshot(service, monkeypatch)
    model = _model(monkeypatch)
    job = service.request_composition(original['id'], original['revision'], [], 'Summarize this period.', 'cancel-in-flight-model')
    original_generate = model.generate
    def generate_and_cancel(*args):
        result = original_generate(*args)
        service.store.cancel(job['id'])
        return result
    monkeypatch.setattr(model, 'generate', generate_and_cancel)
    completed = _finish(service, job)
    assert completed['status'] == 'cancelled'
    assert len(model.calls) == 1
    assert service.store.list('model_call') == []
    assert service.store.list('snapshot') == [original]


def test_cached_provider_response_is_integrity_checked(service, monkeypatch):
    original = _snapshot(service, monkeypatch)
    model = _model(monkeypatch)
    service.request_composition(original['id'], original['revision'], [], 'Summarize this period.', 'cache-integrity')
    job = service.store.claim()
    provider = service._captured_provider(job)
    provider.generate('instructions', {}, {'type': 'object'}, 'report_commentary')
    stored = service.store.list('model_call')[0]
    path = service.store.blob_path(stored['response_digest'])
    path.chmod(0o600); path.write_bytes(b'changed cached model output')
    with pytest.raises(DomainError) as error:
        provider.generate('instructions', {}, {'type': 'object'}, 'report_commentary')
    assert error.value.code == 'ARTIFACT_INTEGRITY'
    assert len(model.calls) == 1


@pytest.mark.parametrize('kind,cancel', [('generation', False), ('generation', True), ('composition', False), ('composition', True)])
def test_prior_job_bound_assets_cannot_become_reserved_even_after_cancellation(service, monkeypatch, kind, cancel):
    if kind == 'generation':
        report_type, _ = _release(service, monkeypatch)
        target, source, period = _assets(service, 'discriminating')
        pinned = service.request_run(report_type['id'], source['id'], period, 'previously-bound-source')
    else:
        original = _snapshot(service, monkeypatch)
        _model(monkeypatch)
        report_type = service.store.get('report_type', original['report_type_id'])
        target, source, period = _assets(service, 'discriminating')
        pinned = service.request_composition(original['id'], original['revision'], [target['id']],
                                              'Summarize the supplied commentary.', 'previously-bound-target')
    if cancel:
        service.store.cancel(pinned['id'])
        # Exposure history must extend beyond the recent-jobs UI page.
        for index in range(101):
            service.store.enqueue('export', 'unrelated-history', str(index), {'snapshot_id': 'unrelated'})
        assert pinned['id'] not in {job['id'] for job in service.store.jobs()}
    with pytest.raises(DomainError) as error:
        service.add_example(report_type['id'], target['id'], [source['id']], period, 'reserved')
    assert error.value.code == 'EVIDENCE_ALREADY_EXPOSED'
    assert service.store.list('example') == []


def test_historical_target_registered_during_model_call_blocks_revision_commit(service, monkeypatch):
    original = _snapshot(service, monkeypatch)
    model = _model(monkeypatch)
    target, source, period = _assets(service, 'ambiguous')
    job = service.request_composition(original['id'], original['revision'], [target['id']],
                                       'Summarize the supplied commentary.', 'mid-call-source-role-drift')
    original_generate = model.generate
    def generate_and_register(*args):
        result = original_generate(*args)
        service.add_example(original['report_type_id'], target['id'], [source['id']], period, 'authoring')
        return result
    monkeypatch.setattr(model, 'generate', generate_and_register)
    completed = _finish(service, job)
    assert completed['status'] == 'blocked'
    assert completed['error']['code'] == 'COMPOSITION_SOURCE_INVALID'
    assert len(model.calls) == 1
    assert service.store.list('snapshot') == [original]


def _change_installed_identity(monkeypatch):
    import foundry.service as service_module
    changed = copy.deepcopy(service_module.code_identity())
    changed['test_installed_revision'] = 'changed-after-process-start'
    monkeypatch.setattr(service_module, 'code_identity', lambda: copy.deepcopy(changed))


@pytest.mark.parametrize('phase', ['request', 'execution'])
def test_composition_restart_required_before_model_invocation(service, monkeypatch, phase):
    original = _snapshot(service, monkeypatch)
    model = _model(monkeypatch)
    if phase == 'execution':
        job = service.request_composition(original['id'], original['revision'], [],
                                          'Summarize this period.', 'restart-before-execution')
    _change_installed_identity(monkeypatch)
    if phase == 'request':
        with pytest.raises(DomainError) as error:
            service.request_composition(original['id'], original['revision'], [],
                                         'Summarize this period.', 'restart-before-request')
        assert error.value.code == 'RUNTIME_RESTART_REQUIRED'
        assert not any(job['kind'] == 'composition' for job in service.store.jobs())
    else:
        completed = _finish(service, job)
        assert completed['status'] == 'blocked'
        assert completed['error']['code'] == 'RUNTIME_RESTART_REQUIRED'
    assert model.calls == []
    assert service.store.list('model_call') == []
    assert service.store.list('snapshot') == [original]


def test_installed_code_changed_during_model_call_prevents_draft_commit(service, monkeypatch):
    original = _snapshot(service, monkeypatch)
    model = _model(monkeypatch)
    job = service.request_composition(original['id'], original['revision'], [],
                                      'Summarize this period.', 'restart-during-model-call')
    original_generate = model.generate
    def generate_and_change_runtime(*args):
        result = original_generate(*args)
        _change_installed_identity(monkeypatch)
        return result
    monkeypatch.setattr(model, 'generate', generate_and_change_runtime)
    completed = _finish(service, job)
    assert completed['status'] == 'blocked'
    assert completed['error']['code'] == 'RUNTIME_RESTART_REQUIRED'
    assert len(model.calls) == 1
    assert service.store.list('snapshot') == [original]
    assert not any(event['action'] == 'commentary_composed' for event in service.store.audit_for(original['id']))


def test_legacy_reserved_pairs_mask_original_assets_and_same_content_aliases(tmp_path):
    with TestClient(create_app(tmp_path, start_worker=False)) as client:
        service = client.app.state.service
        report_type, _ = _type(service)
        target, source, period = _assets(service, 'discriminating')
        target_alias = service.upload(service.store.read_blob(target['digest']), 'legacy-target-alias.docx')
        source_alias = service.upload(service.store.read_blob(source['digest']), 'legacy-source-alias.csv')
        # This is the exact pre-authoring catalog shape: no frozen asset_digests,
        # observation inspection or example digest was stored by the old host.
        legacy = {'id': 'legacy_reserved_pair', 'report_type_id': report_type['id'],
                  'report_asset_id': target['id'], 'source_asset_ids': [source['id']],
                  'period': period, 'corpus_role': 'reserved', 'state': 'inspected',
                  'learning_status': 'not_implemented', 'created_at': '2026-09-08T12:00:00Z'}
        service.store.insert('example', legacy)
        for asset in [target, source, target_alias, source_alias]:
            response = client.get(f'/api/assets/{asset["id"]}')
            assert response.status_code == 200
            assert response.json()['reserved'] is True
            assert response.json()['profile']['regions'] == []
            assert response.json()['profile']['eligible_roles'] == []
            assert 'sample_rows' not in response.json()['profile']
            denied = client.get(f'/api/assets/{asset["id"]}/download')
            assert denied.status_code == 403
            assert denied.json()['error']['code'] == 'RESERVED_EVIDENCE'
        assert all(asset.get('reserved') for asset in client.get('/api/bootstrap').json()['assets'])
        assert client.get('/api/examples/legacy_reserved_pair').json()['inspection'] is None
        assert service.store.get('example', legacy['id']) == legacy


def test_reserved_unknown_component_fails_reconstruction_without_disclosing_its_text(service, monkeypatch):
    from docx import Document
    from foundry.learning import compare_snapshot
    from foundry.runtime import prepare

    report_type, program = _type(service)
    _pair(service, report_type, 'ambiguous')
    folder = FIXTURES / 'discriminating'
    document = Document(BytesIO((folder / 'report.docx').read_bytes()))
    sentinel = 'Confidential withheld explanation that the supported reporting profile cannot reproduce.'
    document.add_paragraph(sentinel)
    stream = BytesIO(); document.save(stream)
    target = service.upload(stream.getvalue(), 'reserved-report-with-extra-content.docx')
    source = service.upload((folder / 'transactions.csv').read_bytes(), 'reserved-transactions.csv')
    period = json.loads((folder / 'period.json').read_text())
    example = service.add_example(report_type['id'], target['id'], [source['id']], period, 'reserved')
    learned, _ = _learn(service, program)
    reviewed = _review(service, learned)

    # The evaluator can reconstruct every extracted observation. The required
    # unknown region is the independent reason the reserved case must still fail.
    entry = next(e for e in reviewed['learning']['corpus'] if e['id'] == example['id'])
    case = service._case(entry, evaluator=True)
    prepared = prepare(case['rows'], case['period'], {k: source[k] for k in ('id', 'digest', 'filename')},
                       reporting_policy={'selection': 'largest_absolute_change'})
    assert compare_snapshot(prepared, case['inspection'])['passed'] is True
    assert any(region['status'] != 'mapped' and sentinel in region.get('text', '') for region in case['inspection']['regions'])

    evaluated = _evaluate(service, reviewed, monkeypatch)
    withheld = [check for check in evaluated['evaluation']['checks'] if check['name'].startswith('Reserved pair')]
    assert len(withheld) == 1 and withheld[0]['passed'] is False
    assert evaluated['evaluation']['passed'] is False
    assert all(check['passed'] for check in evaluated['evaluation']['checks'] if not check['name'].startswith('Reserved pair'))
    assert sentinel not in json.dumps(evaluated)
    assert sentinel not in json.dumps(service.bootstrap())
    assert 'discrepancies' not in withheld[0]
