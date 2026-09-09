"""Independent exported-file checks, deliberately separate from runtime maths."""
import copy
import hashlib
import json
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

from foundry.exporters import ExportError, export_snapshot
from foundry.runtime import fixture_snapshot, render_runs
from foundry.exporters.preview import render_office_preview


@pytest.fixture
def snapshot():
    return fixture_snapshot().model_dump(mode="json")


@pytest.mark.parametrize("format", ["docx", "xlsx", "pdf", "pptx"])
def test_exports_real_bytes_same_frozen_snapshot(format, snapshot, tmp_path):
    before = copy.deepcopy(snapshot)
    result = export_snapshot(snapshot, format, tmp_path)
    file = Path(result["path"])
    assert file.is_file() and file.stat().st_size > 1000
    assert snapshot == before
    manifest = result["manifest"]
    assert hashlib.sha256(file.read_bytes()).hexdigest() == result["sha256"]
    assert manifest["artifact"]["sha256"] == result["sha256"]
    assert json.loads(Path(result["manifest_path"]).read_text()) == manifest
    assert manifest["snapshot_digest"] == hashlib.sha256(json.dumps(snapshot, sort_keys=True, separators=(",", ":"), ensure_ascii=False).encode()).hexdigest()
    assert {c["node_id"] for c in manifest["components"]} == {"summary", "commentary", "regional_table", "regional_chart", "regional_pivot"}
    assert manifest["renderer"]["version"] and manifest["renderer"]["adapter_sha256"]


def test_docx_values_styles_and_static_fallbacks(snapshot, tmp_path):
    result = export_snapshot(snapshot, "docx", tmp_path)
    doc = Document(result["path"])
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "€1,200.00" in text and "20.0%" in text and "North" in text
    assert len(doc.tables) == 2
    table_rows = [[cell.text for cell in row.cells] for row in doc.tables[0].rows]
    north = next(row for row in table_rows if row[0] == "North")
    assert north[1:4] == ["500.00", "400.00", "100.00"]
    with ZipFile(result["path"]) as z:
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        xml = ET.fromstring(z.read("word/document.xml"))
        assert len(xml.findall(".//w:tblHeader", ns)) == 2
        assert any(e.get(f"{{{ns['w']}}}left") == "360" for e in xml.findall(".//w:ind", ns))
        assert len([name for name in z.namelist() if name.startswith("word/media/")]) == 1
    by_id = {c["node_id"]: c for c in result["manifest"]["components"]}
    assert by_id["regional_pivot"]["representation"] == "equivalent_static_table"
    assert by_id["regional_chart"]["representation"] == "raster_fallback"


def test_native_xlsx_pivot_cache_reconciles_and_contains_only_approved_data(snapshot, tmp_path):
    result = export_snapshot(snapshot, "xlsx", tmp_path)
    wb = load_workbook(result["path"], data_only=True)
    assert wb.sheetnames == ["Overview", "Analysis", "Approved data"]
    overview_values = [cell.value for row in wb["Overview"] for cell in row]
    assert any(isinstance(v, str) and "€1,200.00" in v for v in overview_values)
    assert wb["Analysis"]["A6"].value == "North"
    assert wb["Analysis"]["B6"].value == 400
    assert wb["Analysis"]["C6"].value == 500
    assert len(wb["Analysis"]._pivots) == 1
    pivot = wb["Analysis"]._pivots[0]
    assert pivot.rowFields[0].x == 0 and pivot.colFields[0].x == 1
    assert pivot.dataFields[0].subtotal == "sum" and pivot.dataFields[0].fld == 2
    assert [f.name for f in pivot.cache.cacheFields] == ["region", "period", "revenue"]
    assert pivot.cache.recordCount == 8
    assert pivot.cache.cacheSource.worksheetSource.sheet == "Approved data"
    with ZipFile(result["path"]) as z:
        ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        records = ET.fromstring(z.read("xl/pivotCache/pivotCacheRecords1.xml"))
        values = records.findall("s:r", ns)
        assert len(values) == 8
        assert sum(Decimal(r.find("s:n", ns).get("v")) for r in values) == Decimal("2200")
        assert all(len(list(r)) == 3 for r in values)
        assert b"transaction_id" not in z.read("xl/pivotCache/pivotCacheDefinition1.xml")
        assert not any("externalLink" in name for name in z.namelist())
    assert any(w["code"] == "native_certification_pending" for w in result["manifest"]["warnings"])
    assert result["manifest"]["state"] == "verified_with_limitations"
    wb.close()


def test_strict_pivot_gate_blocks_before_writing(snapshot, tmp_path):
    with pytest.raises(ExportError) as error:
        export_snapshot(snapshot, "xlsx", tmp_path, "strict")
    assert error.value.code == "native_certification_pending"
    assert not list(tmp_path.iterdir())


def test_static_policy_is_an_explicit_downgrade(snapshot, tmp_path):
    result = export_snapshot(snapshot, "xlsx", tmp_path, "static")
    wb = load_workbook(result["path"])
    assert not wb["Analysis"]._pivots
    assert wb["Analysis"]["C6"].value == 500
    component = next(c for c in result["manifest"]["components"] if c["node_id"] == "regional_pivot")
    assert component["representation"] == "equivalent_static_table"
    assert any(w["code"] == "static_pivot_approved" for w in result["manifest"]["warnings"])
    wb.close()


def test_pptx_two_slides_native_chart_and_editable_tables(snapshot, tmp_path):
    result = export_snapshot(snapshot, "pptx", tmp_path)
    prs = Presentation(result["path"])
    assert len(prs.slides) == 2
    charts = [s.chart for slide in prs.slides for s in slide.shapes if s.has_chart]
    tables = [s.table for slide in prs.slides for s in slide.shapes if s.has_table]
    assert len(charts) == 1 and len(tables) == 2
    assert sorted(float(v) for v in charts[0].series[0].values) == [150, 200, 350, 500]
    assert sum(charts[0].series[0].values) == 1200
    with ZipFile(result["path"]) as z:
        embedded = [n for n in z.namelist() if n.startswith("ppt/embeddings/")]
        assert len(embedded) == 1
        chart_wb = load_workbook(BytesIO(z.read(embedded[0])))
        all_values = [c.value for ws in chart_wb for row in ws for c in row]
        assert "transaction_id" not in all_values
        assert "North" in all_values
        chart_wb.close()


def test_pdf_native_route_and_required_text(snapshot, tmp_path):
    result = export_snapshot(snapshot, "pdf", tmp_path)
    reader = PdfReader(result["path"])
    assert set(reader.named_destinations) == {"summary", "commentary", "regional_table", "regional_chart", "regional_pivot"}
    text = " ".join(p.extract_text() for p in reader.pages)
    assert "€1,200.00" in text and "20.0%" in text
    assert "North" in text and "500.00" in text
    assert result["manifest"]["details"]["pdf_origin"] == "native_flow_reportlab"
    assert result["manifest"]["details"]["office_pagination_equivalence"] is False


def test_export_does_not_recompose_accepted_text(snapshot, tmp_path):
    commentary = next(n for n in snapshot["nodes"] if n["id"] == "commentary")
    commentary["runs"] = [{"type": "text", "text": "The operations team reviewed this period."}]
    result = export_snapshot(snapshot, "docx", tmp_path)
    actual = "\n".join(p.text for p in Document(result["path"]).paragraphs)
    assert "The operations team reviewed this period." in actual


def test_formula_like_text_remains_text_in_excel(snapshot, tmp_path):
    commentary = next(n for n in snapshot["nodes"] if n["id"] == "commentary")
    commentary["runs"] = [{"type": "text", "text": '=HYPERLINK("https://example.invalid","untrusted text")'}]
    result = export_snapshot(snapshot, "xlsx", tmp_path)
    wb = load_workbook(result["path"], data_only=False)
    matches = [c for ws in wb for row in ws for c in row if isinstance(c.value, str) and c.value.startswith("=HYPERLINK")]
    assert len(matches) == 1 and matches[0].data_type == "s"
    assert all(c.data_type != "f" for ws in wb for row in ws for c in row)
    wb.close()


def test_pivot_cache_disagreement_cannot_export(snapshot, tmp_path):
    snapshot["datasets"]["pivot_source"]["rows"][0]["revenue"] = "999.00"
    with pytest.raises(ExportError) as error:
        export_snapshot(snapshot, "xlsx", tmp_path)
    assert error.value.code == "pivot_reconciliation_failed"


def test_incomplete_view_cannot_disappear_during_export(snapshot, tmp_path):
    snapshot["views"][0]["node_ids"].remove("regional_pivot")
    with pytest.raises(ExportError) as error:
        export_snapshot(snapshot, "docx", tmp_path)
    assert error.value.code == "invalid_snapshot"


def test_snapshot_blocked_state_prevents_export(snapshot, tmp_path):
    snapshot["status"] = "blocked"
    with pytest.raises(ExportError) as error:
        export_snapshot(snapshot, "pdf", tmp_path)
    assert error.value.code == "snapshot_blocked"


def test_unknown_format_and_policy_are_actionable(snapshot, tmp_path):
    with pytest.raises(ExportError, match="Choose DOCX"):
        export_snapshot(snapshot, "html", tmp_path)
    with pytest.raises(ExportError, match="Unknown fidelity"):
        export_snapshot(snapshot, "pdf", tmp_path, "best")


def test_existing_bytes_are_not_overwritten(snapshot, tmp_path):
    result = export_snapshot(snapshot, "pdf", tmp_path)
    before = Path(result["path"]).read_bytes()
    with pytest.raises(ExportError) as error:
        export_snapshot(snapshot, "pdf", tmp_path)
    assert error.value.code == "artifact_exists"
    assert Path(result["path"]).read_bytes() == before


def test_precision_preview_missing_backend_is_explicit(snapshot, tmp_path, monkeypatch):
    artifact = export_snapshot(snapshot, "docx", tmp_path / "export")
    monkeypatch.setenv("FOUNDRY_SOFFICE", str(tmp_path / "does-not-exist"))
    with pytest.raises(ExportError) as error:
        render_office_preview(Path(artifact["path"]), tmp_path / "preview")
    assert error.value.code == "preview_backend_unavailable"
    assert not (tmp_path / "preview").exists()


def test_precision_preview_timeout_cannot_leave_successful_bytes(snapshot, tmp_path, monkeypatch):
    import subprocess
    from foundry.exporters import preview
    artifact = export_snapshot(snapshot, "docx", tmp_path / "export")
    executable = tmp_path / "soffice"
    executable.write_text("#!/bin/sh\nexit 0\n")
    executable.chmod(0o755)
    monkeypatch.setenv("FOUNDRY_SOFFICE", str(executable))
    calls = []
    def fake_run(command, **kwargs):
        calls.append(command)
        if "--version" in command:
            return subprocess.CompletedProcess(command, 0, "TestOffice 1.0\n", "")
        raise subprocess.TimeoutExpired(command, kwargs["timeout"])
    monkeypatch.setattr(preview.subprocess, "run", fake_run)
    with pytest.raises(ExportError) as error:
        render_office_preview(Path(artifact["path"]), tmp_path / "preview", timeout_seconds=.1)
    assert error.value.code == "preview_timeout"
    assert len(calls) == 2
    assert "--headless" in calls[1] and any(arg.startswith("-env:UserInstallation=file:") for arg in calls[1])
    assert not list((tmp_path / "preview").glob("*.pdf"))


def test_precision_preview_rejects_external_office_links(tmp_path, monkeypatch):
    executable = tmp_path / "soffice"
    executable.write_text("#!/bin/sh\nexit 0\n")
    executable.chmod(0o755)
    monkeypatch.setenv("FOUNDRY_SOFFICE", str(executable))
    file = tmp_path / "linked.docx"
    with ZipFile(file, "w") as z:
        z.writestr("word/_rels/document.xml.rels", '<Relationships><Relationship Target="https://example.invalid/" TargetMode="External"/></Relationships>')
    with pytest.raises(ExportError) as error:
        render_office_preview(file, tmp_path / "preview")
    assert error.value.code == "preview_external_relationship"


def test_unsupported_pdf_glyphs_are_not_silently_replaced(snapshot, tmp_path):
    node = next(n for n in snapshot["nodes"] if n["id"] == "commentary")
    node["runs"] = [{"type": "text", "text": "東京"}]
    with pytest.raises(ExportError) as error:
        export_snapshot(snapshot, "pdf", tmp_path)
    assert error.value.code == "unsupported_font_glyph"
    assert not list(tmp_path.glob("*.pdf"))


@pytest.mark.parametrize("format", ["docx", "xlsx", "pdf", "pptx"])
def test_review_required_exports_are_visibly_marked_draft(format, snapshot, tmp_path):
    snapshot["status"] = "review_required"
    result = export_snapshot(snapshot, format, tmp_path)
    assert result["filename"].startswith("draft-")
    if format == "docx":
        text = "\n".join(p.text for p in Document(result["path"]).paragraphs)
    elif format == "xlsx":
        workbook = load_workbook(result["path"])
        text = "\n".join(str(c.value) for ws in workbook for row in ws for c in row)
        workbook.close()
    elif format == "pdf":
        text = "\n".join(page.extract_text() for page in PdfReader(result["path"]).pages)
    else:
        text = "\n".join(shape.text for slide in Presentation(result["path"]).slides for shape in slide.shapes if shape.has_text_frame)
    assert "DRAFT / Awaiting review" in text
    assert any(w["code"] == "draft_snapshot" for w in result["manifest"]["warnings"])
