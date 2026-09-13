"""Prove the frozen future-period oracle and its export readback catch damage.

Expected answers come from independently authored JSON/DOCX fixtures. Runtime
preparation supplies only the candidate output being checked. No model is used.
"""
from __future__ import annotations

import copy
import hashlib
import importlib.util
import json
from pathlib import Path
import shutil

import pytest

from foundry.exporters import export_snapshot
from foundry.ingestion import inspect_asset
from foundry.observations import compare_snapshot, inspect_target
from foundry.runtime import prepare


ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "fixtures" / "evolution"
MANIFEST = json.loads((FIXTURES / "manifest.json").read_text())
FUTURES = [(corpus, case) for corpus in MANIFEST["corpora"] for case in corpus["future"]]
SPEC = importlib.util.spec_from_file_location("evolution_checks_test_subject", ROOT / "scripts" / "evolution_checks.py")
checks = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(checks)


def _case(corpus, case):
    data = (FIXTURES / case["source"]).read_bytes()
    source = {"id": "evolution_source", "digest": hashlib.sha256(data).hexdigest(),
              "filename": Path(case["source"]).name}
    rows = checks.read_rows(data)
    expected = json.loads((FIXTURES / case["expected"]).read_text())
    period = json.loads((FIXTURES / case["period"]).read_text())
    snapshot = prepare(rows, period, source, reporting_policy={"selection": corpus["expected_selection"]}).model_dump(mode="json")
    return snapshot, expected, corpus["expected_selection"], source, rows


def _find(corpus_id, case_id):
    return next((corpus, case) for corpus, case in FUTURES if corpus["id"] == corpus_id and case["id"] == case_id)


@pytest.mark.parametrize("corpus,case", FUTURES,
                         ids=[f'{corpus["id"]}/{case["id"]}' for corpus, case in FUTURES])
def test_all_twelve_unseen_periods_match_independent_exact_answers(corpus, case):
    snapshot, expected, selection, source, rows = _case(corpus, case)
    assert expected["oracle"]["candidate_output_used"] is False
    assert expected["oracle"]["runtime_imported"] is False
    assert checks.snapshot_errors(snapshot, expected, selection, source, rows) == []


@pytest.mark.parametrize("corpus,case", FUTURES,
                         ids=[f'{corpus["id"]}/{case["id"]}' for corpus, case in FUTURES])
def test_independently_written_future_docx_targets_reconstruct(corpus, case, monkeypatch):
    snapshot, expected, _, _, _ = _case(corpus, case)
    data = (FIXTURES / case["report"]).read_bytes()
    asset = {"id": "independent_future_target", "digest": hashlib.sha256(data).hexdigest(),
             **inspect_asset(data, "independent-future-report.docx")}
    def forbidden(*args, **kwargs):
        raise AssertionError("Historical target inspection cannot compute expected values with runtime preparation")
    monkeypatch.setattr("foundry.runtime.prepare", forbidden)
    inspection = inspect_target(asset)
    assert inspection["observations"] and not inspection["issues"]
    assert all(region["status"] == "mapped" for region in inspection["regions"])
    comparison = compare_snapshot(snapshot, inspection)
    assert comparison["passed"], comparison
    assert comparison["observation_count"] >= 5 + 4 * len(expected["regions"])


def test_every_matrix_input_and_expected_artifact_matches_its_frozen_digest():
    frozen = json.loads((FIXTURES / "frozen-sha256.json").read_text())
    assert len(FUTURES) == 12 and len(frozen["files"]) == 72
    for relative, digest in frozen["files"].items():
        assert hashlib.sha256((FIXTURES / relative).read_bytes()).hexdigest() == digest, relative


@pytest.mark.parametrize("mutation", [
    "dataset_number_only", "wrong_driver", "wrong_change_sign", "missing_region",
    "pivot_number_only", "source_digest", "provenance_digest", "physical_csv_line",
    "undefined_as_known", "unknown_status", "missing_undefined_status", "unknown_undefined_value",
    "missing_undefined_fact", "total_growth_unknown_status", "total_growth_missing_status",
    "total_growth_missing_value", "wrong_driver_dependencies",
])
def test_checker_rejects_plausible_snapshot_corruption(mutation):
    snapshot, expected, selection, source, rows = _case(*_find("current_revenue", "future_q3_2026"))
    original = copy.deepcopy(snapshot)
    undefined_id = next(fid for fid, fact in snapshot["facts"].items() if fid.startswith("region.")
                        and fid.endswith(".growth") and fact["value"] is None)
    if mutation == "dataset_number_only":
        snapshot["datasets"]["regional_totals"]["rows"][0]["current"] = "999.00"
        assert snapshot["facts"] == original["facts"]
    elif mutation == "wrong_driver":
        snapshot["facts"]["driver.region"]["value"] = "Atrium"
    elif mutation == "wrong_change_sign":
        snapshot["facts"]["revenue.change"]["value"] = "-215.00"
    elif mutation == "missing_region":
        snapshot["datasets"]["regional_totals"]["rows"].pop()
    elif mutation == "pivot_number_only":
        snapshot["datasets"]["pivot_source"]["rows"][0]["revenue"] = "999.00"
        assert snapshot["facts"] == original["facts"]
    elif mutation == "source_digest":
        snapshot["source_assets"][0]["digest"] = "0" * 64
    elif mutation == "provenance_digest":
        snapshot["facts"]["region.atrium.current"]["sources"][0]["artifact_sha256"] = "0" * 64
    elif mutation == "physical_csv_line":
        snapshot["facts"]["region.atrium.current"]["sources"][0]["locator"] = "csv:lines=999;columns=date,region,status,amount,currency"
    elif mutation == "undefined_as_known":
        snapshot["facts"][undefined_id]["status"] = "known"
    elif mutation == "unknown_status":
        snapshot["facts"][undefined_id]["status"] = "unknown"
    elif mutation == "missing_undefined_status":
        snapshot["facts"][undefined_id].pop("status")
    elif mutation == "unknown_undefined_value":
        snapshot["facts"][undefined_id]["value"] = "Unknown"
    elif mutation == "missing_undefined_fact":
        snapshot["facts"].pop(undefined_id)
    elif mutation == "total_growth_unknown_status":
        snapshot["facts"]["revenue.growth"]["status"] = "unknown"
    elif mutation == "total_growth_missing_status":
        snapshot["facts"]["revenue.growth"].pop("status")
    elif mutation == "total_growth_missing_value":
        snapshot["facts"]["revenue.growth"]["value"] = None
    elif mutation == "wrong_driver_dependencies":
        snapshot["facts"]["driver.region"]["inputs"] = ["region.atrium.change"]
    errors = checks.snapshot_errors(snapshot, expected, selection, source, rows)
    assert errors, mutation
    assert checks.snapshot_errors(original, expected, selection, source, rows) == []


@pytest.fixture(scope="module")
def exported_future(tmp_path_factory):
    snapshot, expected, _, _, _ = _case(*_find("current_revenue", "future_q3_2026"))
    directory = tmp_path_factory.mktemp("evolution-export-readback")
    artifacts = {fmt: Path(export_snapshot(snapshot, fmt, directory / fmt)["path"])
                 for fmt in ("docx", "xlsx", "pdf", "pptx")}
    return snapshot, expected, artifacts


@pytest.mark.parametrize("fmt", ["docx", "xlsx", "pdf", "pptx"])
def test_real_export_bytes_match_independent_future_values(exported_future, fmt):
    snapshot, expected, artifacts = exported_future
    assert checks.export_errors(artifacts[fmt], fmt, snapshot, expected) == []


def _damage_table(path, fmt):
    if fmt == "docx":
        from docx import Document
        document = Document(path)
        row = next(row for table in document.tables if len(table.columns) == 5
                   for row in table.rows if row.cells[0].text == "Briar")
        row.cells[1].text = "999.00"
        document.save(path)
    elif fmt == "xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path)
        sheet = workbook["Overview"]
        row = next(row for row in sheet if row[0].value == "Briar")
        row[1].value = 999
        workbook.save(path)
        workbook.close()
    elif fmt == "pptx":
        from pptx import Presentation
        presentation = Presentation(path)
        row = next(row for slide in presentation.slides for shape in slide.shapes
                   if shape.has_table and len(shape.table.columns) == 5
                   for row in shape.table.rows if row.cells[0].text == "Briar")
        row.cells[1].text = "999.00"
        presentation.save(path)
    else:
        from pypdf import PdfWriter
        from pypdf.generic import ContentStream, TextStringObject
        writer = PdfWriter(clone_from=path)
        replaced = 0
        for page in writer.pages:
            content = ContentStream(page.get_contents(), writer)
            for operands, operator in content.operations:
                if operator == b"Tj" and operands[0] == "90.00":
                    operands[0] = TextStringObject("999.00")
                    replaced += 1
                    break
            if replaced:
                page.replace_contents(content)
                break
        assert replaced == 1
        writer.write(path)
        writer.close()


@pytest.mark.parametrize("fmt", ["docx", "xlsx", "pdf", "pptx"])
def test_export_readback_rejects_changed_table_numbers_with_original_prose(exported_future, tmp_path, fmt):
    snapshot, expected, artifacts = exported_future
    target = tmp_path / f"corrupted.{fmt}"
    shutil.copyfile(artifacts[fmt], target)
    original_text = checks.exported_text(target, fmt)
    _damage_table(target, fmt)
    # The native snapshot remains authoritative, and the altered file retains
    # the same summary. Only a real table-cell readback catches this corruption.
    errors = checks.export_errors(target, fmt, snapshot, expected)
    assert "export Briar current" in errors, errors
    assert not any(error.startswith("exported ") for error in errors), (original_text, errors)
    assert checks.export_errors(artifacts[fmt], fmt, snapshot, expected) == []
