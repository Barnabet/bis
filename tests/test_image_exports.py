"""Inspect image pixels, placement and accessibility in the actual export files."""
import copy
import hashlib
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from docx import Document
from openpyxl import load_workbook
from PIL import Image, PngImagePlugin
from pptx import Presentation
from pypdf import PdfReader
from pypdf.generic import ContentStream

from foundry.exporters import ExportError, export_snapshot
from foundry.runtime import fixture_snapshot

FORMATS = ["docx", "xlsx", "pdf", "pptx"]
ALT = "A green rectangle beside a smaller blue rectangle."
ORIGINAL_DIGEST = hashlib.sha256(b"original uploaded bytes, distinct from normalized PNG").hexdigest()


def png_bytes(size=(320, 160), **save_options):
    picture = Image.new("RGB", size, "#29785C")
    picture.paste("#17416D", (size[0]//2, size[1]//4, size[0]*3//4, size[1]*3//4))
    stream = BytesIO()
    picture.save(stream, format="PNG", **save_options)
    return stream.getvalue()


def image_snapshot(data, *, size=(320, 160), count=1, decorative=False):
    snapshot = fixture_snapshot().model_dump(mode="json")
    snapshot["source_assets"].append({"id": "image_source", "filename": "original-upload.png", "digest": ORIGINAL_DIGEST})
    for index in range(count):
        node = {"id": "report_image" if index == 0 else f"report_image_{index}", "kind": "image",
                "title": "Report image" if index == 0 else f"Report image {index+1}", "asset_id": "image_source",
                "width_px": size[0], "height_px": size[1], "render_digest": hashlib.sha256(data).hexdigest(),
                "media_type": "image/png", "alt_text": "" if decorative else ALT, "decorative": decorative}
        snapshot["nodes"].append(node)
        snapshot["nodes"][0]["children"].append(node["id"])
        for view in snapshot["views"]:
            view["node_ids"].append(node["id"])
            view["coverage"]["required_node_ids"].append(node["id"])
            if view["family"] == "grid":
                view["recipe"]["sheets"].append({"name": "Images", "node_ids": [node["id"]]})
            elif view["family"] == "canvas":
                view["recipe"]["slides"].append({"title": node["title"], "node_ids": [node["id"]]})
    return snapshot


def component(result):
    return next(item for item in result["manifest"]["components"] if item["node_id"] == "report_image")


def office_picture_properties(path, format):
    with ZipFile(path) as archive:
        if format == "docx":
            roots = [ET.fromstring(archive.read("word/document.xml"))]
        elif format == "pptx":
            roots = [ET.fromstring(archive.read(name)) for name in archive.namelist()
                     if name.startswith("ppt/slides/slide") and name.endswith(".xml")]
        else:
            roots = [ET.fromstring(archive.read(name)) for name in archive.namelist()
                     if name.startswith("xl/drawings/drawing") and name.endswith(".xml")]
    return [element for root in roots for element in root.iter()
            if element.tag.rsplit("}", 1)[-1] in {"cNvPr", "docPr"} and element.get("title") == "Report image"]


@pytest.mark.parametrize("format", FORMATS)
def test_real_image_bytes_and_manifest_are_bound_to_frozen_snapshot(format, tmp_path):
    data = png_bytes()
    snapshot = image_snapshot(data)
    before = copy.deepcopy(snapshot)
    calls = []
    def resolver(digest):
        calls.append(digest)
        return data
    result = export_snapshot(snapshot, format, tmp_path, asset_resolver=resolver)
    assert snapshot == before
    assert calls == [hashlib.sha256(data).hexdigest()]
    entry = component(result)
    assert entry["asset_id"] == "image_source"
    assert entry["original_digest"] == ORIGINAL_DIGEST != entry["render_digest"]
    assert entry["render_digest"] == calls[0]
    assert (entry["width_px"], entry["height_px"]) == (320, 160)
    assert entry["alt_text"] == ALT and not entry["decorative"]
    assert entry["pixel_content_editable"] is False and entry["editable"] is False
    assert entry["aspect_ratio_preserved"] and entry["crop"] == "none"
    assert entry["rendered_size"]["width"] / entry["rendered_size"]["height"] == pytest.approx(2)
    if format == "pdf":
        assert entry["representation"] == "embedded_image"
        assert entry["tagged_pdf"] is False
        assert entry["alt_text_support"] == "manifest_and_visible_caption"
        reader = PdfReader(result["path"])
        embedded = [image.image for page in reader.pages for image in page.images if image.image.size == (320, 160)]
        assert len(embedded) == 1
        assert embedded[0].convert("RGB").tobytes() == Image.open(BytesIO(data)).tobytes()
        assert ALT in "\n".join(page.extract_text() for page in reader.pages)
        assert "report_image" in reader.named_destinations
        assert any(w["code"] == "pdf_image_tagging_unavailable" for w in result["manifest"]["warnings"])
    else:
        assert entry["representation"] == "native_image" and entry["placement_editable"]
        with ZipFile(result["path"]) as archive:
            assert data in [archive.read(name) for name in archive.namelist() if "/media/" in name]
            for name in archive.namelist():
                if name.endswith(".rels"):
                    assert all(rel.get("TargetMode") != "External" for rel in ET.fromstring(archive.read(name)))
        properties = office_picture_properties(result["path"], format)
        assert properties and all(item.get("descr") == ALT for item in properties)
        assert all(item.get("title") == "Report image" for item in properties)


@pytest.mark.parametrize("size", [(1600, 400), (400, 1600), (32, 16)])
@pytest.mark.parametrize("format", ["docx", "xlsx", "pptx"])
def test_actual_office_geometry_preserves_ratio_and_stays_in_layout(format, size, tmp_path):
    data = png_bytes(size)
    result = export_snapshot(image_snapshot(data, size=size), format, tmp_path, asset_resolver=lambda _: data)
    ratio = size[0] / size[1]
    if format == "docx":
        document = Document(result["path"])
        shape = document.inline_shapes[-1]
        section = document.sections[0]
        assert shape.width <= section.page_width - section.left_margin - section.right_margin
        assert shape.height / 914400 <= 6.2
        assert shape.width / shape.height == pytest.approx(ratio, rel=1e-5)
    elif format == "pptx":
        presentation = Presentation(result["path"])
        assert len(presentation.slides) == 3
        pictures = [shape for shape in presentation.slides[-1].shapes if shape.shape_type == 13]
        assert len(pictures) == 1
        shape = pictures[0]
        assert shape.left >= 0 and shape.top >= 0
        assert shape.left + shape.width <= presentation.slide_width
        assert shape.top + shape.height <= presentation.slide_height
        assert shape.width / shape.height == pytest.approx(ratio, rel=1e-5)
        assert shape.crop_left == shape.crop_right == shape.crop_top == shape.crop_bottom == 0
    else:
        workbook = load_workbook(result["path"])
        assert workbook.sheetnames == ["Overview", "Analysis", "Approved data", "Images"]
        assert len(workbook["Images"]._images) == 1
        with ZipFile(result["path"]) as archive:
            namespace = {"xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
                         "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
            pictures = [picture for name in archive.namelist() if name.startswith("xl/drawings/drawing") and name.endswith(".xml")
                        for picture in ET.fromstring(archive.read(name)).findall(".//xdr:pic", namespace)]
            assert len(pictures) == 1
            extent = pictures[0].find("xdr:spPr/a:xfrm/a:ext", namespace)
            width, height = int(extent.get("cx")), int(extent.get("cy"))
            assert width / height == pytest.approx(ratio, rel=1e-5)
            assert width / 9525 <= 800 and height / 9525 <= 480
        workbook.close()


@pytest.mark.parametrize("format", FORMATS)
def test_missing_image_resolver_blocks_before_output_directory_exists(format, tmp_path):
    output = tmp_path / "must-not-exist"
    with pytest.raises(ExportError) as error:
        export_snapshot(image_snapshot(png_bytes()), format, output)
    assert error.value.code == "image_resolver_required"
    assert not output.exists()


@pytest.mark.parametrize("failure,expected", [
    ("legacy", "image_render_unavailable"),
    ("digest", "image_integrity"),
    ("dimensions", "image_dimensions_mismatch"),
    ("undecodable", "image_decode_failed"),
    ("format", "image_format_unsupported"),
    ("missing", "image_asset_unavailable"),
    ("resolver_type", "image_resolver_invalid"),
    ("metadata", "image_metadata_unsupported"),
    ("count", "image_count_limit"),
])
def test_invalid_or_missing_bound_image_fails_closed(failure, expected, tmp_path):
    data = png_bytes()
    if failure == "undecodable":
        data = b"not an image"
    elif failure == "format":
        buffer = BytesIO()
        Image.new("RGB", (320, 160)).save(buffer, format="JPEG")
        data = buffer.getvalue()
    elif failure == "metadata":
        data = png_bytes(dpi=(192, 96))
    snapshot = image_snapshot(data, count=5 if failure == "count" else 1)
    if failure == "legacy":
        del snapshot["nodes"][-1]["render_digest"]
    elif failure == "dimensions":
        snapshot["nodes"][-1]["width_px"] += 1
    calls = []
    def resolver(digest):
        calls.append(digest)
        if failure == "missing":
            raise KeyError(digest)
        if failure == "resolver_type":
            return "/arbitrary/local/file.png"
        return b"changed bytes" if failure == "digest" else data
    output = tmp_path / "must-not-exist"
    with pytest.raises(ExportError) as error:
        export_snapshot(snapshot, "docx", output, asset_resolver=resolver)
    assert error.value.code == expected
    assert not output.exists()
    if failure in {"legacy", "count"}:
        assert not calls


@pytest.mark.parametrize("format", ["docx", "xlsx", "pptx"])
def test_decorative_images_have_native_decorative_metadata(format, tmp_path):
    data = png_bytes()
    result = export_snapshot(image_snapshot(data, decorative=True), format, tmp_path, asset_resolver=lambda _: data)
    assert component(result)["decorative"] and component(result)["alt_text"] == ""
    properties = office_picture_properties(result["path"], format)
    namespace = "http://schemas.microsoft.com/office/drawing/2017/decorative"
    assert properties
    for item in properties:
        assert item.get("descr") == ""
        decorative = item.find(f".//{{{namespace}}}decorative")
        assert decorative is not None and decorative.get("val") == "1"


def test_shared_png_is_resolved_once_and_every_node_is_embedded(tmp_path):
    data = png_bytes()
    calls = []
    def resolver(digest):
        calls.append(digest)
        return data
    result = export_snapshot(image_snapshot(data, count=2), "xlsx", tmp_path, asset_resolver=resolver)
    workbook = load_workbook(result["path"])
    assert len(calls) == 1
    assert len(workbook["Images"]._images) == 2
    assert len([entry for entry in result["manifest"]["components"] if entry["representation"] == "native_image"]) == 2
    workbook.close()


def test_png_text_metadata_is_rejected_without_rendering(tmp_path):
    metadata = PngImagePlugin.PngInfo()
    metadata.add_text("Comment", "https://invalid.example/not-an-image-resource")
    data = png_bytes(pnginfo=metadata)
    with pytest.raises(ExportError) as error:
        export_snapshot(image_snapshot(data), "pdf", tmp_path, asset_resolver=lambda _: data)
    assert error.value.code == "image_metadata_unsupported"
    assert not list(tmp_path.iterdir())


@pytest.mark.parametrize("format", FORMATS)
def test_transparent_image_and_optional_schema_defaults(format, tmp_path):
    picture = Image.new("RGBA", (320, 160), (12, 70, 45, 80))
    picture.paste((60, 100, 180, 255), (100, 40, 180, 120))
    stream = BytesIO()
    picture.save(stream, format="PNG")
    data = stream.getvalue()
    snapshot = image_snapshot(data)
    del snapshot["nodes"][-1]["decorative"]
    del snapshot["nodes"][-1]["media_type"]
    result = export_snapshot(snapshot, format, tmp_path, asset_resolver=lambda _: data)
    assert component(result)["decorative"] is False
    if format == "pdf":
        reader = PdfReader(result["path"])
        matching = [obj.get_object() for page in reader.pages for obj in page["/Resources"]["/XObject"].values()
                    if obj.get_object().get("/Width") == 320 and obj.get_object().get("/Height") == 160]
        assert len(matching) == 1 and "/SMask" in matching[0]
        assert matching[0]["/SMask"].get_object().get_data() == picture.getchannel("A").tobytes()
    else:
        with ZipFile(result["path"]) as archive:
            assert data in [archive.read(name) for name in archive.namelist() if "/media/" in name]


@pytest.mark.parametrize("size", [(1600, 400), (400, 1600)])
def test_pdf_image_content_transform_preserves_physical_aspect_ratio(size, tmp_path):
    data = png_bytes(size)
    result = export_snapshot(image_snapshot(data, size=size), "pdf", tmp_path, asset_resolver=lambda _: data)
    reader = PdfReader(result["path"])
    transforms = []
    for page in reader.pages:
        resources = page["/Resources"].get("/XObject", {})
        matrix = None
        for operands, operator in ContentStream(page.get_contents(), reader).operations:
            if operator == b"cm":
                matrix = operands
            elif operator == b"Do":
                image = resources[operands[0]].get_object()
                if (image.get("/Width"), image.get("/Height")) == size:
                    transforms.append(matrix)
    assert len(transforms) == 1
    a, b, c, d, _, _ = (float(value) for value in transforms[0])
    assert b == c == 0
    assert a/d == pytest.approx(size[0]/size[1], rel=1e-5)
    assert a <= 494 and d <= 6.2*72
