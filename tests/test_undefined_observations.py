"""Undefined ratios in independent targets must retain their explicit meaning."""

import copy
import hashlib
from io import BytesIO

from docx import Document
import pytest

from foundry.ingestion import inspect_asset
from foundry.observations import compare_snapshot, inspect_target


def target(*, growth="Not defined", money="40.00", table_growth="Not defined"):
    document = Document()
    document.add_heading("Quarterly revenue", 0)
    document.add_paragraph(f"Current revenue (EUR): {money}")
    document.add_paragraph("Comparison revenue (EUR): 0.00")
    document.add_paragraph(f"Growth: {growth}")
    table = document.add_table(rows=1, cols=5)
    for cell, text in zip(table.rows[0].cells,
                          ["Region", "Current (EUR)", "Comparison (EUR)", "Change (EUR)", "Growth"]):
        cell.text = text
    for cell, text in zip(table.add_row().cells,
                          ["Cedar", "40.00", "0.00", "+40.00", table_growth]):
        cell.text = text
    stream = BytesIO()
    document.save(stream)
    data = stream.getvalue()
    return {"id": "independent-target", "digest": hashlib.sha256(data).hexdigest(),
            **inspect_asset(data, "independent-target.docx")}


def independent_snapshot():
    # The expected records are manually specified, never prepared by the runtime.
    return {
        "facts": {
            "revenue.current": {"value": "40.00", "status": "known"},
            "revenue.comparison": {"value": "0.00", "status": "known"},
            "revenue.growth": {"value": None, "status": "undefined"},
            "region.cedar.current": {"value": "40.00", "status": "known"},
            "region.cedar.comparison": {"value": "0.00", "status": "known"},
            "region.cedar.change": {"value": "40.00", "status": "known"},
            "region.cedar.growth": {"value": None, "status": "undefined"},
        },
        "datasets": {"regional_totals": {"rows": [{"region": "Cedar"}]}},
    }


def test_real_docx_label_and_table_keep_explicit_undefined_ratio(monkeypatch):
    def forbidden(*args, **kwargs):
        raise AssertionError("Target observations must not derive answers from preparation")

    monkeypatch.setattr("foundry.runtime.prepare", forbidden)
    inspection = inspect_target(target())
    assert not inspection["issues"]
    assert all(region["status"] == "mapped" for region in inspection["regions"])
    undefined = [observation for observation in inspection["observations"]
                 if observation["fact_id"] in {"revenue.growth", "region.cedar.growth"}]
    assert len(undefined) == 2
    assert all(observation["value"] is None and observation["status"] == "undefined"
               and observation["display"] == "Not defined" and observation["unit"] == "ratio"
               for observation in undefined)
    assert next(item for item in undefined if item["fact_id"] == "region.cedar.growth")["locator"].endswith("/row[2]/column[5]")
    assert compare_snapshot(independent_snapshot(), inspection)["passed"]


@pytest.mark.parametrize("fact", [
    {"value": None, "status": "missing"},
    {"value": None, "status": "withheld"},
    {"value": None, "status": "not_applicable"},
    {"value": None, "status": "known"},
    {"value": None},
    {"value": "0", "status": "known"},
    {"value": "0", "status": "undefined"},
    None,
])
def test_undefined_observation_requires_null_value_and_undefined_fact_status(fact):
    inspection = inspect_target(target())
    snapshot = independent_snapshot()
    if fact is None:
        snapshot["facts"].pop("region.cedar.growth")
    else:
        snapshot["facts"]["region.cedar.growth"] = copy.deepcopy(fact)
    result = compare_snapshot(snapshot, inspection)
    assert not result["passed"]
    failures = [check for check in result["checks"] if not check["passed"]]
    assert len(failures) == 1 and failures[0]["fact_id"] == "region.cedar.growth"


@pytest.mark.parametrize("token", ["N/A", "undefined", "not defined", "Not Defined", "", "0", "—", "Not defined%"])
def test_other_missing_or_ambiguous_ratio_notations_remain_unresolved(token):
    inspection = inspect_target(target(growth=token, table_growth=token))
    unresolved = [region for region in inspection["regions"] if region["status"] == "needs_decision"]
    assert len(unresolved) == 2
    assert all(not region["observation_ids"] for region in unresolved)
    assert not any(item["fact_id"].endswith(".growth") for item in inspection["observations"])
    assert not any(item["fact_id"].startswith("region.cedar.") for item in inspection["observations"])


def test_undefined_marker_cannot_supply_money_or_replace_explicit_zero_percentage():
    inspection = inspect_target(target(money="Not defined", growth="0.0%", table_growth="0.0%"))
    assert len(inspection["issues"]) == 1
    assert not any(item["fact_id"] == "revenue.current" for item in inspection["observations"])
    ratios = [item for item in inspection["observations"] if item["unit"] == "ratio"]
    assert len(ratios) == 2
    assert all(item["value"] == "0.0" and item.get("status") != "undefined" for item in ratios)
    result = compare_snapshot(independent_snapshot(), inspection)
    assert not result["passed"]
    assert {check["fact_id"] for check in result["checks"] if not check["passed"]} == {
        "revenue.growth", "region.cedar.growth"}
