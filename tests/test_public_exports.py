"""Public aggregate exports retain fact displays and independent chart data."""
import copy
from io import BytesIO
from pathlib import Path
import re
from zipfile import ZipFile

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader
import pytest

from foundry.contracts import Snapshot
from foundry.exporters import ExportError, export_snapshot


def public_snapshot(unit="percent_points"):
    source = {"id": "public-source", "digest": "a" * 64, "filename": "published.xlsx"}
    values = [
        ("sales", "715417", "USD million", "$715,417 million", "known"),
        ("movement", "-0.9", "percent_points", "−0.9%", "known"),
        ("precision", "1234.567890123456789", "units", "1,234.567890123456789 units", "known"),
        ("missing", None, "units", "Not available (NA)", "missing"),
        ("suppressed", None, "units", "Suppressed (S)", "withheld"),
        ("undefined", None, "ratio", "Not defined", "undefined"),
        ("inapplicable", None, "units", "Not applicable (*)", "not_applicable"),
        ("index", "104.2", "index", "104.2 (2023=100)", "known"),
    ]
    facts = {key: {"id": key, "kind": "number", "value": value, "unit": fact_unit,
                   "display": display, "status": status, "definition": "Published aggregate.",
                   "scope": "May release", "sources": [{"asset_id": source["id"],
                   "artifact_sha256": source["digest"], "locator": f"sheet=Table 1;cell=J{index + 1}"}]}
             for index, (key, value, fact_unit, display, status) in enumerate(values)}
    chart_values = ["-0.9", "0.4"] if unit == "percent_points" else ["715417", "692635"]
    nodes = [
        {"id": "summary", "kind": "rich_text", "title": "Summary", "editable": False,
         "mode": "computed", "runs": [{"type": "text", "text": "Published sales: "}, {"type": "fact", "fact_id": "sales"}]},
        {"id": "commentary", "kind": "rich_text", "title": "Commentary", "editable": True,
         "mode": "literal", "runs": [{"type": "text", "text": "Published aggregates retain their reported measurement basis."}]},
        {"id": "published_table", "kind": "table", "title": "Published measures", "dataset_id": "measures"},
        {"id": "published_chart", "kind": "chart", "title": "Published comparison", "dataset_id": "chart_values",
         "category_column": "category", "chart_type": "bar", "axis_unit": unit,
         "series": [{"column_id": "value", "label": "Published value"}]},
    ]
    ids = [n["id"] for n in nodes]
    result = {"schema_version": "1.0", "id": "public-snapshot", "revision": 1, "parent_id": None,
        "title": "Public retail report", "report_type_id": "public-retail", "program": {"id": "public-program", "version": "1.0.0", "digest": "b" * 64},
        "period": {"label": "May 2025", "start": "2025-05-01", "end_exclusive": "2025-06-01",
                   "comparison": {"start": "2024-05-01", "end_exclusive": "2024-06-01"},
                   "timezone": "America/New_York", "as_of": "2025-06-17T08:30:00-04:00"},
        "source_snapshot_digest": "c" * 64, "source_assets": [source], "facts": facts,
        "datasets": {
            "measures": {"id": "measures", "columns": [{"id": "metric", "label": "Metric", "type": "text"},
                {"id": "reported", "label": "Reported value and status", "type": "fact"}],
                "rows": [{"metric": key.title(), "reported": key} for key, *_ in values]},
            "chart_values": {"id": "chart_values", "columns": [{"id": "category", "label": "Category", "type": "text"},
                {"id": "value", "label": "Published value", "type": "decimal", "unit": unit},
                {"id": "private_note", "label": "Unused", "type": "text"}],
                "rows": [{"category": category, "value": value, "private_note": "UNUSED_COLUMN_MUST_NOT_BE_EMBEDDED"}
                         for category, value in zip(["Retail sales", "Comparison"], chart_values)]}},
        "nodes": [{"id": "root", "kind": "section", "title": "Public report", "children": ids}, *nodes],
        "views": [{"id": family, "family": family, "title": family, "node_ids": ids,
                   "coverage": {"scope": "complete", "required_node_ids": ids, "omitted_node_ids": []}, "recipe": {}}
                  for family in ("flow", "grid", "canvas")],
        "findings": [], "metadata": {}, "status": "review_required", "created_at": "2025-06-17T13:00:00Z"}
    return Snapshot.model_validate(result).model_dump(mode="json")


def exported_text(path, fmt):
    if fmt == "pdf":
        return " ".join(page.extract_text() for page in PdfReader(path).pages)
    if fmt == "docx":
        doc = Document(path)
        return "\n".join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
    if fmt == "pptx":
        deck = Presentation(path)
        return "\n".join([s.text for slide in deck.slides for s in slide.shapes if s.has_text_frame] +
                         [c.text for slide in deck.slides for s in slide.shapes if s.has_table for r in s.table.rows for c in r.cells])
    workbook = load_workbook(path, data_only=False)
    try:
        return "\n".join(str(c.value) for ws in workbook for row in ws for c in row if c.value is not None)
    finally:
        workbook.close()


@pytest.mark.parametrize("fmt", ["docx", "xlsx", "pdf", "pptx"])
@pytest.mark.parametrize("unit", ["percent_points", "usd_million"])
def test_public_fact_table_and_separate_chart_preserve_original_meaning(fmt, unit, tmp_path):
    snapshot = public_snapshot(unit)
    before = copy.deepcopy(snapshot)
    result = export_snapshot(snapshot, fmt, tmp_path)
    text = re.sub(r"\s+", " ", exported_text(result["path"], fmt))
    for fact in snapshot["facts"].values():
        assert fact["display"] in text
    assert "EUR" not in text and "Regional" not in text and "regional aggregates" not in text
    assert "UNUSED_COLUMN_MUST_NOT_BE_EMBEDDED" not in text
    assert {entry["node_id"] for entry in result["manifest"]["components"]} == set(snapshot["views"][0]["node_ids"])
    assert "table_display_roundtrip" in {check["check"] for check in result["manifest"]["validation_results"]}
    assert snapshot == before
    expected_values = [-0.9, 0.4] if unit == "percent_points" else [715417, 692635]
    if fmt == "xlsx":
        wb = load_workbook(result["path"], data_only=False)
        try:
            assert not any(c.data_type == "f" for ws in wb for row in ws for c in row)
            chart = wb["Analysis"]._charts[0]
            assert [point.v for point in chart.series[0].val.numRef.numCache.pt] == expected_values
            assert not any(ws._pivots for ws in wb)
            assert "native_chart_values_roundtrip" in {check["check"] for check in result["manifest"]["validation_results"]}
        finally:
            wb.close()
    if fmt == "pptx":
        charts = [s.chart for slide in Presentation(result["path"]).slides for s in slide.shapes if s.has_chart]
        assert list(charts[0].series[0].values) == expected_values
        assert charts[0].series[0].invert_if_negative is False
        with ZipFile(result["path"]) as archive:
            wb = load_workbook(BytesIO(archive.read(next(n for n in archive.namelist() if n.startswith("ppt/embeddings/")))))
            assert "UNUSED_COLUMN_MUST_NOT_BE_EMBEDDED" not in [c.value for ws in wb for row in ws for c in row]
            wb.close()


@pytest.mark.parametrize("fmt", ["docx", "xlsx", "pdf", "pptx"])
def test_fact_cell_missing_reference_cannot_be_exported_as_an_identifier(fmt, tmp_path):
    snapshot = public_snapshot()
    snapshot["datasets"]["measures"]["rows"][0]["reported"] = "unknown-fact"
    with pytest.raises(ExportError, match="schema or reference"):
        export_snapshot(snapshot, fmt, tmp_path)
    assert not list(tmp_path.iterdir())


@pytest.mark.parametrize("fmt", ["docx", "xlsx", "pdf", "pptx"])
def test_explicit_percentage_points_override_legacy_growth_column_name(fmt, tmp_path):
    snapshot = public_snapshot()
    snapshot["datasets"]["measures"] = {"id": "measures", "columns": [
        {"id": "metric", "label": "Metric", "type": "text", "unit": ""},
        {"id": "growth", "label": "Change (percentage points)", "type": "decimal", "unit": "percent_points"}],
        "rows": [{"metric": "Retail", "growth": "-0.9"}]}
    result = export_snapshot(snapshot, fmt, tmp_path)
    text = exported_text(result["path"], fmt)
    assert "-0.9" in text and "-90.0%" not in text
    if fmt == "xlsx":
        wb = load_workbook(result["path"])
        matching = [cell for row in wb["Overview"] for cell in row if cell.value == -0.9]
        assert len(matching) == 1 and "%" not in matching[0].number_format
        wb.close()


def test_formula_like_fact_display_remains_literal_text_in_excel(tmp_path):
    snapshot = public_snapshot()
    display = '=HYPERLINK("https://example.invalid", "text")'
    snapshot["facts"]["missing"].update(kind="text", status="known", value=display, display=display)
    result = export_snapshot(snapshot, "xlsx", tmp_path)
    wb = load_workbook(result["path"], data_only=False)
    matching = [cell for row in wb["Overview"] for cell in row if cell.value == display]
    assert len(matching) == 1 and matching[0].data_type == "s" and matching[0].hyperlink is None
    wb.close()


def test_unsupported_fact_display_glyph_still_blocks_pdf(tmp_path):
    snapshot = public_snapshot()
    snapshot["facts"]["missing"]["display"] = "東京"
    with pytest.raises(ExportError) as error:
        export_snapshot(snapshot, "pdf", tmp_path)
    assert error.value.code == "unsupported_font_glyph"
    assert not list(tmp_path.glob("*.pdf"))


@pytest.mark.parametrize("fmt,kind", [("xlsx", "chart"), ("pptx", "chart"), ("pptx", "table")])
def test_unsupported_multiple_native_regions_block_instead_of_omitting_or_overlaying(fmt, kind, tmp_path):
    snapshot = public_snapshot()
    extra = copy.deepcopy(next(n for n in snapshot["nodes"] if n["kind"] == kind))
    extra["id"] = "additional-component"
    snapshot["nodes"].append(extra)
    snapshot["nodes"][0]["children"].append(extra["id"])
    for view in snapshot["views"]:
        view["node_ids"].append(extra["id"])
        view["coverage"]["required_node_ids"].append(extra["id"])
    with pytest.raises(ExportError) as error:
        export_snapshot(snapshot, fmt, tmp_path)
    assert error.value.code == "unsupported_native_layout"
    assert not list(tmp_path.iterdir())


def test_negative_static_chart_value_does_not_overlap_its_category(monkeypatch):
    from PIL import ImageDraw
    from foundry.exporters import _chart_png
    original = ImageDraw.ImageDraw.text
    bounds = {}
    def recording(draw, xy, text, *args, **kwargs):
        if text in {"Retail sales", "-0.90"}:
            bounds[text] = draw.textbbox(xy, text, font=kwargs.get("font"))
        return original(draw, xy, text, *args, **kwargs)
    monkeypatch.setattr(ImageDraw.ImageDraw, "text", recording)
    snapshot = public_snapshot()
    _chart_png(snapshot, next(n for n in snapshot["nodes"] if n["kind"] == "chart"))
    category, value = bounds["Retail sales"], bounds["-0.90"]
    assert category[2] < value[0] or value[2] < category[0] or category[3] < value[1] or value[3] < category[1]


@pytest.mark.parametrize("fmt", ["xlsx", "pptx"])
def test_chart_display_axis_does_not_change_unscaled_percentage_values(fmt, tmp_path):
    snapshot = public_snapshot()
    next(n for n in snapshot["nodes"] if n["kind"] == "chart")["axis_unit"] = "Published figure (%)"
    result = export_snapshot(snapshot, fmt, tmp_path)
    if fmt == "pptx":
        chart = next(s.chart for slide in Presentation(result["path"]).slides for s in slide.shapes if s.has_chart)
        assert list(chart.series[0].values) == [-0.9, 0.4]
        assert chart.value_axis.tick_labels.number_format == "#,##0.0"
    else:
        workbook = load_workbook(result["path"])
        chart = workbook["Analysis"]._charts[0]
        assert [point.v for point in chart.series[0].val.numRef.numCache.pt] == [-0.9, 0.4]
        assert chart.y_axis.numFmt.formatCode == "#,##0.0"
        workbook.close()


@pytest.mark.parametrize("fmt", ["xlsx", "pptx"])
@pytest.mark.parametrize("negative", [False, True])
def test_native_bar_axes_include_zero_for_one_sided_series(fmt, negative, tmp_path):
    snapshot = public_snapshot("usd_million")
    if negative:
        for row in snapshot["datasets"]["chart_values"]["rows"]:
            row["value"] = "-" + row["value"]
    result = export_snapshot(snapshot, fmt, tmp_path)
    if fmt == "pptx":
        chart = next(s.chart for slide in Presentation(result["path"]).slides for s in slide.shapes if s.has_chart)
        assert (chart.value_axis.maximum_scale if negative else chart.value_axis.minimum_scale) == 0
    else:
        workbook = load_workbook(result["path"])
        axis = workbook["Analysis"]._charts[0].y_axis.scaling
        assert (axis.max if negative else axis.min) == 0
        workbook.close()


def test_docx_static_chart_and_caption_are_kept_on_the_same_page(tmp_path):
    result = export_snapshot(public_snapshot(), "docx", tmp_path)
    paragraphs = Document(result["path"]).paragraphs
    captions = [i for i, p in enumerate(paragraphs) if p.text.startswith("Chart uses the accepted report values.")]
    assert len(captions) == 1
    picture = paragraphs[captions[0] - 1]
    assert picture._p.xpath(".//w:drawing")
    assert picture.paragraph_format.keep_with_next is True


def test_workbook_chart_source_labels_have_room_to_wrap_without_changing_values(tmp_path):
    snapshot = public_snapshot()
    labels = ["Online value: monthly change", "Rolling annual change in published retail sales"]
    for row, label in zip(snapshot["datasets"]["chart_values"]["rows"], labels):
        row["category"] = label
    result = export_snapshot(snapshot, "xlsx", tmp_path)
    workbook = load_workbook(result["path"])
    try:
        source = workbook["Approved data"]
        assert [source.cell(row, 1).value for row in (2, 3)] == labels
        assert [source.cell(row, 2).value for row in (2, 3)] == [-0.9, 0.4]
        assert source.column_dimensions["A"].width >= 38
        assert all(source.cell(row, 1).alignment.wrap_text for row in (2, 3))
        assert source.row_dimensions[3].height >= 30
        chart = workbook["Analysis"]._charts[0]
        assert [point.v for point in chart.series[0].val.numRef.numCache.pt] == [-0.9, 0.4]
        assert "UNUSED_COLUMN_MUST_NOT_BE_EMBEDDED" not in [cell.value for sheet in workbook for row in sheet for cell in row]
    finally:
        workbook.close()
