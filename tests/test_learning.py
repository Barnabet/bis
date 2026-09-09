"""Independent historical reconstruction and registered-policy behavior."""
import copy
import hashlib
from io import BytesIO
import json
from pathlib import Path

from docx import Document
import pytest

from foundry.errors import DomainError
from foundry.ingestion import inspect_asset, rows_from_csv
from foundry.learning import analyze_examples, SELECTION_OPTIONS
from foundry.observations import inspect_target, compare_snapshot
from foundry.runtime import prepare, render_runs, RuntimeBlocked

FIXTURES = Path(__file__).resolve().parents[1] / "fixtures" / "learning"


def asset(data, filename, id):
    return {"id": id, "digest": hashlib.sha256(data).hexdigest(), **inspect_asset(data, filename)}


def case(name):
    directory = FIXTURES / name
    report = asset((directory / "report.docx").read_bytes(), "report.docx", f"{name}_target")
    source_data = (directory / "transactions.csv").read_bytes()
    source = asset(source_data, "transactions.csv", f"{name}_source")
    return {"id": name, "corpus_role": "authoring", "period": json.loads((directory / "period.json").read_text()),
            "report_asset": report, "source_asset": source, "rows": rows_from_csv(source_data)[1],
            "inspection": inspect_target(report)}


def prepared(example, selection="largest_absolute_change"):
    return prepare(example["rows"], example["period"],
                   {key: example["source_asset"][key] for key in ("id", "digest", "filename")},
                   reporting_policy={"selection": selection})


def change_target(example, old, new):
    document = Document(FIXTURES / example["id"] / "report.docx")
    assert any(paragraph.text == old for paragraph in document.paragraphs)
    for paragraph in document.paragraphs:
        if paragraph.text == old:
            paragraph.text = new
    stream = BytesIO()
    document.save(stream)
    changed = copy.deepcopy(example)
    changed["report_asset"] = asset(stream.getvalue(), "changed-report.docx", example["report_asset"]["id"] + "_changed")
    changed["inspection"] = inspect_target(changed["report_asset"])
    return changed


def supported(result):
    return {hypothesis["value"] for hypothesis in result["hypotheses"] if hypothesis["supported"]}


def test_real_docx_observations_are_extracted_without_preparation(monkeypatch):
    def forbidden(*args, **kwargs):
        raise AssertionError("Target inspection must not obtain expected answers from runtime")
    monkeypatch.setattr("foundry.runtime.prepare", forbidden)
    example = case("discriminating")
    inspection = example["inspection"]
    assert inspection["asset_digest"] == hashlib.sha256((FIXTURES / "discriminating" / "report.docx").read_bytes()).hexdigest()
    assert inspection["asset_id"] == example["report_asset"]["id"]
    by_fact = {observation["fact_id"]: observation for observation in inspection["observations"]}
    assert by_fact["revenue.current"]["value"] == "1600.00"
    assert by_fact["revenue.comparison"]["value"] == "1100.00"
    assert by_fact["driver.region"]["value"] == "South"
    assert by_fact["region.north.current"]["display"] == "1,000.00"
    assert by_fact["regional_totals.region_order"]["value"] == ["East", "North", "South"]
    assert by_fact["region.north.current"]["method"] == "native_table_cell"
    assert "/row[" in by_fact["region.north.current"]["locator"]
    assert len(inspection["regions"]) == len(example["report_asset"]["profile"]["regions"])
    assert not inspection["issues"]


def test_one_ambiguous_example_preserves_all_compatible_policies():
    example = case("ambiguous")
    before = copy.deepcopy(example)
    result = analyze_examples([example])
    assert example == before
    assert supported(result) == {option["value"] for option in SELECTION_OPTIONS}
    assert any("distinguishing" in item for item in result["assumptions"])
    assert len(result["coverage"]) == len(example["inspection"]["regions"])
    assert all(item["status"] == "mapped" for item in result["coverage"])


def test_second_example_discriminates_with_located_failures():
    result = analyze_examples([case("ambiguous"), case("discriminating")])
    assert supported(result) == {"largest_absolute_change"}
    failures = {h["value"]: [check for check in h["checks"] if not check["passed"]] for h in result["hypotheses"]}
    assert not failures["largest_absolute_change"]
    for policy, actual in [("largest_current_revenue", "North"), ("largest_percentage_change", "East")]:
        assert len(failures[policy]) == 1
        check = failures[policy][0]
        assert check["example_id"] == "discriminating"
        assert check["expected"] == "South" and check["actual"] == actual
        assert check["fact_id"] == "driver.region" and "word/document.xml" in check["locator"]


def test_contradictory_targets_do_not_silently_choose_best_fit():
    first = case("discriminating")
    conflicting = change_target(first, "Selected region: South", "Selected region: East")
    conflicting["id"] = "contradictory_case"
    result = analyze_examples([first, conflicting])
    assert not supported(result)
    assert any("contradictory" in message for message in result["assumptions"])
    assert all(any(not check["passed"] for check in h["checks"]) for h in result["hypotheses"])


def test_wrong_numeric_target_defeats_every_selection_policy():
    example = change_target(case("discriminating"), "Current revenue (EUR): 1,600.00", "Current revenue (EUR): 1,999.00")
    result = analyze_examples([example])
    assert not supported(result)
    for hypothesis in result["hypotheses"]:
        numeric_failure = next(check for check in hypothesis["checks"] if check.get("fact_id") == "revenue.current")
        assert not numeric_failure["passed"]
        assert numeric_failure["expected"] == "1999.00" and numeric_failure["actual"] == "1600.00"


def test_reserved_example_rejected_before_any_candidate_runs(monkeypatch):
    public, reserved = case("ambiguous"), case("discriminating")
    reserved["corpus_role"] = "reserved"
    def forbidden(*args, **kwargs):
        raise AssertionError("Must reject the entire request before executing public or reserved cases")
    monkeypatch.setattr("foundry.learning.prepare", forbidden)
    with pytest.raises(DomainError) as error:
        analyze_examples([public, reserved])
    assert error.value.code == "RESERVED_EXAMPLE_DENIED" and error.value.status_code == 403


def test_development_case_is_allowed_but_not_described_as_holdout():
    example = case("discriminating")
    example["corpus_role"] = "development"
    result = analyze_examples([example])
    assert supported(result) == {"largest_absolute_change"}
    assert any("not a protected holdout" in item for item in result["limitations"])


@pytest.mark.parametrize("kind", ["image", "chart", "header", "footer", "footnote", "textbox", "unknown"])
def test_unsupported_target_regions_cannot_disappear_from_coverage(kind):
    example = case("ambiguous")
    example["report_asset"]["profile"]["regions"].append({"id": "unsupported", "kind": kind,
        "text": "Required historical content", "locator": f"original/{kind}"})
    example["inspection"] = inspect_target(example["report_asset"])
    result = analyze_examples([example])
    row = next(row for row in result["coverage"] if row["region_id"] == "unsupported")
    assert row["status"] == "needs_decision" and row["component_id"] is None
    assert row["text"] == "Required historical content" and row["locator"] == f"original/{kind}"
    assert any("Unresolved historical regions" in message for message in result["assumptions"])


def test_full_summary_recognition_rejects_unrecognized_trailing_claim():
    example = case("discriminating")
    paragraph = "Posted revenue reached €1,600.00. Growth against the comparison period was 45.5%. South recorded the largest absolute regional movement, +€400.00, against a total change of +€500.00."
    target = copy.deepcopy(example["report_asset"])
    target["profile"]["regions"] = [{"id": "summary", "kind": "paragraph", "locator": "paragraph[0]", "text": paragraph}]
    inspection = inspect_target(target)
    assert compare_snapshot(prepared(example), inspection)["passed"]
    assert next(o["value"] for o in inspection["observations"] if o["fact_id"] == "policy.selection") == "largest_absolute_change"
    target["profile"]["regions"][0]["text"] += " A new campaign explains this performance."
    rejected = inspect_target(target)
    assert not rejected["observations"] and rejected["regions"][0]["status"] == "needs_decision"


@pytest.mark.parametrize("bad_value", ["1.600,00", "1,60.00", "1600,00", "1 600.00", "1.600"])
def test_locale_or_ambiguous_number_syntax_remains_unresolved(bad_value):
    example = change_target(case("discriminating"), "Current revenue (EUR): 1,600.00", "Current revenue (EUR): " + bad_value)
    region = next(region for region in example["inspection"]["regions"] if region.get("text", "").startswith("Current revenue"))
    assert region["status"] == "needs_decision"
    assert not region["observation_ids"]


def test_empty_observation_set_cannot_pass_comparison():
    example = case("ambiguous")
    inspection = {"asset_id": "empty", "observations": [], "regions": [], "issues": []}
    result = compare_snapshot(prepared(example), inspection)
    assert not result["passed"] and result["observation_count"] == 0
    assert result["checks"][0]["method"] == "observation_coverage"


def test_target_heading_does_not_supply_selection_evidence():
    example = case("ambiguous")
    example["report_asset"]["profile"]["regions"] = [{"id": "heading", "kind": "paragraph", "text": "Quarterly revenue", "locator": "paragraph[0]"}]
    example["inspection"] = inspect_target(example["report_asset"])
    result = analyze_examples([example])
    assert not supported(result)
    assert result["coverage"][0]["status"] == "mapped"
    assert any("no observed selected region" in item for item in result["assumptions"])


def test_display_rounding_compares_observed_precision_not_exact_ratio():
    example = case("discriminating")
    snapshot = prepared(example).model_dump(mode="json")
    assert snapshot["facts"]["revenue.growth"]["value"].startswith("0.454545")
    assert compare_snapshot(snapshot, example["inspection"])["passed"]
    snapshot["facts"]["revenue.growth"]["value"] = "0.4544"
    result = compare_snapshot(snapshot, example["inspection"])
    assert not next(check for check in result["checks"] if check["fact_id"] == "revenue.growth")["passed"]


@pytest.mark.parametrize("mutation", ["missing", "order"])
def test_table_membership_and_order_are_checked_independently_of_cells(mutation):
    example = case("discriminating")
    snapshot = prepared(example).model_dump(mode="json")
    rows = snapshot["datasets"]["regional_totals"]["rows"]
    if mutation == "missing":
        rows.pop()
    else:
        rows.reverse()
    checks = compare_snapshot(snapshot, example["inspection"])["checks"]
    assert not next(check for check in checks if check["fact_id"] == "regional_totals.region_order")["passed"]
    assert all(check["passed"] for check in checks if check["fact_id"] != "regional_totals.region_order")


@pytest.mark.parametrize("selection,winner,dependency,wording", [
    ("largest_absolute_change", "South", ".change", "largest absolute regional movement"),
    ("largest_percentage_change", "East", ".growth", "largest absolute percentage revenue movement"),
    ("largest_current_revenue", "North", ".current", "highest current-period regional revenue"),
])
def test_new_period_selection_changes_facts_dependencies_and_computed_prose(selection, winner, dependency, wording):
    training = analyze_examples([case("ambiguous")])
    assert selection in supported(training)
    snapshot = prepared(case("discriminating"), selection)
    assert snapshot.facts["driver.region"].value == winner
    assert all(fid.endswith(dependency) for fid in snapshot.facts["driver.region"].inputs)
    assert snapshot.facts["revenue.current"].value == "1600.00"
    assert snapshot.facts["revenue.change"].value == "500.00"
    summary = next(node for node in snapshot.nodes if node.id == "summary")
    assert wording in render_runs(summary.runs, snapshot.facts)
    assert winner in render_runs(summary.runs, snapshot.facts)
    assert snapshot.metadata["reporting_policy"]["selection"] == selection


def test_different_explicit_policies_have_distinct_default_program_identity():
    example = case("discriminating")
    assert len({prepared(example, option["value"]).program.digest for option in SELECTION_OPTIONS}) == 3


def test_percentage_policy_excludes_undefined_ratios():
    example = case("discriminating")
    for row in example["rows"]:
        if row["region"] == "East" and "comparison" in row["transaction_id"]:
            row["amount"] = "0.00"
        if row["region"] == "East" and "current" in row["transaction_id"]:
            row["amount"] = "999999.00"
    snapshot = prepared(example, "largest_percentage_change")
    assert snapshot.facts["region.east.growth"].value is None
    assert snapshot.facts["driver.region"].value == "South"
    assert "undefined ratios are excluded" in snapshot.facts["driver.region"].definition


def test_all_undefined_ratios_explicitly_block_percentage_selection():
    example = case("discriminating")
    for row in example["rows"]:
        if "comparison" in row["transaction_id"]:
            row["amount"] = "0.00"
    with pytest.raises(RuntimeBlocked) as error:
        prepared(example, "largest_percentage_change")
    assert error.value.findings[0].code == "SELECTION_UNDEFINED"
    assert prepared(example, "largest_current_revenue").facts["driver.region"].value == "North"


def test_percentage_selection_uses_absolute_growth_and_stable_ties():
    example = case("discriminating")
    amounts = {"East": "15.00", "North": "500.00", "South": "10.00"}
    for row in example["rows"]:
        if "current" in row["transaction_id"]:
            row["amount"] = amounts[row["region"]]
    assert prepared(example, "largest_percentage_change").facts["driver.region"].value == "South"
    for row in example["rows"]:
        if "current" in row["transaction_id"] and row["region"] == "East":
            row["amount"] = "1.00"
    assert prepared(example, "largest_percentage_change").facts["driver.region"].value == "East"
    example["rows"].reverse()
    assert prepared(example, "largest_percentage_change").facts["driver.region"].value == "East"


@pytest.mark.parametrize("policy", [{}, {"selection": "arbitrary_code"}, {"selection": []}, {"selection": True},
                                  {"selection": "largest_absolute_change", "code": "ignored"}, "largest_absolute_change"])
def test_reporting_policy_allowlist_rejects_unsupported_inputs(policy):
    example = case("ambiguous")
    with pytest.raises(RuntimeBlocked) as error:
        prepare(example["rows"], example["period"],
                {key: example["source_asset"][key] for key in ("id", "digest", "filename")}, reporting_policy=policy)
    assert error.value.findings[0].code == "REPORTING_POLICY_INVALID"


def test_inspection_must_match_immutable_target_identity():
    example = case("ambiguous")
    example["inspection"]["asset_digest"] = "f" * 64
    with pytest.raises(DomainError) as error:
        analyze_examples([example])
    assert error.value.code == "TARGET_IDENTITY_INVALID"


def test_invalid_source_is_reported_as_failed_reconstruction():
    example = case("ambiguous")
    example["rows"][0]["amount"] = "invalid"
    result = analyze_examples([example])
    assert not supported(result)
    assert all(h["checks"][0]["code"] == "INPUT_DRIFT" for h in result["hypotheses"])


def test_recognized_literal_disclosure_is_an_exact_evaluation_obligation():
    example = case("discriminating")
    observation = next(item for item in example["inspection"]["observations"] if item["fact_id"] == "node.commentary.text")
    assert observation["method"] == "exact_literal_text"
    snapshot = prepared(example).model_dump(mode="json")
    assert compare_snapshot(snapshot, example["inspection"])["passed"]
    commentary = next(node for node in snapshot["nodes"] if node["id"] == "commentary")
    commentary["runs"] = [{"type": "text", "text": "An unsupported explanation replaced the required wording."}]
    checks = compare_snapshot(snapshot, example["inspection"])["checks"]
    assert not next(check for check in checks if check["fact_id"] == "node.commentary.text")["passed"]
    assert all(check["passed"] for check in checks if check["fact_id"] != "node.commentary.text")
