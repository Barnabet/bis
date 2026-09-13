"""Static Census MARTS observations, independent of generated report output.

Only the explicit published table contract is supported. External relationships
are evidence, never resources to retrieve; formulas and their caches never
supply business values. This module has no manifest, model or network dependency.
"""
from __future__ import annotations

import calendar
import copy
from datetime import date, datetime, time, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import hashlib
from io import BytesIO
import re
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

from .errors import DomainError
from .ingestion import MAX_BYTES, MAX_ROWS, inspect_asset, inspect_zip

FAMILY = 'census_marts'
NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
      'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
BASIS = 'seasonally_adjusted_holiday_trading_day_adjusted_not_price_adjusted'
MAX_STATIC_CELLS = 200_000
_MONTH = {name.lower(): index for index, name in enumerate(calendar.month_name) if name}
_MONTH.update({name.lower(): index for index, name in enumerate(calendar.month_abbr) if name})
_LEVELS = [
    ('sales_usd_million', 12, 'total', None),
    ('retail_sales_usd_million', 17, 'Retail', None),
    ('nonstore_sales_usd_million', 64, 'Nonstore retailers', '454'),
    ('food_services_sales_usd_million', 67, 'Food services & drinking places', '722'),
]
_RATES = [
    ('month_on_month_pct', 'C15', 'J12', 'K12'),
    ('year_on_year_pct', 'D15', 'J12', 'M12'),
    ('retail_month_on_month_pct', 'C21', 'J17', 'K17'),
    ('retail_year_on_year_pct', 'D21', 'J17', 'M17'),
    ('nonstore_year_on_year_pct', 'D51', 'J64', 'M64'),
    ('food_services_year_on_year_pct', 'D53', 'J67', 'M67'),
]


def _digest(data):
    return hashlib.sha256(data).hexdigest()


def _safe_value(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return str(value)
    return value


def _numeric_state(raw, formula=None):
    if formula is not None:
        return None, 'undefined'
    if raw is None or isinstance(raw, str) and not raw.strip():
        return None, 'missing'
    if raw == '(S)':
        return None, 'withheld'
    if raw == '(NA)':
        return None, 'undefined'
    if raw == '(*)':
        return None, 'missing'
    if isinstance(raw, bool):
        return None, 'undefined'
    token = str(raw)
    if not re.fullmatch(r'[+-]?(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.\d+)?', token):
        return None, 'undefined'
    try:
        result = Decimal(token.replace(',', ''))
    except InvalidOperation:
        return None, 'undefined'
    return (format(result, 'f'), 'known') if result.is_finite() else (None, 'undefined')


def extract_static_xlsx(data: bytes) -> dict:
    """Read bounded workbook bytes without formula evaluation or external I/O.

    The explicit static mode belongs to the archive validator; generic uploads
    retain their stricter relationship gate. Returned cells preserve original
    values, formulas, untrusted caches and exact sheet/cell locators separately.
    """
    if not isinstance(data, bytes) or not data or len(data) > MAX_BYTES:
        raise DomainError('UPLOAD_LIMIT', 'Supply nonempty workbook bytes no larger than 20 MB.')
    digest = _digest(data)
    archive = inspect_zip(data, 'xlsx', mode='static_xlsx')
    external, formula_xml, sheet_parts = [], {}, {}
    try:
        for name in archive.namelist():
            if name.lower().endswith('.rels'):
                for relation in ET.fromstring(archive.read(name)):
                    if relation.get('TargetMode', '').lower() == 'external':
                        external.append({'part': name, 'id': relation.get('Id'),
                                         'type': relation.get('Type'), 'target': relation.get('Target'),
                                         'target_mode': relation.get('TargetMode'), 'resolved': False,
                                         'source_digest': digest})
        workbook = ET.fromstring(archive.read('xl/workbook.xml'))
        relations = {r.get('Id'): r for r in ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))}
        for sheet in workbook.findall('s:sheets/s:sheet', NS):
            relation = relations.get(sheet.get(f"{{{NS['r']}}}id"))
            if relation is None or relation.get('TargetMode', '').lower() == 'external':
                raise DomainError('FORMAT_INVALID', 'A sheet must bind to a local workbook part.')
            target = relation.get('Target', '')
            if not target or '\\' in target or ':' in target or '..' in target.split('/'):
                raise DomainError('ARCHIVE_INVALID', 'Unsafe internal worksheet relationship.')
            part = target.lstrip('/') if target.startswith('/') else 'xl/' + target
            if not part.startswith('xl/worksheets/') or part not in archive.namelist():
                raise DomainError('FORMAT_INVALID', 'A worksheet relationship does not identify a local worksheet.')
            sheet_parts[sheet.get('name')] = part
            root = ET.fromstring(archive.read(part))
            formula_xml[sheet.get('name')] = {
                c.get('r'): {'formula': c.find('s:f', NS), 'cache': c.findtext('s:v', default=None, namespaces=NS),
                             'raw_xml_value': c.findtext('s:v', default=None, namespaces=NS)}
                for c in root.findall('s:sheetData/s:row/s:c', NS)
            }
        book = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        cells, sheets = [], []
        try:
            estimated = 0
            for sheet in book:
                estimated += sheet.max_row * sheet.max_column
                if sheet.max_row > MAX_ROWS or sheet.max_column > 200 or estimated > MAX_STATIC_CELLS:
                    raise DomainError('INPUT_LIMIT', 'Workbook dimensions exceed the static extraction limit.')
                sheets.append({'name': sheet.title, 'rows': sheet.max_row, 'columns': sheet.max_column,
                               'hidden': sheet.sheet_state != 'visible', 'part': sheet_parts[sheet.title]})
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.value is None:
                            continue
                        xml = formula_xml[sheet.title].get(cell.coordinate, {})
                        f = xml.get('formula')
                        formula = str(cell.value) if cell.data_type == 'f' else ('=' + (f.text or '') if f is not None else None)
                        raw = _safe_value(cell.value)
                        value, status = _numeric_state(raw, formula)
                        cells.append({'sheet': sheet.title, 'address': cell.coordinate,
                                      'raw_value': raw, 'raw_xml_value': xml.get('raw_xml_value'),
                                      'value': value, 'status': status, 'value_type': cell.data_type,
                                      'formula': formula, 'cached_value': xml.get('cache') if formula is not None else None,
                                      'cached_value_trusted': False if formula is not None else None,
                                      'number_format': cell.number_format,
                                      'locator': f'sheet={sheet.title};cell={cell.coordinate}',
                                      'source_digest': digest})
        finally:
            book.close()
    except DomainError:
        raise
    except Exception as exc:
        raise DomainError('FORMAT_INVALID', 'The static workbook structure could not be read.') from exc
    finally:
        archive.close()
    return {'source_digest': digest, 'byte_count': len(data), 'sheets': sheets, 'cells': cells,
            'external_relationships': external, 'formula_evaluation': 'never', 'external_resolution': 'never'}


def _fail(message, locator=None):
    raise DomainError('PUBLIC_SOURCE_DRIFT', message, details={'locator': locator} if locator else None)


def _month_label(text):
    match = re.fullmatch(r'\s*([A-Za-z]+)\.?\s*(?:3)?\s*', str(text))
    if not match or match[1].lower() not in _MONTH:
        _fail('Unrecognized monthly column label.')
    return _MONTH[match[1].lower()]


def _shift(period, offset):
    year, month = map(int, period.split('-'))
    number = year * 12 + month - 1 + offset
    return f'{number // 12:04d}-{number % 12 + 1:02d}'


def _date_text(text):
    match = re.search(r'\b([A-Za-z]+) (\d{1,2}), (\d{4})\b', text)
    if not match or match[1].lower() not in _MONTH:
        _fail('A publication date must be explicit in the source footnote.')
    return date(int(match[3]), _MONTH[match[1].lower()], int(match[2])).isoformat()


def _observation(metric, period, value, status, unit, decimals, locator, definition, **extra):
    return {'metric_id': 'census.' + metric, 'period': period, 'value': value, 'status': status,
            'unit': unit, 'display_decimals': decimals, 'locator': locator, 'definition': definition, **extra}


def _blocking_issues(result):
    for issue in result['issues']:
        issue.setdefault('severity', 'block')
    return result


def inspect_source(data: bytes) -> dict:
    static = extract_static_xlsx(data)
    lookup = {(c['sheet'], c['address']): c for c in static['cells']}

    def cell(sheet, address):
        return lookup.get((sheet, address), {'raw_value': None, 'value': None, 'status': 'missing',
                                            'formula': None, 'locator': f'sheet={sheet};cell={address}'})

    def raw(sheet, address):
        return cell(sheet, address)['raw_value']

    def require(sheet, address, predicate, message):
        if not predicate(raw(sheet, address)):
            _fail(message, f'sheet={sheet};cell={address}')

    if [s['name'] for s in static['sheets']] != ['Table 1.', 'Table 2.', 'Table 3.']:
        _fail('Census MARTS requires the three explicitly named published tables.')
    require('Table 1.', 'A4', lambda x: 'millions of dollars' in str(x), 'Table 1 must state millions of dollars.')
    require('Table 2.', 'A3', lambda x: 'shown as percents' in str(x), 'Table 2 must state percent units.')
    require('Table 1.', 'J6', lambda x: x == 'Adjusted2', 'Adjusted and unadjusted column identities changed.')
    require('Table 1.', 'A76', lambda x: all(t in str(x) for t in ['seasonal variation', 'holiday and trading day', 'not for price changes']),
            'The adjustment basis must be explicit and unchanged.')
    require('Table 2.', 'A58', lambda x: 'derived from adjusted estimates' in str(x), 'Change estimates must refer to adjusted levels.')
    year = raw('Table 1.', 'J7')
    if not isinstance(year, int) or not 1900 <= year <= 2200:
        _fail('Current-year header is not an explicit supported year.', 'sheet=Table 1.;cell=J7')
    month = _month_label(raw('Table 1.', 'J8'))
    period = f'{year:04d}-{month:02d}'
    for address, offset in [('K8', -1), ('L8', -2), ('M8', -12), ('N8', -13)]:
        if _month_label(raw('Table 1.', address)) != int(_shift(period, offset)[5:]):
            _fail('Comparison column period changed.', f'sheet=Table 1.;cell={address}')
    require('Table 1.', 'M7', lambda x: x == year - 1, 'Year-ago column year changed.')
    for address, status in [('J9', '(a)'), ('K9', '(p)'), ('L9', '(r)'), ('M9', '(r)'), ('N9', '(r)')]:
        require('Table 1.', address, lambda x, status=status: x == status, 'Publication maturity marker changed.')
    match = re.fullmatch(r'([A-Za-z]+)\.? (\d{4}) Advance', str(raw('Table 2.', 'C8')))
    if not match or _MONTH.get(match[1].lower()) != month or int(match[2]) != year:
        _fail('Table 2 current period disagrees with Table 1.', 'sheet=Table 2.;cell=C8')
    dates = [_date_text(str(raw(sheet, address))) for sheet, address in
             [('Table 1.', 'A85'), ('Table 2.', 'A59'), ('Table 3.', 'A42')]]
    if len(set(dates)) != 1:
        _fail('The three source table publication dates disagree.')
    vintage = dates[0]
    if vintage[:7] <= period:
        _fail('A Census monthly release must follow its reference month.', 'sheet=Table 1.;cell=A85')
    for sheet, row, label, code in [
        ('Table 1.', 11, 'Retail & food services', None), *[('Table 1.', r, label, code) for _, r, label, code in _LEVELS],
        ('Table 2.', 14, 'Retail & food services', None), ('Table 2.', 15, 'total', None),
        ('Table 2.', 21, 'Retail', None), ('Table 2.', 51, 'Nonstore retailers', '454'),
        ('Table 2.', 53, 'Food services & drinking places', '722')]:
        actual = str(raw(sheet, f'B{row}')).strip().rstrip(' .,…')
        if actual != label or code and str(raw(sheet, f'A{row}')) != code:
            _fail('Census category or NAICS identity changed.', f'sheet={sheet};row={row}')
    observations, issues = [], []
    # These cells are sales levels, not signed movements. Validate their domain
    # independently of the change-table reconciliation: a coherently altered
    # percentage table must not legitimize a negative sales estimate.
    for _, row, _, _ in _LEVELS:
        for column in 'JKLMN':
            address = f'{column}{row}'
            source = cell('Table 1.', address)
            if source['status'] == 'known' and Decimal(source['value']) < 0:
                issues.append({'code': 'PUBLIC_VALUE_INVALID', 'locator': source['locator'],
                               'value': source['value'], 'unit': 'usd_million',
                               'message': 'Published sales levels must be nonnegative; signed changes are separate measures.'})
                lookup[('Table 1.', address)] = {**source, 'value': None, 'status': 'undefined'}

    def add_cell(metric, address, observed_period, unit='usd_million', decimals=0, sheet='Table 1.', **extra):
        c = cell(sheet, address)
        observation = _observation(metric, observed_period, c['value'], c['status'], unit, decimals,
                                   c['locator'], BASIS, source_digest=static['source_digest'], **extra)
        observations.append(observation)
        if c['status'] != 'known':
            issues.append({'code': 'PUBLIC_VALUE_UNAVAILABLE', 'metric_id': observation['metric_id'],
                           'locator': c['locator'], 'status': c['status'], 'raw_value': c['raw_value'],
                           'message': 'Required published value is missing, withheld, undefined or formula-dependent.'})
        return observation

    for metric, row, _, _ in _LEVELS:
        for column, offset, maturity in [('J', 0, 'advance'), ('K', -1, 'preliminary'), ('L', -2, 'revised'),
                                        ('M', -12, 'revised'), ('N', -13, 'revised')]:
            add_cell(metric, f'{column}{row}', _shift(period, offset), publication_status=maturity)
    add_cell('sales_headline_usd_million', 'J12', period, decimals=-2,
             display_unit='usd_billion', display_scale='1000', display_unit_decimals=1)
    for metric, published, current, prior in _RATES:
        numerator, denominator = cell('Table 1.', current), cell('Table 1.', prior)
        published_cell = cell('Table 2.', published)
        value, status = None, 'undefined'
        if numerator['status'] == denominator['status'] == 'known':
            base = Decimal(denominator['value'])
            if base != 0:
                value = format(((Decimal(numerator['value']) / base - 1) * 100).quantize(Decimal('.1'), rounding=ROUND_HALF_UP), 'f')
                status = 'known'
        else:
            status = next(c['status'] for c in (numerator, denominator) if c['status'] != 'known')
        locator = f"{numerator['locator']}|{denominator['locator']}"
        observations.append(_observation(metric, period, value, status, 'percent_points', 1, locator, BASIS,
                            source_digest=static['source_digest'], method='change_from_published_levels',
                            inputs=[{k: c[k] for k in ('locator', 'value', 'status')} for c in (numerator, denominator)],
                            source_locators=[numerator['locator'], denominator['locator']],
                            published_value=published_cell['value'], published_locator=published_cell['locator']))
        if status != 'known' or published_cell['status'] != 'known' or Decimal(value) != Decimal(published_cell['value']):
            issues.append({'code': 'PUBLIC_SOURCE_RECONCILIATION', 'metric_id': 'census.' + metric,
                           'computed': value, 'published': published_cell['value'], 'locator': locator,
                           'message': 'Calculated change is unavailable or does not reconcile with the published change table.'})
    add_cell('rolling_three_month_year_on_year_pct', 'H15', period, unit='percent_points', decimals=1, sheet='Table 2.',
             method='published_rate', independent_recomputation_available=False,
             limitation='The workbook does not contain all six level months needed for independent recomputation.')
    return _blocking_issues({'family': FAMILY, 'parser_version': 1, 'source_digest': static['source_digest'], 'period': period,
            'vintage': vintage, 'observations': observations, 'issues': issues,
            'metadata': {'basis': BASIS, 'units': ['usd_million', 'percent_points'],
                         'data_level': 'published_monthly_survey_estimates_not_transactions',
                         'static_extraction': static,
                         'vintage_locators': ['sheet=Table 1.;cell=A85', 'sheet=Table 2.;cell=A59', 'sheet=Table 3.;cell=A42']}})


def inspect_report(asset: dict) -> dict:
    """Read only PDF text regions; no source workbook or expected values enter."""
    digest = str(asset.get('digest', '')).removeprefix('sha256:')
    if not isinstance(asset.get('id'), str) or not re.fullmatch('[0-9a-f]{64}', digest):
        raise DomainError('TARGET_IDENTITY_INVALID', 'A report needs its original asset id and digest.')
    regions = copy.deepcopy(asset.get('profile', {}).get('regions', []))
    observations, issues = [], []
    result = {'asset_id': asset['id'], 'asset_digest': digest, 'family': FAMILY, 'parser_version': 1,
              'observations': observations, 'regions': regions, 'issues': issues,
              'report_period': None, 'vintage': None}
    for index, region in enumerate(regions):
        region.setdefault('id', f'page{index + 1}')
        region.setdefault('locator', f'page={index + 1}')
        region.update(status='needs_decision', component_id=None, observation_ids=[])
    title = r'ADVANCE MONTHLY SALES FOR RETAIL AND FOOD SERVICES,\s*([A-Z]+) (\d{4})'
    candidates = [(r, re.search(title, r.get('text', ''))) for r in regions]
    candidates = [(r, m) for r, m in candidates if m]
    if len(candidates) != 1 or candidates[0][1][1].lower() not in _MONTH:
        issues.append({'code': 'PUBLIC_REPORT_IDENTITY', 'message': 'Expected exactly one explicit Census report title and period.'})
        return _blocking_issues(result)
    region, heading = candidates[0]
    period = f'{int(heading[2]):04d}-{_MONTH[heading[1].lower()]:02d}'
    result['report_period'] = period
    text = ' '.join(region['text'].split())
    release = re.search(r'FOR RELEASE AT .*?([A-Z]+ \d{1,2}, \d{4})', text)
    if not release:
        issues.append({'code': 'PUBLIC_REPORT_VINTAGE', 'message': 'Report release date is missing.'})
        return _blocking_issues(result)
    result['vintage'] = _date_text(release[1].title())
    if result['vintage'][:7] <= period:
        issues.append({'code': 'PUBLIC_REPORT_VINTAGE',
                       'message': 'A Census monthly report publication date must follow its reference month.'})
        return _blocking_issues(result)
    if not all(t in text.lower() for t in ['seasonal variation', 'trading-day differences', 'not for price changes']):
        issues.append({'code': 'PUBLIC_REPORT_BASIS', 'message': 'Report adjustment basis is not explicit.'})
        return _blocking_issues(result)

    def add(r, metric, value, unit, decimals, locator, display, observed_period=period, method='published_report_narrative'):
        numeric, status = _numeric_state(value)
        observation = _observation(metric, observed_period, numeric, status, unit, decimals, locator, BASIS,
                                   source_digest=digest, method=method, display=display,
                                   id=f"{asset['id']}:{r['id']}:{len(observations) + 1}", region_id=r['id'])
        observations.append(observation)
        r['observation_ids'].append(observation['id'])
        r.update(status='partially_mapped', component_id='census_published_metrics')
        if status != 'known':
            issues.append({'code': 'PUBLIC_REPORT_VALUE_UNAVAILABLE', 'metric_id': observation['metric_id'],
                           'status': status, 'locator': locator, 'message': 'The report does not publish a known numeric value.'})

    number = r'(\d+(?:\.\d+)?)'
    change = rf'(up|down) {number} percent'
    patterns = [
        ('month_on_month_pct', rf'were \$\d+(?:\.\d+)? billion, {change}'),
        ('year_on_year_pct', rf'from the previous month, and {change}'),
        ('rolling_three_month_year_on_year_pct', rf'Total sales for the .*? period were {change}'),
        ('retail_month_on_month_pct', rf'Retail trade sales were {change}'),
        ('retail_year_on_year_pct', rf'Retail trade sales were .*?, and {change}'),
        ('nonstore_year_on_year_pct', rf'Nonstore retailers were {change}'),
        ('food_services_year_on_year_pct', rf'food service(?:s)? and drinking places were {change}'),
    ]
    sales = list(re.finditer(r'were \$(\d+(?:\.\d+)?) billion', text))
    if len(sales) == 1:
        match = sales[0]
        add(region, 'sales_headline_usd_million', str(Decimal(match[1]) * 1000), 'usd_million', -2,
            f"{region['locator']};normalized_text_chars={match.start()}:{match.end()}", match[0])
    else:
        issues.append({'code': 'PUBLIC_REPORT_METRIC', 'metric_id': 'census.sales_headline_usd_million', 'message': 'Missing or ambiguous headline sales.'})
    for metric, pattern in patterns:
        matches = list(re.finditer(pattern, text, re.IGNORECASE))
        if len(matches) != 1:
            issues.append({'code': 'PUBLIC_REPORT_METRIC', 'metric_id': 'census.' + metric, 'message': 'Missing or ambiguous signed narrative value.'})
            continue
        match = matches[0]
        value = Decimal(match[2]) * (-1 if match[1].lower() == 'down' else 1)
        add(region, metric, str(value), 'percent_points', 1,
            f"{region['locator']};normalized_text_chars={match.start()}:{match.end()}", match[0])

    def table_lines(title_prefix):
        matched = [r for r in regions if title_prefix in r.get('text', '')]
        if len(matched) != 1:
            issues.append({'code': 'PUBLIC_REPORT_TABLE', 'message': f'Missing or ambiguous {title_prefix}.'})
            return None
        return matched[0]

    table1 = table_lines('Table 1.  Estimated Monthly Sales')
    table2 = table_lines('Table 2.  Estimated Change in Monthly Sales')
    for table in (table1, table2):
        if table:
            source_note = re.search(r'Source:\s*U\.S\. Census Bureau, Advance Monthly Retail Trade Survey,\s*([A-Za-z]+ \d{1,2}, \d{4})', table['text'])
            if not source_note or _date_text(source_note[1]) != result['vintage']:
                issues.append({'code': 'PUBLIC_REPORT_VINTAGE', 'locator': table['locator'],
                               'message': 'The table publication date is missing or disagrees with the report.'})
    if table2:
        header = re.search(r'([A-Za-z]+)\.? (\d{4}) Advance', table2['text'])
        if not header or _MONTH.get(header[1].lower()) != int(period[5:]) or header[2] != period[:4]:
            issues.append({'code': 'PUBLIC_REPORT_PERIOD', 'locator': table2['locator'],
                           'message': 'The change table period disagrees with the report title.'})
    row_patterns = [('sales_usd_million', r'^\s*total\s'), ('retail_sales_usd_million', r'^\s*Retail\s+[.·…]'),
                    ('nonstore_sales_usd_million', r'^\s*454\s+Nonstore retailers'),
                    ('food_services_sales_usd_million', r'^\s*722\s+Food services & drinking places')]

    def tokens(r, pattern, count):
        matches = [(n, line) for n, line in enumerate(r['text'].splitlines(), 1) if re.search(pattern, line)]
        if len(matches) != 1:
            raise ValueError('Expected one explicitly labeled table row.')
        n, line = matches[0]
        tail = re.sub(pattern, '', line)
        values = re.findall(r'(?<![A-Za-z])[+-]?\d[\d,]*(?:\.\d+)?|\(S\)|\(NA\)|\(\*\)', tail)
        if len(values) != count:
            raise ValueError('The numeric table row has an unexpected column count.')
        return n, line, values

    if table1:
        if 'millions of dollars' not in table1['text']:
            issues.append({'code': 'PUBLIC_REPORT_UNITS', 'message': 'Table 1 units missing.'})
        else:
            for metric, pattern in row_patterns:
                try:
                    n, line, values = tokens(table1, pattern, 12)
                    for column, offset in [(7, 0), (8, -1), (9, -2), (10, -12), (11, -13)]:
                        add(table1, metric, values[column], 'usd_million', 0,
                            f"{table1['locator']};line={n};numeric_column={column + 1}", line,
                            _shift(period, offset), method='published_report_table')
                except ValueError as exc:
                    issues.append({'code': 'PUBLIC_REPORT_TABLE', 'metric_id': 'census.' + metric, 'message': str(exc)})
    if table2:
        if 'shown as percents' not in table2['text']:
            issues.append({'code': 'PUBLIC_REPORT_UNITS', 'message': 'Table 2 units missing.'})
        else:
            rate_rows = [
                (r'^\s*total\s', [(0, 'month_on_month_pct'), (1, 'year_on_year_pct'), (5, 'rolling_three_month_year_on_year_pct')]),
                (r'^\s*Retail\s+[.·…]', [(0, 'retail_month_on_month_pct'), (1, 'retail_year_on_year_pct')]),
                (r'^\s*454\s+Nonstore retailers', [(1, 'nonstore_year_on_year_pct')]),
                (r'^\s*722\s+Food services & drinking places', [(1, 'food_services_year_on_year_pct')]),
            ]
            for pattern, metrics in rate_rows:
                try:
                    n, line, values = tokens(table2, pattern, 6)
                    for column, metric in metrics:
                        add(table2, metric, values[column], 'percent_points', 1,
                            f"{table2['locator']};line={n};numeric_column={column + 1}", line,
                            method='published_report_table')
                except ValueError as exc:
                    issues.append({'code': 'PUBLIC_REPORT_TABLE', 'message': str(exc)})
    grouped = {}
    for observation in observations:
        key = (observation['metric_id'], observation['period'])
        previous = grouped.setdefault(key, observation)
        if observation['value'] != previous['value'] or observation['status'] != previous['status']:
            issues.append({'code': 'PUBLIC_REPORT_CONTRADICTION', 'metric_id': key[0], 'period': key[1],
                           'locators': [previous['locator'], observation['locator']],
                           'message': 'Independent narrative/table observations disagree.'})
    headline = grouped.get(('census.sales_headline_usd_million', period))
    level = grouped.get(('census.sales_usd_million', period))
    if headline and level and headline['status'] == level['status'] == 'known':
        if Decimal(level['value']).quantize(Decimal('1E2'), rounding=ROUND_HALF_UP) != Decimal(headline['value']):
            issues.append({'code': 'PUBLIC_REPORT_CONTRADICTION', 'metric_id': headline['metric_id'], 'period': period,
                           'locators': [headline['locator'], level['locator']],
                           'message': 'Headline sales disagree with the independently printed level at headline precision.'})
    return _blocking_issues(result)


def inspect_report_bytes(data: bytes, *, asset_id='census-report') -> dict:
    return inspect_report({'id': asset_id, 'digest': hashlib.sha256(data).hexdigest(),
                           **inspect_asset(data, 'report.pdf')})
