"""Bounded native-format inspection. No macros, formula execution or network reads."""
from __future__ import annotations
import csv
from datetime import date, datetime, time, timedelta
from io import BytesIO, StringIO
from pathlib import PurePosixPath, Path
import unicodedata
from zipfile import ZipFile, BadZipFile
from xml.etree import ElementTree as ET
from .errors import DomainError

MAX_BYTES = 20 * 1024 * 1024
MAX_EXPANDED = 80 * 1024 * 1024
MAX_ROWS = 100_000
COLUMNS = ['transaction_id', 'date', 'region', 'status', 'amount', 'currency']
MIME = {'csv': 'text/csv', 'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', 'pdf': 'application/pdf',
        'png': 'image/png', 'jpg': 'image/jpeg', 'jpeg': 'image/jpeg'}


def _cell_observation(value):
    """Preserve native temporal values in JSON-safe structural observations."""
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return str(value)
    return value


def inspect_zip(data, expected_format=None):
    try:
        z = ZipFile(BytesIO(data))
        entries = z.infolist()
        if len(entries) > 2500 or sum(i.file_size for i in entries) > MAX_EXPANDED:
            raise DomainError('ARCHIVE_LIMIT', 'The Office archive exceeds the inspection limit.')
        for i in entries:
            path = PurePosixPath(i.filename)
            if path.is_absolute() or '..' in path.parts or '\\' in i.filename:
                raise DomainError('ARCHIVE_INVALID', 'Unsafe archive member path.')
            if i.flag_bits & 1 or (i.external_attr >> 16) & 0o170000 == 0o120000:
                raise DomainError('ARCHIVE_INVALID', 'Encrypted or symbolic-link archive members are not supported.')
            lower = i.filename.lower()
            if any(k in lower for k in ('vbaproject', 'activex', '/embeddings/', 'externallinks/')):
                raise DomainError('ACTIVE_CONTENT', 'Macros, embedded executable objects and external workbook links are not supported.')
            if lower.endswith('.rels'):
                root = ET.fromstring(z.read(i))
                for rel in root:
                    if rel.attrib.get('TargetMode') == 'External':
                        raise DomainError('EXTERNAL_RESOURCE', 'The document contains external relationships. Supply a self-contained copy.')
        required = {'[Content_Types].xml', '_rels/.rels'}
        if expected_format == 'xlsx':
            required.add('xl/workbook.xml')
        elif expected_format == 'docx':
            required.add('word/document.xml')
        if expected_format and not required.issubset(set(z.namelist())):
            raise DomainError('FORMAT_INVALID', 'The archive is missing required Office document parts.')
        return z
    except (BadZipFile, ET.ParseError) as exc:
        raise DomainError('FORMAT_INVALID', 'This is not a valid supported Office document.') from exc


def rows_from_csv(data):
    try:
        text = data.decode('utf-8-sig')
    except UnicodeDecodeError as exc:
        raise DomainError('INPUT_ENCODING', 'CSV inputs must use UTF-8 encoding.') from exc
    if '\x00' in text:
        raise DomainError('INPUT_ENCODING', 'CSV contains null bytes.')
    try:
        reader = csv.DictReader(StringIO(text), strict=True)
        columns = reader.fieldnames or []
        if not columns or len(set(columns)) != len(columns) or any(not x.strip() for x in columns):
            raise DomainError('INPUT_DRIFT', 'CSV headers must be nonempty and unique.')
        rows = []
        while True:
            start_line = reader.line_num + 1
            try:
                row = next(reader)
            except StopIteration:
                break
            if len(rows) >= MAX_ROWS:
                raise DomainError('INPUT_LIMIT', f'The local adapter supports up to {MAX_ROWS:,} records.')
            if None in row or any(v is None for v in row.values()):
                raise DomainError('INPUT_DRIFT', f'CSV row {reader.line_num} does not match its header.')
            row['_line'] = start_line
            if reader.line_num > start_line:
                row['_locator'] = f'csv:lines={start_line}-{reader.line_num}'
            rows.append(row)
        return columns, rows
    except csv.Error as exc:
        raise DomainError('INPUT_PARSE', f'CSV could not be parsed: {exc}') from exc


def xlsx_structure(data):
    from openpyxl import load_workbook
    inspect_zip(data, 'xlsx')
    wb = load_workbook(BytesIO(data), read_only=False, data_only=False, keep_links=False)
    cached = load_workbook(BytesIO(data), read_only=True, data_only=True, keep_links=False)
    sheets, candidates = [], []
    for ws in wb:
        if ws.max_row > MAX_ROWS + 1 or ws.max_column > 200:
            raise DomainError('INPUT_LIMIT', 'The worksheet exceeds the local row/column limit.')
        cells, rows, formula_count = [], [], 0
        headers = [str(c.value) if c.value is not None else '' for c in ws[1]]
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                if cell.data_type == 'f':
                    formula_count += 1
                if len(cells) < 1000:
                    cells.append({'address': cell.coordinate, 'value': _cell_observation(cell.value),
                                  'value_type': type(cell.value).__name__,
                                  'formula': cell.value if cell.data_type == 'f' else None,
                                  'cached_value': _cell_observation(cached[ws.title][cell.coordinate].value) if cell.data_type == 'f' else None,
                                  'number_format': cell.number_format})
        if set(COLUMNS) == set(headers) and len(headers) == len(COLUMNS) and not ws.merged_cells.ranges:
            for row_index, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                if all(v is None for v in row):
                    continue
                item = dict(zip(headers, row))
                for k, v in item.items():
                    if k == 'date' and isinstance(v, datetime) and (v.time() != time.min or v.tzinfo is not None):
                        raise DomainError('INPUT_DRIFT',
                            f'Transaction date at sheet={ws.title};row={row_index} contains a timestamp. '
                            'This program accepts local business dates only; timestamp conversion is not declared.')
                    if isinstance(v, (date, datetime)):
                        v = v.date().isoformat() if isinstance(v, datetime) else v.isoformat()
                    item[k] = '' if v is None else str(v)
                item['_locator'] = f'sheet={ws.title};row={row_index}'
                rows.append(item)
            if not formula_count:
                candidates.append((ws.title, rows))
        sheets.append({'name': ws.title, 'row_count': max(0, ws.max_row - 1), 'columns': headers,
                       'hidden': ws.sheet_state != 'visible', 'merged_cells': [str(r) for r in ws.merged_cells.ranges],
                       'formula_count': formula_count, 'cells_sample': cells,
                       'chart_count': len(ws._charts), 'pivot_count': len(ws._pivots),
                       'tables': list(ws.tables)})
    wb.close()
    cached.close()
    return sheets, candidates


def inspect_asset(data: bytes, filename: str, *, render_sink=None):
    if len(data) > MAX_BYTES or not data:
        raise DomainError('UPLOAD_LIMIT', 'Supply a nonempty file no larger than 20 MB.')
    filename = unicodedata.normalize('NFC', Path(filename.replace('\\', '/')).name)[:200]
    ext = filename.rsplit('.', 1)[-1].lower()
    if ext not in MIME:
        raise DomainError('FORMAT_UNSUPPORTED', 'Supported inputs are CSV, XLSX, DOCX, text-based PDF, PNG and JPEG.')
    profile = {'format': ext, 'warnings': [], 'eligible_roles': [], 'row_count': None, 'columns': [], 'regions': []}
    if ext in {'png', 'jpg', 'jpeg'}:
        from .images import normalize_image
        rendered, image_profile = normalize_image(data, ext)
        if render_sink is not None and render_sink(rendered) != image_profile['render_digest']:
            raise DomainError('IMAGE_INTEGRITY', 'The image store returned a different render identity.', 409)
        profile.update(image=image_profile, eligible_roles=['report_image'], parser='pillow/bounded-raster-v1')
        profile['warnings'].append('Display and export use a PNG with orientation applied and metadata removed. The original remains available as evidence; image content is not interpreted as computed facts.')
    elif ext == 'csv':
        columns, rows = rows_from_csv(data)
        profile.update(columns=columns, row_count=len(rows), sample_rows=rows[:8], parser='stdlib-csv/utf8-comma-v1')
        if set(columns) == set(COLUMNS) and len(columns) == len(COLUMNS):
            profile['eligible_roles'] = ['transactions']
        else:
            profile['warnings'].append('Catalogued as evidence. The revenue adapter requires exactly: ' + ', '.join(COLUMNS) + '.')
    elif ext == 'xlsx':
        sheets, candidates = xlsx_structure(data)
        profile.update(sheets=sheets, parser='openpyxl/3.1', row_count=sum(s['row_count'] for s in sheets))
        if len(candidates) == 1:
            name, rows = candidates[0]
            profile.update(eligible_roles=['transactions'], bound_sheet=name, columns=COLUMNS, sample_rows=rows[:8], row_count=len(rows))
        elif len(candidates) > 1:
            profile['warnings'].append('Several sheets match transactions. Bindings are ambiguous; export the intended sheet as CSV.')
        else:
            profile['warnings'].append('No unambiguous value-only transaction sheet. Formula results are retained as observations, never silently used as current values.')
    elif ext == 'docx':
        from docx import Document
        archive = inspect_zip(data, 'docx')
        doc = Document(BytesIO(data))
        regions = []
        body_order = {element: index for index, element in enumerate(doc.element.body)}
        for index, p in enumerate(doc.paragraphs):
            if p.text.strip():
                regions.append({'id': f'p{index}', 'kind': 'paragraph', 'text': p.text, 'style': p.style.name,
                                'locator': f'word/document.xml/paragraph[{index}]', 'order': body_order.get(p._p, index), 'status': 'uninvestigated'})
        for index, t in enumerate(doc.tables):
            regions.append({'id': f't{index}', 'kind': 'table', 'rows': [[c.text for c in r.cells] for r in t.rows],
                            'locator': f'word/document.xml/table[{index}]', 'order': body_order.get(t._tbl, index), 'status': 'uninvestigated'})
        regions.sort(key=lambda r: r['order'])
        for index, sec in enumerate(doc.sections):
            for kind in ('header', 'footer', 'first_page_header', 'first_page_footer', 'even_page_header', 'even_page_footer'):
                part = getattr(sec, kind)
                txt = '\n'.join(p.text for p in part.paragraphs if p.text)
                if txt:
                    regions.append({'id': f'{kind}{index}', 'kind': kind, 'text': txt, 'locator': f'section[{index}]/{kind}', 'status': 'uninvestigated'})
                for table_index, table in enumerate(part.tables):
                    regions.append({'id': f'{kind}{index}_table{table_index}', 'kind': kind,
                                    'rows': [[c.text for c in r.cells] for r in table.rows],
                                    'locator': f'section[{index}]/{kind}/table[{table_index}]', 'status': 'uninvestigated'})
        # Preserve difficult structures as explicit inventory items, not discarded extraction warnings.
        ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        for name in archive.namelist():
            if name.startswith('word/') and name.endswith('.xml') and any(x in name for x in ('document', 'header', 'footer', 'footnotes', 'endnotes')):
                root = ET.fromstring(archive.read(name))
                for index, drawing in enumerate(root.findall('.//w:drawing', ns) + root.findall('.//w:pict', ns)):
                    regions.append({'id': f'drawing_{len(regions)}', 'kind': 'image',
                                    'text': 'Embedded drawing or image; its content and placement require review.',
                                    'locator': f'{name}/drawing[{index}]', 'status': 'uninvestigated'})
                for index, box in enumerate(root.findall('.//w:txbxContent', ns)):
                    regions.append({'id': f'textbox_{len(regions)}', 'kind': 'textbox',
                                    'text': ' '.join(e.text or '' for e in box.findall('.//w:t', ns)),
                                    'locator': f'{name}/textbox[{index}]', 'status': 'uninvestigated'})
                if 'footnotes' in name or 'endnotes' in name:
                    for index, note in enumerate(root):
                        txt = ' '.join(e.text or '' for e in note.findall('.//w:t', ns)).strip()
                        if txt:
                            regions.append({'id': f'note_{len(regions)}', 'kind': 'footnote', 'text': txt,
                                            'locator': f'{name}/note[{index}]', 'status': 'uninvestigated'})
        if len(regions) > 500 or sum(len(str(r)) for r in regions) > 1_000_000:
            raise DomainError('DOCUMENT_LIMIT', 'The local document inspector supports at most 500 regions and 1 MB of structural text.')
        profile.update(regions=regions, parser='python-docx/native-v2', eligible_roles=['historical_target', 'commentary'])
        profile['warnings'].append('Located structure is preserved. Learning recognizes a bounded regional-report profile; other regions and exact layout require review.')
    else:
        from pypdf import PdfReader
        if not data.startswith(b'%PDF-'):
            raise DomainError('FORMAT_INVALID', 'The uploaded bytes are not a PDF.')
        try:
            pdf = PdfReader(BytesIO(data), strict=False)
            if pdf.is_encrypted:
                raise DomainError('PDF_ENCRYPTED', 'Encrypted PDF inputs are not supported.')
            if len(pdf.pages) > 100:
                raise DomainError('INPUT_LIMIT', 'The local inspector supports PDF inputs up to 100 pages.')
            for i, page in enumerate(pdf.pages):
                content = page.extract_text() or ''
                profile['regions'].append({'id': f'page{i+1}', 'kind': 'page', 'text': content,
                                           'locator': f'page={i+1}', 'status': 'uninvestigated'})
            profile.update(page_count=len(pdf.pages), parser='pypdf/text-v1', eligible_roles=['historical_target', 'commentary'])
            if not any(r['text'].strip() for r in profile['regions']):
                profile['eligible_roles'] = []
                profile['warnings'].append('No extractable text. Scanned PDFs require an OCR adapter and cannot be marked usable.')
            else:
                profile['warnings'].append('Page text retained with original PDF. Reading order, charts and table reconstruction require review.')
        except DomainError:
            raise
        except Exception as exc:
            raise DomainError('FORMAT_INVALID', 'The PDF could not be inspected.') from exc
    return {'filename': filename, 'media_type': MIME[ext], 'size': len(data), 'profile': profile,
            'status': 'usable' if profile['eligible_roles'] else 'blocked'}


def transaction_rows(data, profile):
    if 'transactions' not in profile.get('eligible_roles', []):
        raise DomainError('INPUT_DRIFT', 'This source cannot bind to the transaction role. Inspect its format findings.')
    if profile['format'] == 'csv':
        columns, rows = rows_from_csv(data)
        if set(columns) != set(COLUMNS) or len(columns) != len(COLUMNS):
            raise DomainError('INPUT_DRIFT', 'Transaction schema changed.')
        return rows
    if profile['format'] == 'xlsx':
        _, candidates = xlsx_structure(data)
        if len(candidates) != 1 or candidates[0][0] != profile['bound_sheet']:
            raise DomainError('INPUT_DRIFT', 'The approved transaction sheet no longer matches.')
        return candidates[0][1]
    raise DomainError('INPUT_DRIFT', 'Only CSV and value-only XLSX can supply transactions.')
