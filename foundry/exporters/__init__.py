"""Compile one frozen native snapshot into an explicitly scoped export.

No exporter reads original sources, runs models, executes customer code, or
resolves network resources. Facts and accepted prose come only from the
snapshot. Export planning rejects capabilities outside the tested core.
"""
from __future__ import annotations

import hashlib
import importlib.metadata
import json
import math
import re
import textwrap
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape
from zipfile import ZipFile
from typing import Callable

from .pivot import add_flat_pivot
from .images import resolve_images, set_picture_properties, annotate_workbook_images

VERSION = "1.1.0"
MEDIA_TYPES = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}
FAMILIES = {"docx": "flow", "pdf": "flow", "xlsx": "grid", "pptx": "canvas"}
INK, GREEN, MUTED, PALE = "203A36", "397B66", "65726E", "EEF3EF"


def _period_caption(snapshot):
    label = snapshot["period"]["label"]
    return label + ("   DRAFT / Awaiting review" if snapshot.get("status") == "review_required" else "")


class ExportError(ValueError):
    def __init__(self, message: str, code="export_blocked", status_code=422, findings=None):
        super().__init__(message)
        self.message, self.code, self.status_code = message, code, status_code
        self.findings = findings or [{"severity": "block", "phase": "export_plan", "code": code, "message": message}]


def _digest(value) -> str:
    return hashlib.sha256(json.dumps(value, sort_keys=True, separators=(",", ":"), ensure_ascii=False).encode()).hexdigest()


def _number(value, *, approximate=False) -> float:
    if value is None:
        raise ExportError("An undefined value cannot become an export number.", "undefined_numeric_value")
    try:
        exact = Decimal(str(value))
        converted = float(exact)
        if not exact.is_finite() or not math.isfinite(converted) or (not approximate and Decimal(str(converted)) != exact):
            raise ValueError
        return converted
    except (ValueError, InvalidOperation, OverflowError):
        raise ExportError("A value exceeds the supported native numeric precision.", "numeric_precision") from None


def _text(node, facts) -> str:
    result = []
    for run in node.get("runs", []):
        if run.get("type") == "text":
            result.append(run["text"])
        elif run.get("type") == "fact":
            fact = facts.get(run.get("fact_id"))
            if fact is None:
                raise ExportError("A text node references a missing fact.", "missing_fact")
            result.append(str(fact["display"]))
        else:
            raise ExportError("The text contains an unsupported run.", "unsupported_rich_text")
    return "".join(result)


def _display(value, column) -> str:
    if value is None:
        return "Undefined"
    if column.get("type") in {"decimal", "integer"}:
        amount = Decimal(str(value))
        if column.get("unit") in {"ratio", "percent", "%"} or column["id"] == "growth":
            rounded = (amount * 100).quantize(Decimal(".1"), rounding=ROUND_HALF_UP)
            return f"{rounded:+.1f}%"
        if column.get("unit") == "EUR":
            return f"{amount.quantize(Decimal('.01'), rounding=ROUND_HALF_UP):,.2f}"
        return format(amount, "f")
    return str(value)


def _table(snapshot, node, pivot=False):
    identifier = node["materialized_dataset_id"] if pivot else node["dataset_id"]
    data = snapshot["datasets"][identifier]
    columns = data["columns"]
    if pivot:
        columns = [next(c for c in columns if c["id"] == key) for key in ["region", "comparison", "current"]]
    headers = [c["label"] + (" (EUR)" if c.get("unit") == "EUR" and "EUR" not in c["label"] else "") for c in columns]
    rows = [[_display(row[c["id"]], c) for c in columns] for row in data["rows"]]
    return headers, rows, columns, data


def _chart_data(snapshot, node):
    data = snapshot["datasets"][node["dataset_id"]]
    if node.get("chart_type") != "bar":
        raise ExportError("Only the declared grouped bar chart is supported.", "unsupported_chart")
    if not 1 <= len(data["rows"]) <= 20:
        raise ExportError("This chart profile supports one to twenty categories.", "chart_size_limit")
    categories = [str(row[node["category_column"]]) for row in data["rows"]]
    if any(len(label) > 60 for label in categories):
        raise ExportError("A chart label exceeds this layout profile.", "chart_label_limit")
    series = [(s["label"], [_number(row[s["column_id"]]) for row in data["rows"]]) for s in node["series"]]
    return categories, series


def _chart_png(snapshot, node) -> bytes:
    """Static bar representation for document media; includes values and units."""
    from PIL import Image, ImageDraw, ImageFont
    categories, series = _chart_data(snapshot, node)
    for text in categories + [node["title"], node["axis_unit"]] + [name for name, _ in series]:
        try:
            text.encode("cp1252")
        except UnicodeEncodeError:
            raise ExportError("The static chart font profile does not cover these labels. A font-capable renderer is required.", "unsupported_font_glyph") from None
    width, height = 1200, max(500, len(categories) * 90 + 150)
    font = ImageFont.load_default(size=24)
    small = ImageFont.load_default(size=21)
    left, right, top, bottom = 210, 1070, 75, height - 80
    # Preserve full series labels and size the legend from actual font metrics.
    # Fixed column offsets overlap when the period has a descriptive label.
    legend, legend_x, legend_y, row_height = [], left, 0, 0
    for name, _ in series:
        available = width - left - 60
        for wrap_width in range(min(80, max(1, len(name))), 0, -1):
            lines = textwrap.wrap(name, width=wrap_width, break_long_words=True,
                                  break_on_hyphens=False) or [""]
            text_width = max(small.getlength(line) for line in lines)
            if text_width <= available:
                break
        if legend_x > left and legend_x + 27 + text_width > width - 24:
            legend_x, legend_y = left, legend_y + row_height + 10
            row_height = 0
        legend.append((legend_x, legend_y, lines))
        row_height = max(row_height, len(lines) * 26)
        legend_x += 27 + text_width + 32
    legend_top = bottom + 44
    height = max(height, legend_top + legend_y + row_height + 10)
    result = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(result)
    values = [v for _, row in series for v in row]
    low, high = min(0, min(values)), max(0, max(values))
    span = high - low or 1
    project = lambda value: left + (value - low) / span * (right - left - 80)
    baseline = project(0)
    draw.line((baseline, top, baseline, bottom), fill="#A6B4AE", width=2)
    for i in range(5):
        v = low + span * i / 4
        x = project(v)
        draw.line((x, top, x, bottom), fill="#E4EAE6", width=1)
        draw.text((x - 10, bottom + 12), f"{v:,.0f}", fill=f"#{MUTED}", font=small)
    palette = [f"#{GREEN}", "#A3B9AF", "#738AA1"]
    row_h = (bottom - top) / len(categories)
    bar_h = min(24, row_h / (len(series) + 1))
    for i, category in enumerate(categories):
        cy = top + row_h * (i + .5)
        label = "\n".join(textwrap.wrap(category, width=17, break_long_words=True))
        draw.multiline_text((8, cy - 12 * len(label.splitlines())), label, fill=f"#{INK}", font=small, spacing=2)
        for j, (_, vals) in enumerate(series):
            y = cy + (j - len(series) / 2) * bar_h
            x = project(vals[i])
            draw.rectangle((min(baseline, x), y, max(baseline, x), y + bar_h - 3), fill=palette[j % len(palette)])
            label_x = x + 8 if vals[i] >= 0 else max(0, x - 100)
            draw.text((label_x, y - 2), f"{vals[i]:,.2f}", fill=f"#{INK}", font=small)
    draw.text((8, 14), f"{node['title']} ({node['axis_unit']})", fill=f"#{INK}", font=font)
    for j, (x, y, lines) in enumerate(legend):
        y += legend_top
        draw.rectangle((x, y + 3, x + 18, y + 21), fill=palette[j % len(palette)])
        for line_index, line in enumerate(lines):
            draw.text((x + 27, y + line_index * 26), line, fill=f"#{MUTED}", font=small)
    buffer = BytesIO()
    result.save(buffer, format="PNG")
    return buffer.getvalue()


def _plan(snapshot, fmt, policy):
    if fmt not in MEDIA_TYPES:
        raise ExportError("Choose DOCX, XLSX, PDF, or PPTX.", "unsupported_format")
    if policy not in {"compatible", "strict", "static"}:
        raise ExportError("Unknown fidelity policy.", "unsupported_policy")
    if snapshot.get("status") == "blocked" or any(f.get("severity") == "block" for f in snapshot.get("findings", [])):
        raise ExportError("Resolve blocking findings before exporting this snapshot.", "snapshot_blocked")
    family = FAMILIES[fmt]
    views = snapshot["views"].values() if isinstance(snapshot["views"], dict) else snapshot["views"]
    view = next((v for v in views if v["family"] == family), None)
    if not view:
        raise ExportError("This snapshot has no presentation for the requested format.", "missing_view")
    nodes = snapshot["nodes"]
    if isinstance(nodes, list):
        nodes = {n["id"]: n for n in nodes}
    identifiers = view["node_ids"]
    if len(identifiers) != len(set(identifiers)):
        raise ExportError("The view contains duplicate component IDs.", "duplicate_coverage")
    if set(identifiers) - set(nodes):
        raise ExportError("The view references a missing component.", "missing_component")
    required = set(view["coverage"]["required_node_ids"])
    if not required.issubset(identifiers):
        raise ExportError("The view omits required content.", "missing_coverage")
    leaves = {n["id"] for n in nodes.values() if n["kind"] != "section"}
    if view["coverage"]["scope"] == "complete" and set(identifiers) != leaves:
        raise ExportError("A complete view must represent every content component.", "missing_coverage")
    selected = [nodes[i] for i in identifiers]
    for node in selected:
        if node["kind"] not in {"rich_text", "table", "chart", "pivot", "image"}:
            raise ExportError(f"The {node['kind']} component {node['id']} has no certified asset adapter.", "unsupported_component")
        if node["kind"] == "rich_text":
            _text(node, snapshot["facts"])
        if node["kind"] in {"table", "chart", "pivot"}:
            if node.get("dataset_id") not in snapshot["datasets"]:
                raise ExportError("A component references a missing dataset.", "missing_dataset")
            data = snapshot["datasets"][node["dataset_id"]]
            if len(data["rows"]) > 1000:
                raise ExportError("This export profile supports at most 1,000 aggregate rows.", "export_size_limit")
        if node["kind"] == "chart":
            _chart_data(snapshot, node)
        if node["kind"] == "pivot":
            if (node["row_dimensions"] != ["region"] or node["column_dimensions"] != ["period"]
                    or node["measures"] != [{"column_id": "revenue", "aggregation": "sum"}]):
                raise ExportError("Only a region by period sum pivot is implemented.", "unsupported_pivot")
            source_rows = snapshot["datasets"][node["dataset_id"]]["rows"]
            expected = {}
            for record in source_rows:
                if set(record) != {"region", "period", "revenue"} or record["period"] not in {"current", "comparison"}:
                    raise ExportError("The pivot must use only approved region, period and revenue aggregates.", "unapproved_embedded_data")
                key = (record["region"], record["period"])
                expected[key] = expected.get(key, Decimal(0)) + Decimal(record["revenue"])
            materialized = snapshot["datasets"][node["materialized_dataset_id"]]["rows"]
            actual = {(row["region"], period): Decimal(row[period]) for row in materialized for period in ["current", "comparison"]}
            if expected != actual:
                raise ExportError("The pivot cache does not reconcile to its materialized result.", "pivot_reconciliation_failed")
            if fmt == "xlsx" and policy == "strict" and "grid" in node.get("native_required_in", []):
                raise ExportError("Native XLSX pivot structure is implemented, but Excel interaction certification is pending. Choose the compatible preview or explicitly permit a static workbook.", "native_certification_pending")
    warnings = []
    if fmt == "pdf":
        warnings.append({"code": "native_pdf_font_profile", "message": "The native PDF uses its Helvetica flow profile. It is independent of Word pagination and font rendering."})
        if any(node["kind"] == "image" for node in selected):
            warnings.append({"code": "pdf_image_tagging_unavailable", "message": "Report image descriptions are preserved in captions and the manifest. This PDF does not provide tagged image accessibility."})
    if fmt == "xlsx" and any(n["kind"] == "pivot" for n in selected):
        warnings.append({"code": "static_pivot_approved" if policy == "static" else "native_certification_pending",
                         "message": "The chosen policy permits a static pivot result." if policy == "static" else
                         "Native pivot definition and caches are included. Interaction in Microsoft Excel has not been certified."})
    if snapshot.get("status") == "review_required":
        warnings.append({"code": "draft_snapshot", "message": "This exported revision still requires content review."})
    return view, selected, warnings


def _location(node, representation, location, **extra):
    return {"node_id": node["id"], "representation": representation, "location": location,
            "semantic_coverage": "complete", **extra}


def _docx(snapshot, nodes, path, policy, view, images):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor, Mm
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    # Bundled default templates may carry a themed title border. The flow
    # profile is explicit, so remove inherited paragraph decoration as well.
    for paragraph_style in doc.styles:
        for border in list(paragraph_style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    sec = doc.sections[0]
    recipe = view.get("recipe", {})
    margin = float(recipe.get("margin_mm", 18))
    indent = float(recipe.get("paragraph_indent_pt", 18))
    if not 10 <= margin <= 30 or not 0 <= indent <= 72:
        raise ExportError("The requested flow spacing is outside the tested layout profile.", "unsupported_flow_layout")
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Mm(margin)
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    style = doc.styles["Normal"]
    style.font.name, style.font.size = recipe.get("font", "Calibri"), Pt(float(recipe.get("body_pt", 10)))
    style.paragraph_format.space_after = Pt(7)
    style.paragraph_format.line_spacing = 1.15
    for name in ["Title", "Heading 1", "Heading 2"]:
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
        doc.styles[name].font.name = "Calibri"
    for name in ["Subtitle", "Caption"]:
        doc.styles[name].font.color.rgb = RGBColor.from_string(MUTED)
        doc.styles[name].font.bold = False
    doc.styles["Caption"].font.size = Pt(8)
    doc.core_properties.title = snapshot["title"]
    doc.core_properties.subject = f"Frozen native report {snapshot['id']} revision {snapshot['revision']}"
    doc.add_heading(snapshot["title"], 0)
    doc.add_paragraph(_period_caption(snapshot), "Subtitle")
    mapping, validations = [], []
    for index, node in enumerate(nodes):
        paragraph = doc.add_paragraph()
        bookmark = OxmlElement("w:bookmarkStart")
        bookmark.set(qn("w:id"), str(index + 1))
        bookmark.set(qn("w:name"), "rf_" + re.sub(r"[^A-Za-z0-9_]", "_", node["id"])[:35])
        paragraph._p.append(bookmark)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(index + 1))
        paragraph._p.append(end)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.line_spacing = Pt(1)
        paragraph.paragraph_format.keep_with_next = True
        if node["kind"] == "rich_text":
            if node["id"] != "summary":
                doc.add_heading(node["title"], 2)
            p = doc.add_paragraph(_text(node, snapshot["facts"]))
            # Physical paragraph formatting, never leading spaces.
            p.paragraph_format.left_indent = Pt(indent)
            p.paragraph_format.first_line_indent = Inches(0)
            mapping.append(_location(node, "native_text", f"bookmark:rf_{node['id']}", editable=bool(node.get("editable"))))
        elif node["kind"] in {"table", "pivot"}:
            doc.add_heading(node["title"], 2)
            headers, rows, columns, _ = _table(snapshot, node, node["kind"] == "pivot")
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = "Table Grid"
            table.autofit = False
            borders = OxmlElement("w:tblBorders")
            for side in ["top", "left", "bottom", "right", "insideH", "insideV"]:
                border = OxmlElement("w:"+side)
                for key, value in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
                    border.set(qn("w:"+key), value)
                borders.append(border)
            table._tbl.tblPr.append(borders)
            for cell, label in zip(table.rows[0].cells, headers):
                cell.text = label
                fill = OxmlElement("w:shd")
                fill.set(qn("w:fill"), INK)
                cell._tc.get_or_add_tcPr().append(fill)
                for run in cell.paragraphs[0].runs:
                    run.font.bold, run.font.color.rgb, run.font.size = True, RGBColor(255, 255, 255), Pt(9)
            repeat = OxmlElement("w:tblHeader")
            table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for rindex, row in enumerate(rows):
                table_row = table.add_row()
                cant_split = OxmlElement("w:cantSplit")
                table_row._tr.get_or_add_trPr().append(cant_split)
                cells = table_row.cells
                for cindex, (cell, text) in enumerate(zip(cells, row)):
                    cell.text = text
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after = Pt(4)
                        p.paragraph_format.space_before = Pt(4)
                        for run in p.runs:
                            run.font.size = Pt(9)
                    if rindex % 2 == 0:
                        shade = OxmlElement("w:shd")
                        shade.set(qn("w:fill"), PALE)
                        cell._tc.get_or_add_tcPr().append(shade)
            rep = "equivalent_static_table" if node["kind"] == "pivot" else "native_table"
            mapping.append(_location(node, rep, f"bookmark:rf_{node['id']}", editable=True,
                                     native_behavior="static result" if node["kind"] == "pivot" else "table"))
            if node["kind"] == "pivot":
                doc.add_paragraph("Static pivot result. Interactive pivot behavior is available only through the workbook capability profile.", "Caption")
        elif node["kind"] == "chart":
            doc.add_heading(node["title"], 2)
            doc.add_picture(BytesIO(_chart_png(snapshot, node)), width=Inches(6.6))
            doc.add_paragraph("Chart uses the accepted report values. The document contains a static image.", "Caption")
            mapping.append(_location(node, "raster_fallback", f"bookmark:rf_{node['id']}", editable=False,
                                     native_behavior="static chart", fallback_approval="flow view policy"))
        elif node["kind"] == "image":
            bound = images[node["id"]]
            doc.add_heading(node["title"], 2)
            width, height = bound.fit(float(sec.page_width-sec.left_margin-sec.right_margin)/914400, 6.2,
                                      units_per_pixel=1/96)
            picture = doc.add_picture(BytesIO(bound.data), width=Inches(width), height=Inches(height))
            set_picture_properties(picture._inline.docPr, node)
            set_picture_properties(picture._inline.xpath(".//pic:cNvPr")[0], node)
            if node["alt_text"] and not node.get("decorative", False):
                doc.paragraphs[-1].paragraph_format.keep_with_next = True
                doc.add_paragraph(node["alt_text"], "Caption")
            mapping.append(_location(node, "native_image", f"bookmark:rf_{node['id']}",
                                     editable=False, placement_editable=True, alt_text_support="drawing_properties",
                                     rendered_size={"width": width, "height": height, "unit": "in"}, **bound.manifest(node)))
    footer = sec.footer.paragraphs[0]
    footer.text = f"{_period_caption(snapshot)}   Revision {snapshot['revision']}"
    footer.style = doc.styles["Caption"]
    doc.save(path)
    readback = Document(path)
    actual = "\n".join(p.text for p in readback.paragraphs)
    for node in nodes:
        if node["kind"] == "rich_text" and _text(node, snapshot["facts"]) not in actual:
            raise ExportError("DOCX text verification failed.", "render_content_mismatch")
    with ZipFile(path) as z:
        document = ET.fromstring(z.read("word/document.xml"))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        if any(n["kind"] in {"table", "pivot"} for n in nodes) and not document.findall(".//w:tblHeader", ns):
            raise ExportError("DOCX repeated table headers are missing.", "render_structure_mismatch")
        indents = document.findall(".//w:ind", ns)
        if any(n["kind"] == "rich_text" for n in nodes) and not any(i.attrib.get(f"{{{ns['w']}}}left") == str(round(indent*20)) for i in indents):
            raise ExportError("DOCX precision indentation was not retained.", "render_structure_mismatch")
    validations.extend([{"check": "accepted_text_roundtrip", "status": "pass"},
                        {"check": "repeated_table_headers", "status": "pass"},
                        {"check": "declared_paragraph_left_indent", "status": "pass", "points": indent},
                        {"check": "target_application_pagination", "status": "not_certified"}])
    return mapping, validations, "python-docx", {"font": style.font.name, "font_substitution_certification": "pending", "paragraph_indent_pt": indent, "margin_mm": margin}


def _xlsx(snapshot, nodes, path, policy, view, images):
    import xlsxwriter
    from openpyxl import load_workbook
    workbook = xlsxwriter.Workbook(path, {"strings_to_urls": False, "strings_to_formulas": False})
    workbook.set_properties({"title": snapshot["title"], "subject": f"Snapshot {snapshot['id']} revision {snapshot['revision']}"})
    overview = workbook.add_worksheet("Overview")
    analysis = workbook.add_worksheet("Analysis")
    source = workbook.add_worksheet("Approved data")
    title = workbook.add_format({"font_name": "Calibri", "font_size": 22, "bold": True, "font_color": INK})
    subtitle = workbook.add_format({"font_size": 11, "font_color": MUTED})
    body = workbook.add_format({"font_size": 11, "text_wrap": True, "valign": "vcenter", "font_color": INK})
    head = workbook.add_format({"bold": True, "bg_color": INK, "font_color": "FFFFFF", "text_wrap": True, "border": 1, "border_color": "D9D9D9", "valign": "vcenter"})
    number = workbook.add_format({"num_format": '#,##0.00;[Red](#,##0.00)', "border": 1, "border_color": "D9D9D9"})
    ratio = workbook.add_format({"num_format": '+0.0%;[Red]-0.0%;+0.0%', "border": 1, "border_color": "D9D9D9"})
    cell = workbook.add_format({"border": 1, "border_color": "D9D9D9"})
    for sheet in [overview, analysis, source]:
        sheet.hide_gridlines(2)
        sheet.set_column("A:A", 23)
        sheet.set_column("B:E", 19)
        sheet.set_default_row(22)
        sheet.set_landscape()
        sheet.fit_to_pages(1, 0)
        sheet.set_paper(9)
        sheet.set_margins(.3, .3, .5, .5)
    overview.merge_range("A1:E2", snapshot["title"], title)
    overview.merge_range("A3:E3", _period_caption(snapshot), subtitle)
    overview.freeze_panes(4, 1)
    analysis.merge_range("A1:H1", "Regional analysis", title)
    analysis.set_column("B:I", 19)
    source.freeze_panes(1, 0)
    mapping, row_cursor, pivot_meta = [], 4, None
    analysis_last_row = 20
    table_reference = None
    pivot_node = next((n for n in nodes if n["kind"] == "pivot"), None)
    chart_node = next((n for n in nodes if n["kind"] == "chart"), None)
    for node in nodes:
        if node["kind"] == "rich_text":
            text = _text(node, snapshot["facts"])
            height = max(2, math.ceil(len(text) / 92))
            overview.merge_range(row_cursor, 0, row_cursor + height - 1, 4, text, body)
            mapping.append(_location(node, "native_text", f"Overview!A{row_cursor+1}:E{row_cursor+height}", editable=True))
            row_cursor += height + 1
        elif node["kind"] == "table":
            headers, rows, columns, data = _table(snapshot, node)
            overview.merge_range(row_cursor, 0, row_cursor, len(columns)-1, node["title"], head)
            row_cursor += 1
            start = row_cursor
            for c, label in enumerate(headers):
                overview.write_string(row_cursor, c, label, head)
            overview.set_row(row_cursor, 34)
            for offset, row in enumerate(data["rows"], 1):
                for c, column in enumerate(columns):
                    value = row[column["id"]]
                    fmt = ratio if column["id"] == "growth" else number
                    if column["type"] in {"decimal", "integer"} and value is not None:
                        overview.write_number(row_cursor + offset, c, _number(value, approximate=column["id"] == "growth"), fmt)
                    else:
                        overview.write_string(row_cursor + offset, c, _display(value, column), cell)
            overview.autofilter(row_cursor, 0, row_cursor + len(rows), len(columns)-1)
            overview.repeat_rows(row_cursor)
            table_reference = (start, data, columns)
            mapping.append(_location(node, "native_table", f"Overview!A{start+1}:E{start+len(rows)+1}", editable=True))
            row_cursor += len(rows) + 3
    source_rows = snapshot["datasets"][pivot_node["dataset_id"]]["rows"] if pivot_node else []
    for c, name in enumerate(["region", "period", "revenue"]):
        source.write_string(0, c, name, head)
    for r, record in enumerate(source_rows, 1):
        if set(record) != {"region", "period", "revenue"}:
            workbook.close()
            raise ExportError("The workbook cannot embed undeclared pivot source columns.", "unapproved_embedded_data")
        source.write_string(r, 0, str(record["region"]), cell)
        source.write_string(r, 1, str(record["period"]), cell)
        source.write_number(r, 2, _number(record["revenue"]), number)
    source.write_string(len(source_rows)+3, 0, "Approved regional aggregates only. No transaction identifiers are embedded.", subtitle)
    if pivot_node:
        _, pivot_rows, _, data = _table(snapshot, pivot_node, True)
        analysis.write_string(2, 0, "Revenue (EUR)", subtitle)
        analysis.write_string(2, 1, "Period", subtitle)
        for c, label in enumerate(["Region", "comparison", "current"]):
            analysis.write_string(3, c, label, head)
        ordered = sorted(data["rows"], key=lambda r: str(r["region"]).casefold())
        for r, record in enumerate(ordered, 4):
            analysis.write_string(r, 0, record["region"], cell)
            analysis.write_number(r, 1, _number(record["comparison"]), number)
            analysis.write_number(r, 2, _number(record["current"]), number)
        warning_row = max(len(ordered)+5, 17)
        analysis_last_row = warning_row+3
        analysis.merge_range(warning_row, 0, warning_row+2, 8,
                             "Static pivot result approved by the selected policy." if policy == "static" else
                             "Native pivot structure and approved aggregate cache included. Microsoft Excel interaction certification is pending.", body)
        mapping.append(_location(pivot_node, "equivalent_static_table" if policy == "static" else "native_pivot",
                                 f"Analysis!A3:C{len(ordered)+4}", editable=True,
                                 target_certification="not_applicable" if policy == "static" else "pending"))
    if chart_node:
        if not table_reference:
            workbook.close()
            raise ExportError("The workbook chart needs its declared materialized table.", "chart_source_missing")
        start, data, columns = table_reference
        if chart_node["dataset_id"] != data["id"]:
            workbook.close()
            raise ExportError("The chart must reference the approved regional table.", "chart_source_mismatch")
        category_index = next(i for i, c in enumerate(columns) if c["id"] == chart_node["category_column"])
        chart = workbook.add_chart({"type": "bar"})
        for index, series in enumerate(chart_node["series"]):
            ci = next(i for i, c in enumerate(columns) if c["id"] == series["column_id"])
            chart.add_series({"name": series["label"], "categories": ["Overview", start+1, category_index, start+len(data["rows"]), category_index],
                              "values": ["Overview", start+1, ci, start+len(data["rows"]), ci],
                              "fill": {"color": GREEN if index == 0 else "A3B9AF"}, "border": {"none": True}})
        chart.set_title({"name": chart_node["title"]})
        chart.set_x_axis({"name": chart_node["axis_unit"], "num_format": '#,##0'})
        chart.set_y_axis({"reverse": True})
        chart.set_legend({"position": "bottom"})
        analysis.insert_chart(2, 4, chart, {"x_scale": 1.1, "y_scale": 1.18})
        mapping.append(_location(chart_node, "native_chart", "Analysis!E3", editable=True))
    overview.print_area(0, 0, row_cursor, 4)
    analysis.print_area(0, 0, analysis_last_row, 8)
    image_nodes = [node for node in nodes if node["kind"] == "image"]
    if image_nodes:
        image_sheet = workbook.add_worksheet("Images")
        image_sheet.hide_gridlines(2)
        image_sheet.set_column("A:J", 12)
        image_sheet.set_default_row(20)
        image_sheet.set_landscape()
        image_sheet.fit_to_pages(1, 0)
        image_sheet.set_paper(9)
        image_sheet.set_margins(.3, .3, .5, .5)
        image_sheet.set_header("&L" + _period_caption(snapshot))
        image_row, page_breaks = 0, []
        for node in image_nodes:
            bound = images[node["id"]]
            image_sheet.merge_range(image_row, 0, image_row+1, 9, node["title"], title)
            image_sheet.merge_range(image_row+2, 0, image_row+2, 9, _period_caption(snapshot), subtitle)
            width, height = bound.fit(800, 480, units_per_pixel=1)
            image_sheet.insert_image(image_row+4, 0, BytesIO(bound.data), {
                "x_scale": width/bound.width_px, "y_scale": height/bound.height_px,
                "description": node["alt_text"], "decorative": node.get("decorative", False), "object_position": 1})
            bottom_row = image_row+4+math.ceil(height/(20*96/72))
            if node["alt_text"] and not node.get("decorative", False):
                caption_rows = max(2, math.ceil(len(node["alt_text"])/110))
                image_sheet.merge_range(bottom_row+1, 0, bottom_row+caption_rows, 9, node["alt_text"], body)
                bottom_row += caption_rows
            mapping.append(_location(node, "native_image", f"Images!A{image_row+5}",
                                     editable=False, placement_editable=True, alt_text_support="drawing_properties",
                                     rendered_size={"width": width, "height": height, "unit": "px"}, **bound.manifest(node)))
            image_row = bottom_row+4
            page_breaks.append(image_row)
        image_sheet.print_area(0, 0, image_row-2, 9)
        if len(page_breaks) > 1:
            image_sheet.set_h_pagebreaks(page_breaks[:-1])
    workbook.close()
    annotate_workbook_images(path, image_nodes)
    if pivot_node and policy != "static":
        pivot_meta = add_flat_pivot(path, source_rows)
    check = load_workbook(path, data_only=False)
    if pivot_node and policy != "static":
        pivots = check["Analysis"]._pivots
        if len(pivots) != 1 or pivots[0].cache.recordCount != len(source_rows):
            raise ExportError("Independent pivot structure verification failed.", "pivot_verification_failed")
    if any(cell.data_type == "f" for ws in check for row in ws for cell in row):
        raise ExportError("Materialized workbook unexpectedly contains formulas.", "unexpected_formula")
    check.close()
    validations = [{"check": "openpyxl_independent_parse", "status": "pass"},
                   {"check": "materialized_values_no_formulas", "status": "pass"},
                   {"check": "approved_aggregate_cache_columns", "status": "pass"},
                   {"check": "target_application_interaction", "status": "not_certified"}]
    if pivot_meta:
        validations.append({"check": "pivot_definition_cache_records", "status": "pass"})
    return mapping, validations, "XlsxWriter", {"pivot": pivot_meta, "numeric_policy": "Currency retains source decimal display precision. Ratios use native IEEE numeric storage and one-decimal percentage display. Exact decimal authority remains the snapshot."}


def _pdf(snapshot, nodes, path, policy, view, images):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, KeepTogether
    from pypdf import PdfReader, PdfWriter
    # Helvetica's standard PDF encoding covers the declared Latin profile.
    # Reject unsupported text instead of emitting unmarked missing glyphs.
    texts = [snapshot["title"], snapshot["period"]["label"]]
    for node in nodes:
        texts.append(node["title"])
        if node["kind"] == "rich_text":
            texts.append(_text(node, snapshot["facts"]))
        elif node["kind"] in {"table", "pivot"}:
            headers, rows, _, _ = _table(snapshot, node, node["kind"] == "pivot")
            texts.extend(headers + [value for row in rows for value in row])
        elif node["kind"] == "image" and not node.get("decorative", False):
            texts.append(node["alt_text"])
    try:
        "".join(texts).encode("cp1252")
    except UnicodeEncodeError:
        raise ExportError("The native PDF font profile does not cover this text. Export DOCX or install a wider font profile.", "unsupported_font_glyph") from None
    margin = float(view.get("recipe", {}).get("margin_mm", 18)) * 72 / 25.4
    indent = float(view.get("recipe", {}).get("paragraph_indent_pt", 18))
    content_width = A4[0] - margin * 2
    doc = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=margin, leftMargin=margin, topMargin=margin, bottomMargin=margin,
                            title=snapshot["title"], author="Report Foundry")
    styles = {
        "title": ParagraphStyle("title", fontName="Helvetica-Bold", fontSize=23, leading=27, textColor=colors.HexColor("#"+INK), spaceAfter=10),
        "body": ParagraphStyle("body", fontName="Helvetica", fontSize=10, leading=14, leftIndent=indent, spaceAfter=10),
        "head": ParagraphStyle("head", fontName="Helvetica-Bold", fontSize=12, leading=15, spaceBefore=10, spaceAfter=7, keepWithNext=True),
        "small": ParagraphStyle("small", fontName="Helvetica", fontSize=8, leading=11, textColor=colors.HexColor("#"+MUTED), spaceAfter=8),
        "cell": ParagraphStyle("cell", fontName="Helvetica", fontSize=8, leading=11),
        "thead": ParagraphStyle("thead", fontName="Helvetica-Bold", fontSize=8, leading=11, textColor=colors.white),
    }
    node_pages = {}
    class TrackedParagraph(Paragraph):
        def split(self, available_width, available_height):
            pieces = super().split(available_width, available_height)
            if pieces and hasattr(self, "foundry_node_id"):
                pieces[0].foundry_node_id = self.foundry_node_id
            return pieces
    def tracked(text, style, identifier):
        p = TrackedParagraph(text, style)
        p.foundry_node_id = identifier
        return p
    def after_flowable(flowable):
        identifier = getattr(flowable, "foundry_node_id", None)
        if identifier is not None:
            node_pages.setdefault(identifier, doc.page-1)
    doc.afterFlowable = after_flowable
    story = [Paragraph(escape(snapshot["title"]), styles["title"]), Paragraph(escape(_period_caption(snapshot)), styles["small"])]
    mapping = []
    for node in nodes:
        anchor = re.sub(r"[^A-Za-z0-9_]", "_", node["id"])
        if node["kind"] == "rich_text":
            story.append(tracked(escape(_text(node, snapshot["facts"])), styles["body"], anchor))
            mapping.append(_location(node, "static_text", f"named_destination:{anchor}", editable=False))
        elif node["kind"] in {"table", "pivot"}:
            story.append(tracked(escape(node["title"]), styles["head"], anchor))
            headers, rows, _, _ = _table(snapshot, node, node["kind"] == "pivot")
            table = Table([[Paragraph(escape(value), styles["thead"]) for value in headers]] +
                          [[Paragraph(escape(value), styles["cell"]) for value in row] for row in rows],
                          colWidths=[content_width/len(headers)]*len(headers), repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([("BACKGROUND", (0,0),(-1,0), colors.HexColor("#"+INK)),
                                       ("ROWBACKGROUNDS",(0,1),(-1,-1), [colors.HexColor("#"+PALE), colors.white]),
                                       ("GRID",(0,0),(-1,-1),.4,colors.HexColor("#D9D9D9")),
                                       ("VALIGN",(0,0),(-1,-1),"MIDDLE"), ("TOPPADDING",(0,0),(-1,-1),7),
                                       ("BOTTOMPADDING",(0,0),(-1,-1),7)]))
            story.extend([table, Spacer(1, 8)])
            if node["kind"] == "pivot":
                story.append(Paragraph("Static pivot result under the flow view policy.", styles["small"]))
            mapping.append(_location(node, "equivalent_static_table" if node["kind"] == "pivot" else "static_table",
                                     f"named_destination:{anchor}", editable=False))
        elif node["kind"] == "chart":
            picture = Image(BytesIO(_chart_png(snapshot, node)))
            ratio = content_width / picture.imageWidth
            picture.drawWidth, picture.drawHeight = content_width, picture.imageHeight * ratio
            story.append(KeepTogether([tracked(escape(node["title"]), styles["head"], anchor), picture]))
            mapping.append(_location(node, "raster_fallback", f"named_destination:{anchor}", editable=False,
                                     fallback_approval="flow view policy"))
        elif node["kind"] == "image":
            bound = images[node["id"]]
            width, height = bound.fit(content_width, 6.2*72, units_per_pixel=72/96)
            picture = Image(BytesIO(bound.data), width=width, height=height)
            picture.hAlign = "LEFT"
            component = [tracked(escape(node["title"]), styles["head"], anchor), picture, Spacer(1, 7)]
            if node["alt_text"] and not node.get("decorative", False):
                component.append(Paragraph(escape(node["alt_text"]), styles["small"]))
            story.append(KeepTogether(component))
            mapping.append(_location(node, "embedded_image", f"named_destination:{anchor}",
                                     editable=False, placement_editable=False, alt_text_support="manifest_and_visible_caption",
                                     tagged_pdf=False, rendered_size={"width": width, "height": height, "unit": "pt"}, **bound.manifest(node)))
    def footer(canvas, document):
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#"+MUTED))
        canvas.drawString(margin, 22, f"{_period_caption(snapshot)}   Revision {snapshot['revision']}")
        canvas.drawRightString(A4[0]-margin, 22, str(document.page))
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    writer = PdfWriter()
    writer.append(PdfReader(path))
    for identifier, page_index in node_pages.items():
        writer.add_named_destination(identifier, page_index)
    with path.open("wb") as destination:
        writer.write(destination)
    reader = PdfReader(path)
    if set(reader.named_destinations) != {re.sub(r"[^A-Za-z0-9_]", "_", n["id"]) for n in nodes}:
        raise ExportError("PDF component destinations could not be verified.", "render_coverage_mismatch")
    actual = " ".join(page.extract_text() for page in reader.pages)
    normalize = lambda text: re.sub(r"\s+", " ", text).strip()
    for node in nodes:
        if node["kind"] == "rich_text" and normalize(_text(node, snapshot["facts"])) not in normalize(actual):
            raise ExportError("PDF text verification failed.", "render_content_mismatch")
    return mapping, [{"check": "pdf_readback_and_accepted_text", "status": "pass"}], "reportlab", {
        "page_count": len(reader.pages), "pdf_origin": "native_flow_reportlab", "office_pagination_equivalence": False,
        "font_profile": "Helvetica / Windows Latin", "margin_mm": margin*25.4/72, "paragraph_indent_pt": indent}


def _pptx(snapshot, nodes, path, policy, view, images):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.core_properties.title = snapshot["title"]
    prs.core_properties.subject = f"Frozen snapshot {snapshot['id']} revision {snapshot['revision']}"
    mapping = []
    def textbox(slide, text, x, y, width, height, size=16, bold=False, color=INK):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(.02)
        tf.margin_top = tf.margin_bottom = Inches(.02)
        p = tf.paragraphs[0]
        p.text = text
        p.font.name, p.font.size, p.font.bold = "Calibri", Pt(size), bold
        p.font.color.rgb = RGBColor.from_string(color)
        return shape
    def add_slide(title):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        textbox(slide, title, .6, .4, 12, .55, 28, True)
        textbox(slide, _period_caption(snapshot), .6, 1.02, 12, .3, 12, color=MUTED)
        textbox(slide, f"Revision {snapshot['revision']}", .6, 7.1, 11, .2, 9, color=MUTED)
        textbox(slide, str(len(prs.slides)), 12, 7.1, .4, .2, 9, color=MUTED)
        return slide
    def table(slide, headers, rows, x, y, width, height):
        shape = slide.shapes.add_table(len(rows)+1, len(headers), Inches(x), Inches(y), Inches(width), Inches(height))
        for r, row in enumerate([headers] + rows):
            for c, value in enumerate(row):
                cell = shape.table.cell(r,c)
                cell.text = str(value)
                cell.margin_left = cell.margin_right = Inches(.09)
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor.from_string(INK if r == 0 else (PALE if r % 2 else "FFFFFF"))
                for p in cell.text_frame.paragraphs:
                    p.font.name = "Calibri"
                    p.font.size = Pt(12 if len(headers) > 3 else 14)
                    p.font.bold = r == 0
                    p.font.color.rgb = RGBColor.from_string("FFFFFF" if r == 0 else INK)
        return shape
    first = add_slide(snapshot["title"])
    second = add_slide("Regional analysis")
    text_nodes = [n for n in nodes if n["kind"] == "rich_text"]
    y = 1.52
    for node in text_nodes:
        content = _text(node, snapshot["facts"])
        if len(content) > 800:
            raise ExportError("The canvas profile requires a separate layout for long prose.", "canvas_text_overflow")
        height = max(.45, math.ceil(len(content)/130) * .25)
        textbox(first, content, .6, y, 12.05, height, 16 if node["id"] == "summary" else 13, color=INK if node["id"] == "summary" else MUTED)
        mapping.append(_location(node, "native_text", f"slide:1/shape:{len(first.shapes)}", editable=True))
        y += height + .18
    for node in [n for n in nodes if n["kind"] == "table"]:
        headers, rows, _, _ = _table(snapshot, node)
        first_capacity = max(1, min(8, math.floor((6.7-y)/.45)-1))
        chunks = [rows[:first_capacity]]
        chunks.extend(rows[i:i+10] for i in range(first_capacity, len(rows), 10))
        locations = []
        for index, chunk in enumerate(chunks):
            slide = first if index == 0 else add_slide(f"{node['title']} continued")
            table_y = y if index == 0 else 1.6
            shape = table(slide, headers, chunk, .6, table_y, 12.05, min(4.8, .5*(len(chunk)+1)))
            slide_number = list(prs.slides).index(slide)+1
            locations.append(f"slide:{slide_number}/shape:{shape.shape_id}")
        mapping.append(_location(node, "native_table", locations, editable=True))
    for node in [n for n in nodes if n["kind"] == "chart"]:
        categories, series = _chart_data(snapshot, node)
        data = CategoryChartData()
        data.categories = categories
        for name, values in series:
            data.add_series(name, values)
        chart = second.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(.5), Inches(1.6), Inches(6.3), Inches(4.75), data).chart
        chart.has_title = True
        chart.chart_title.text_frame.text = node["title"] + f" ({node['axis_unit']})"
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.font.size = Pt(11)
        chart.category_axis.tick_labels.font.size = Pt(12)
        chart.value_axis.tick_labels.font.size = Pt(10)
        chart.value_axis.tick_labels.number_format = '#,##0'
        for index, series_object in enumerate(chart.series):
            series_object.format.fill.solid()
            series_object.format.fill.fore_color.rgb = RGBColor.from_string(GREEN if index == 0 else "A3B9AF")
        mapping.append(_location(node, "native_chart", "slide:2/chart:1", editable=True,
                                 embedded_data="approved regional aggregates"))
    for node in [n for n in nodes if n["kind"] == "pivot"]:
        headers, rows, _, _ = _table(snapshot, node, True)
        locations = []
        for index in range(0, len(rows), 9):
            chunk = rows[index:index+9]
            slide = second if index == 0 else add_slide(f"{node['title']} continued")
            x, width = (7, 5.6) if index == 0 else (.6, 12.05)
            textbox(slide, node["title"], x, 1.58, width, .5, 18, True)
            shape = table(slide, headers, chunk, x, 2.18, width, .42*(len(chunk)+1))
            textbox(slide, "Static pivot result. Native interaction is a workbook capability.", x, 6.6, width, .32, 10, color=MUTED)
            locations.append(f"slide:{list(prs.slides).index(slide)+1}/shape:{shape.shape_id}")
        mapping.append(_location(node, "equivalent_static_table", locations, editable=True,
                                 native_behavior="static result", fallback_approval="canvas view policy"))
    for node in [n for n in nodes if n["kind"] == "image"]:
        bound = images[node["id"]]
        slide = add_slide(node["title"])
        width, height = bound.fit(12.05, 5.1, units_per_pixel=1/96)
        x, y = .6+(12.05-width)/2, 1.6+(5.1-height)/2
        picture = slide.shapes.add_picture(BytesIO(bound.data), Inches(x), Inches(y), Inches(width), Inches(height))
        set_picture_properties(picture._element.nvPicPr.cNvPr, node)
        mapping.append(_location(node, "native_image", f"slide:{len(prs.slides)}/shape:{picture.shape_id}",
                                 editable=False, placement_editable=True, alt_text_support="drawing_properties",
                                 rendered_size={"width": width, "height": height, "unit": "in"}, **bound.manifest(node)))
    for slide in prs.slides:
        for shape in slide.shapes:
            if min(shape.left, shape.top) < 0 or shape.left+shape.width > prs.slide_width or shape.top+shape.height > prs.slide_height:
                raise ExportError("A slide element exceeds the canvas.", "canvas_bounds")
    prs.save(path)
    readback = Presentation(path)
    text = "\n".join(shape.text for slide in readback.slides for shape in slide.shapes if shape.has_text_frame)
    for node in text_nodes:
        if _text(node, snapshot["facts"]) not in text:
            raise ExportError("Presentation text verification failed.", "render_content_mismatch")
    return mapping, [{"check": "native_chart_and_text_roundtrip", "status": "pass"},
                     {"check": "shape_bounds", "status": "pass"},
                     {"check": "target_application_text_layout", "status": "not_certified"}], "python-pptx", {"slide_count": len(prs.slides)}


def _verify_relationships(path):
    """No external linked content is permitted in these generated core files."""
    if path.suffix == ".pdf":
        return
    with ZipFile(path) as z:
        if z.testzip() is not None:
            raise ExportError("The export ZIP failed its integrity check.", "corrupt_export")
        for name in z.namelist():
            if name.endswith(".rels"):
                for rel in ET.fromstring(z.read(name)):
                    if rel.attrib.get("TargetMode") == "External":
                        raise ExportError("The generated export contains an external relationship.", "external_resource")


def export_snapshot(snapshot: dict, format: str, output_dir: Path, policy="compatible", *,
                    asset_resolver: Callable[[str], bytes] | None = None) -> dict:
    """Render bytes and a fidelity manifest from the same frozen snapshot.

    compatible includes a native XLSX pivot with an explicit pending consumer
    certification. strict blocks that requirement. static is the user's
    explicit permission to replace native pivot behavior with a displayed table.
    """
    if hasattr(snapshot, "model_dump"):
        snapshot = snapshot.model_dump(mode="json")
    # Copy JSON state so rendering cannot mutate the accepted object.
    snapshot = json.loads(json.dumps(snapshot))
    from foundry.contracts import Snapshot
    try:
        Snapshot.model_validate(snapshot)
    except ValueError as error:
        raise ExportError("The native snapshot failed its schema or reference checks.", "invalid_snapshot") from error
    fmt = format.lower().lstrip(".")
    view, nodes, warnings = _plan(snapshot, fmt, policy)
    images = resolve_images(snapshot, nodes, asset_resolver)
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    identifier = re.sub(r"[^A-Za-z0-9_-]", "-", str(snapshot["id"]))[:70] or "report"
    draft_prefix = "draft-" if snapshot.get("status") == "review_required" else ""
    filename = f"{draft_prefix}{identifier}-r{snapshot['revision']}.{fmt}"
    path = output_dir / filename
    if path.exists():
        raise ExportError("An export with this filename already exists in this job directory.", "artifact_exists", 409)
    renderer = {"docx": _docx, "xlsx": _xlsx, "pdf": _pdf, "pptx": _pptx}[fmt]
    try:
        mapping, validations, package, details = renderer(snapshot, nodes, path, policy, view, images)
        if {entry["node_id"] for entry in mapping} != set(view["node_ids"]):
            raise ExportError("Rendered coverage differs from the planned view.", "render_coverage_mismatch")
        _verify_relationships(path)
        validations.extend([{"check": "per_component_coverage", "status": "pass"},
                            {"check": "no_external_relationships", "status": "pass"}])
        if images:
            validations.extend([{"check": "frozen_image_digests_and_decoded_dimensions", "status": "pass"},
                                {"check": "image_aspect_ratio_no_crop", "status": "pass"}])
            if fmt == "pdf":
                validations.append({"check": "tagged_image_accessibility", "status": "not_certified"})
        data = path.read_bytes()
        version = importlib.metadata.version(package)
        manifest = {
            "schema_version": "1.0", "state": "verified_with_limitations" if any(v["status"] == "not_certified" for v in validations) else "verified",
            "snapshot_id": snapshot["id"], "snapshot_revision": snapshot["revision"],
            "snapshot_digest": _digest(snapshot), "program": snapshot["program"],
            "source_snapshot_digest": snapshot["source_snapshot_digest"],
            "format": fmt, "view_id": view["id"], "view_family": view["family"], "fidelity_policy": policy,
            "renderer": {"name": package, "version": version, "adapter_version": VERSION,
                         "adapter_sha256": hashlib.sha256(Path(__file__).read_bytes() + Path(__file__).with_name("pivot.py").read_bytes() + Path(__file__).with_name("images.py").read_bytes()).hexdigest()},
            "artifact": {"filename": filename, "media_type": MEDIA_TYPES[fmt], "sha256": hashlib.sha256(data).hexdigest(), "size_bytes": len(data)},
            "coverage": view["coverage"], "components": mapping, "validation_results": validations,
            "warnings": warnings, "details": details,
        }
        manifest_path = output_dir / f"{filename}.manifest.json"
        manifest_path.write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")
        return {"path": str(path), "filename": filename, "media_type": MEDIA_TYPES[fmt],
                "manifest": manifest, "manifest_path": str(manifest_path),
                "sha256": manifest["artifact"]["sha256"], "size_bytes": len(data)}
    except Exception as error:
        path.unlink(missing_ok=True)
        if isinstance(error, ExportError):
            raise
        raise ExportError("The renderer could not lay out this snapshot under the selected profile.", "renderer_failed") from error


__all__ = ["ExportError", "export_snapshot"]
